# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""Contact imports: CSV / Excel files and the Leads desk.

A file import is two steps: ``preview`` parses the file, guesses which column
is which field and keeps the rows aside; ``commit`` applies the (possibly
edited) mapping. Nothing touches the book until commit.
"""

from __future__ import annotations

import csv
import io
import json
import re
import time
import unicodedata
import uuid
from typing import Any

from navin.campaigns.crm_source import import_people
from navin.campaigns.db import campaigns_root
from navin.campaigns.store import CampaignError

MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
TARGETS = ("email", "phone", "first_name", "last_name", "full_name", "company", "title", "city",
           "country", "tags", "ignore")

_GUESS: dict[str, tuple[str, ...]] = {
    "email": ("email", "e-mail", "mail", "courriel", "adresse email", "email address", "البريد"),
    "phone": ("phone", "telephone", "tel", "mobile", "portable", "whatsapp", "numero", "gsm", "هاتف"),
    "first_name": ("first name", "firstname", "prenom", "given name", "الاسم الأول"),
    "last_name": ("last name", "lastname", "nom de famille", "surname", "family name", "nom"),
    "full_name": ("name", "full name", "contact", "nom complet", "الاسم"),
    "company": ("company", "societe", "entreprise", "organisation", "organization", "account", "الشركة"),
    "title": ("title", "job title", "poste", "fonction", "role", "position"),
    "city": ("city", "ville", "المدينة"),
    "country": ("country", "pays", "البلد"),
    "tags": ("tags", "tag", "labels", "segment"),
}


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode() or str(text or "")
    return re.sub(r"[_\s]+", " ", text.strip().lower())


def guess_mapping(columns: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for column in columns:
        key = _norm(column)
        target = "ignore"
        for field, names in _GUESS.items():
            if field in used and field != "tags":
                continue
            if key in names or any(key == _norm(n) for n in names):
                target = field
                break
        if target == "ignore":
            for field, names in _GUESS.items():
                if field not in used and any(_norm(n) in key for n in names if len(n) > 3):
                    target = field
                    break
        if target == "ignore" and key:
            target = f"attrs.{re.sub(r'[^a-z0-9]+', '_', key).strip('_')[:40] or 'extra'}"
        mapping[column] = target
        if target in TARGETS:
            used.add(target)
    return mapping


def _read_rows(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    if not data:
        raise CampaignError("The file is empty")
    if len(data) > MAX_IMPORT_BYTES:
        raise CampaignError("Import files are limited to 20 MB")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - openpyxl is a core dependency
            raise CampaignError("Excel support is not installed") from exc
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise CampaignError("This Excel file could not be read") from exc
        rows = [["" if v is None else str(v).strip() for v in row] for row in sheet.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt", ".tsv")) or not name:
        text = None
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        text = text or ""
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    elif name.endswith(".xls"):
        raise CampaignError("Old .xls files are not supported: save the sheet as .xlsx or .csv")
    else:
        raise CampaignError("Import a .csv or .xlsx file")
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        raise CampaignError("The file has no rows")
    header = [c or f"Column {i + 1}" for i, c in enumerate(rows[0])]
    body = rows[1:MAX_ROWS + 1]
    width = len(header)
    return header, [(row + [""] * width)[:width] for row in body]


def _imports_dir():
    path = campaigns_root() / "imports"
    path.mkdir(parents=True, exist_ok=True)
    cutoff = time.time() - 24 * 3600
    for old in path.glob("*.json"):
        if old.stat().st_mtime < cutoff:
            old.unlink(missing_ok=True)
    return path


def preview(filename: str, data: bytes) -> dict[str, Any]:
    columns, rows = _read_rows(filename, data)
    import_id = uuid.uuid4().hex
    (_imports_dir() / f"{import_id}.json").write_text(
        json.dumps({"filename": filename, "columns": columns, "rows": rows}, ensure_ascii=False), encoding="utf-8",
    )
    return {"import_id": import_id, "filename": filename, "columns": columns, "sample": rows[:8],
            "rows": len(rows), "mapping": guess_mapping(columns)}


def rows_to_people(columns: list[str], rows: list[list[str]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    people = []
    for row in rows:
        person: dict[str, Any] = {"attrs": {}, "tags": [], "source": "import"}
        for column, value in zip(columns, row, strict=False):
            target = mapping.get(column, "ignore")
            if not value or target == "ignore":
                continue
            if target == "full_name":
                first, _, last = value.partition(" ")
                person.setdefault("first_name", first)
                person.setdefault("last_name", last)
            elif target == "tags":
                person["tags"] += [t.strip() for t in re.split(r"[,;|]", value) if t.strip()]
            elif target.startswith("attrs."):
                person["attrs"][target[6:]] = value
            elif target in TARGETS:
                person[target] = value
        people.append(person)
    return people


def commit(import_id: str, mapping: dict[str, str], *, list_id: str = "", new_list_name: str = "",
           tags: list[str] | None = None) -> dict[str, Any]:
    path = _imports_dir() / f"{re.sub(r'[^a-f0-9]', '', import_id)}.json"
    if not import_id or not path.is_file():
        raise CampaignError("This import expired; upload the file again", status=404)
    data = json.loads(path.read_text(encoding="utf-8"))
    targets = {v for v in mapping.values()}
    if not targets & {"email", "phone"}:
        raise CampaignError("Map at least one column to Email or Phone")
    people = rows_to_people(data["columns"], data["rows"], mapping)
    result = import_people(people, list_id=list_id, new_list_name=new_list_name, tags=tags)
    path.unlink(missing_ok=True)
    return result


def from_leads(*, list_id: str = "", new_list_name: str = "") -> dict[str, Any]:
    """Copy the Leads desk book (people with an email or phone) into Campaigns."""
    from navin.config.paths import get_runtime_subdir
    from navin.leads.store import LeadsStore

    leads = LeadsStore(get_runtime_subdir("leads")).load_leads()
    people = []
    for lead in leads if isinstance(leads, list) else []:
        email = lead.get("email") or ""
        phone = lead.get("phone") or ""
        if not email and not phone:
            continue
        first, _, last = str(lead.get("person") or lead.get("contact_name") or "").partition(" ")
        people.append({
            "email": email, "phone": phone, "first_name": lead.get("first_name") or first,
            "last_name": lead.get("last_name") or last, "company": lead.get("company") or lead.get("company_name") or "",
            "title": lead.get("title") or "", "city": lead.get("city") or "", "country": lead.get("country") or "",
            "source": "leads", "attrs": {k: str(lead[k]) for k in ("score", "stage", "domain", "industry") if lead.get(k)},
        })
    return import_people(people, list_id=list_id, new_list_name=new_list_name or ("" if list_id else "Leads desk"))
