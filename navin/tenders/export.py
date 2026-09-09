# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""Render a ready Word / PowerPoint pack from the written dossier."""

from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
import tempfile
import textwrap
from pathlib import Path
from typing import Any

from loguru import logger

from navin.tenders.bid_pack import SECTION_KEYS, section_title
from navin.tenders.profile import filed_documents

_NAVY = (27, 54, 93)
_SLATE = (51, 65, 85)
_GOLD = (184, 148, 90)
_ROW = "F4F7FA"
_LINE = "D0D7DE"


def _diagram_spec(tender: dict[str, Any], response: dict[str, Any]) -> dict[str, Any] | None:
    """A proposed delivery flow with labels and trace IDs derived from this notice."""
    fr = _lang_of(response) == "fr"
    rows = [row for row in response.get("requirement_responses") or [] if row.get("category") == "technical" and row.get("facets")]
    if not rows:
        return None
    text = "\n".join(str(row.get("text") or "") for row in rows)
    facets = {facet for row in rows for facet in row.get("facets") or []}

    def source_ids(facet: str) -> str:
        return ", ".join(str(row.get("buyer_id") or row["id"]) for row in rows if facet in row.get("facets", []))

    def tag(facet: str) -> str:
        ids = [str(row.get("buyer_id") or row["id"]) for row in rows if facet in row.get("facets", [])]
        return ids[0] if ids else ""

    count = re.search(r"\b\d+\s+sources?\b", text, re.I)
    technology = re.search(r"\b(?:PostgreSQL|Oracle|SQL Server|MySQL)\b", text, re.I)
    interface = "API REST" if re.search(r"\bAPI REST\b", text, re.I) else "Interfaces" if "api" in facets else ("Conception" if fr else "Design")
    recovery_pattern = r"\b(RPO|RTO)\s*(?:de|of|:|=)?\s*(\d+(?:[.,]\d+)?)\s*(heures?|hours?|minutes?|jours?|days?)"
    recovery_row = next((row for row in rows if re.search(recovery_pattern, str(row.get("text") or ""), re.I)), None)
    recovery = re.findall(recovery_pattern, str((recovery_row or {}).get("text") or ""), re.I)
    transfer_row = next((row for row in rows if re.search(r"organis[^.]{0,30}transfert|formation|mise en situation|training|handover", str(row.get("text") or ""), re.I)), None)
    recovery_label = " / ".join(f"{kind.upper()} {value}{'h' if unit.lower().startswith(('heure', 'hour')) else 'min' if unit.lower().startswith('minute') else 'j' if fr else 'd'}" for kind, value, unit in recovery)
    delivery_label = count.group(0) if count else ("Livrables" if fr else "Deliverables")
    nodes = [
        {"id": "bid_scope", "lane": "buyer", "col": 0, "type": "external", "label": "Cadrage" if fr else "Scope", "sublabel": "Exigences et acces" if fr else "Requirements and access", "tag": str(rows[0].get("buyer_id") or rows[0]["id"])},
        {"id": "bid_contracts", "lane": "delivery", "col": 1, "type": "backend", "label": interface, "sublabel": "Contrats a valider" if fr else "Contracts to approve", "tag": tag("api")},
        {"id": "bid_delivery", "lane": "delivery", "col": 2, "type": "backend", "label": delivery_label, "sublabel": technology.group(0) if technology else ("Preuves de tests" if fr else "Test evidence"), "tag": tag("data") or tag("quality")},
        {"id": "bid_acceptance", "lane": "buyer", "col": 3, "type": "security", "label": "Reprise" if recovery and fr else "Recovery" if recovery else "Recette" if fr else "Acceptance", "sublabel": recovery_label or ("Decision et reserves" if fr else "Decision and qualifications"), "tag": str((recovery_row or {}).get("buyer_id") or (recovery_row or {}).get("id") or tag("quality"))},
        {"id": "bid_handover", "lane": "operations", "col": 4, "type": "external", "label": "Transfert" if "transfer" in facets and fr else "Handover" if "transfer" in facets else "Validation", "sublabel": "Supports editables" if fr else "Editable materials", "tag": str((transfer_row or {}).get("buyer_id") or (transfer_row or {}).get("id") or tag("transfer"))},
    ]
    for node in nodes:
        if not node["tag"]:
            node.pop("tag")
    links = [row["id"] for row in nodes]
    cards = []
    for facet, label in (("data", "Donnees" if fr else "Data"), ("api", "Interfaces"), ("continuity", "Continuite" if fr else "Continuity")):
        ids = source_ids(facet)
        if ids:
            cards.append({"dot": "emerald" if facet == "data" else "amber", "title": label, "items": [ids, "Voir les criteres dans les fiches correspondantes." if fr else "See acceptance criteria in the linked responses."]})
    return {
        "schema_version": 2, "diagram_type": "workflow",
        "meta": {"title": "Deroulement propose et preuves" if fr else "Proposed delivery and evidence", "subtitle": _clean(tender.get("title") or ""), "animation": "none", "quality_profile": "showcase", "legend": {"mode": "hidden"}},
        "lanes": [{"id": "buyer", "label": "Acheteur" if fr else "Buyer"}, {"id": "delivery", "label": "Equipe proposee" if fr else "Proposed team"}, {"id": "operations", "label": "Destinataires" if fr else "Recipients"}],
        "mainPath": links, "semanticChecks": {"allowedRoots": [links[0]], "allowedTerminals": [links[-1]], "requiredPaths": [{"from": links[0], "to": links[-1]}]},
        "nodes": nodes, "edges": [{"from": left, "to": right, "role": "main"} for left, right in zip(links, links[1:])], "cards": cards,
    }


def _standalone_archify_svg(html: str) -> bytes:
    """Retain validated SVG geometry and resolve Archify's own light-theme classes."""
    match = re.search(r"<svg\b[\s\S]*?</svg>", html)
    css = re.search(r"<style[^>]*>([\s\S]*?)</style>", html)
    theme = re.search(r'\[data-theme="light"\]\s*\{([^}]+)\}', html)
    if not match or not css or not theme:
        raise ValueError("Archify SVG or theme missing")
    svg = match.group(0)
    variables = dict(re.findall(r"(--[\w-]+)\s*:\s*([^;]+);", theme.group(1)))
    classes = {token for value in re.findall(r'class="([^"]+)"', svg) for token in value.split()}
    styles = []
    for cls in sorted(classes):
        rule = re.search(r"(?:^|\n)\s*(?:svg\s+)?\." + re.escape(cls) + r"\s*\{([^}]+)\}", css.group(1))
        if rule:
            resolved = re.sub(r"var\((--[\w-]+)\)", lambda item: variables.get(item.group(1), "#64748b"), rule.group(1))
            styles.append(f".{cls}{{{resolved}}}")
    if "xmlns=" not in svg.split(">", 1)[0]:
        svg = svg.replace("<svg ", '<svg xmlns="http://www.w3.org/2000/svg" ', 1)
    svg = svg.replace(">", "><style>text{font-family:Arial,sans-serif;}" + "".join(styles) + "</style>", 1)
    return svg.encode("utf-8")


def _attach_diagram_exports(store: Any, tender: dict[str, Any], response: dict[str, Any], exports: dict[str, Any]) -> None:
    """Run the bundled Archify renderer, never a provider, browser or remote service."""
    info = response["export_generation"]
    fr = _lang_of(response) == "fr"
    if "archify" not in info.get("skills", {}).get("loaded", []):
        info["warnings"].append("Archify indisponible ou desactive: diagramme non genere." if fr else "Archify unavailable or disabled: diagram not generated.")
        return
    spec = _diagram_spec(tender, response)
    if spec is None:
        info["warnings"].append("Cahier des charges trop incomplet pour un diagramme technique: preciser les exigences." if fr else "Specification too incomplete for a technical diagram: clarify requirements.")
        return
    node = shutil.which("node")
    cli = Path(__file__).resolve().parents[1] / "skills" / "archify" / "bin" / "archify.mjs"
    if not node or not cli.is_file():
        info["warnings"].append("Moteur Archify local absent: diagramme non genere." if fr else "Local Archify engine missing: diagram not generated.")
        return
    tid = _safe_name(str(tender.get("id") or "notice"), "notice")
    try:
        with tempfile.TemporaryDirectory(prefix="navin-tender-archify-") as tmp:
            source = Path(tmp) / "workflow.json"
            target = Path(tmp) / "workflow.html"
            source.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([node, str(cli), "deliver", "workflow", str(source), str(target), "--quality", "showcase", "--json"], capture_output=True, text=True, timeout=20, check=True)
            receipt = json.loads(result.stdout)
            validation = receipt.get("validation") or {}
            if not receipt.get("ok") or validation.get("checksPassed") != 9 or validation.get("warnings") != 0 or validation.get("errors") != 0:
                raise ValueError("Archify did not pass showcase validation")
            html = target.read_text(encoding="utf-8")
            svg = _standalone_archify_svg(html)
            for kind, data, mime in (("html", html.encode("utf-8"), "text/html"), ("svg", svg, "image/svg+xml"), ("json", source.read_bytes(), "application/json")):
                key = "diagram_spec" if kind == "json" else f"diagram_{kind}"
                record = store.save_bytes(f"{tid}_architecture.{kind}", data, file_id=f"out_archify_{kind}_{tid}")
                exports[key] = {**record, "kind": kind, "mime": mime}
            response["technical_diagram"] = {"type": "workflow", "scope": "proposed_delivery", "requirement_ids": [row["id"] for row in response.get("requirement_responses") or [] if row.get("category") == "technical"], "validation": validation, "specification": receipt.get("specification"), "artifact": receipt.get("artifact"), "visual_review": "not_performed", "caption": "Deroulement propose, rattache aux exigences source. Les choix et les preuves restent a valider." if fr else "Proposed delivery linked to source requirements. Choices and evidence still need approval."}
            try:
                import cairosvg

                png = cairosvg.svg2png(bytestring=svg, background_color="#ffffff", output_width=2000)
                record = store.save_bytes(f"{tid}_architecture.png", png, file_id=f"out_archify_png_{tid}")
                exports["diagram_png"] = {**record, "kind": "png", "mime": "image/png"}
            except (ImportError, OSError, ValueError):
                info["warnings"].append("Diagramme HTML/SVG genere; conversion PNG indisponible pour Word/PPT." if fr else "HTML/SVG diagram generated; PNG conversion unavailable for Word/PPT.")
    except (OSError, ValueError, subprocess.SubprocessError):
        info["warnings"].append("Diagramme Archify non valide ou generation locale echouee; revue requise." if fr else "Archify validation or local rendering failed; review required.")


def _diagram_png(store: Any, response: dict[str, Any]) -> bytes | None:
    record = (response.get("exports") or {}).get("diagram_png") or {}
    if not record.get("path"):
        return None
    path = store.files_dir / Path(str(record["path"])).name
    return path.read_bytes() if path.is_file() else None


def _safe_name(value: str, fallback: str) -> str:
    token = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "").strip())[:80]
    return token or fallback


def _lang_of(response: dict[str, Any]) -> str:
    lang = str(response.get("language") or "en").strip().lower()
    return "fr" if lang.startswith("fr") else "en"


def _clean(text: Any) -> str:
    return str(text or "").replace("\u2014", " - ").replace("\u2013", "-")


def _mapping(tender: dict[str, Any], profile: dict[str, Any], response: dict[str, Any]) -> dict[str, str]:
    fields = {
        "letter": response.get("letter"),
        "executive_summary": response.get("executive_summary"),
        "summary": response.get("executive_summary"),
        "cover": response.get("cover"),
        "toc": response.get("toc"),
        "company": response.get("company"),
        "need": response.get("need"),
        "approach": response.get("approach"),
        "vision": response.get("vision"),
        "functional": response.get("functional"),
        "architecture": response.get("architecture"),
        "methodology": response.get("methodology"),
        "followup_kpi": response.get("followup_kpi"),
        "raci_risks": response.get("raci_risks"),
        "governance": response.get("governance"),
        "planning": response.get("planning"),
        "staffing": response.get("staffing"),
        "references": response.get("references"),
        "compliance_matrix": response.get("compliance_matrix"),
        "matrix": response.get("compliance_matrix"),
        "financial_schedule": response.get("financial_schedule"),
        "title": tender.get("title"),
        "buyer": tender.get("buyer"),
        "deadline": tender.get("deadline"),
        "company_name": profile.get("name"),
        "revision_notes": response.get("revision_notes"),
        "evidence": response.get("evidence"),
        "assumptions": response.get("assumptions"),
    }
    out: dict[str, str] = {}
    for key, value in fields.items():
        text = _clean(value).strip()
        out[f"{{{{{key}}}}}"] = text
        out[f"{{{{{key.upper()}}}}}"] = text
    return out


def _replace_text(text: str, mapping: dict[str, str]) -> tuple[str, bool]:
    changed = False
    body = text or ""
    for key, value in mapping.items():
        if key in body:
            body = body.replace(key, value)
            changed = True
    return body, changed


def _first_template(profile: dict[str, Any], store: Any, bucket: str) -> Path | None:
    for row in filed_documents(profile):
        if row.get("bucket") != bucket or not row.get("path"):
            continue
        path = store.files_dir / Path(str(row["path"])).name
        if path.is_file():
            return path
    return None


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _fill_docx(doc: Any, mapping: dict[str, str]) -> bool:
    changed = False
    for paragraph in doc.paragraphs:
        new, hit = _replace_text(paragraph.text, mapping)
        if hit:
            _set_paragraph_text(paragraph, new)
            changed = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    new, hit = _replace_text(paragraph.text, mapping)
                    if hit:
                        _set_paragraph_text(paragraph, new)
                        changed = True
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                new, hit = _replace_text(paragraph.text, mapping)
                if hit:
                    _set_paragraph_text(paragraph, new)
                    changed = True
    return changed


def doc_mod_pt(size: int) -> Any:
    from docx.shared import Pt

    return Pt(size)


def _style_run(run: Any, *, size: int = 11, bold: bool = False, color: Any = None, name: str = "Calibri") -> None:
    run.font.name = name
    run.font.size = doc_mod_pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    try:
        from docx.oxml.ns import qn

        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        run._element.rPr.rFonts.set(qn("w:ascii"), name)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    except Exception:
        pass


def _apply_theme(doc: Any) -> None:
    from docx.shared import Pt, RGBColor

    try:
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor(*_SLATE)
        normal.paragraph_format.line_spacing = 1.12
        normal.paragraph_format.widow_control = True
    except Exception:
        pass
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        try:
            style = doc.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(*_NAVY)
            style.font.italic = False
            style.paragraph_format.keep_with_next = True
        except Exception:
            pass


def _shade_cell(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_borders(cell: Any, color: str = _LINE) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), color)
        tc_borders.append(item)
    tc_pr.append(tc_borders)


def _add_heading_rule(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1B365D")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_docx_heading(doc: Any, text: str, *, level: int = 1) -> None:
    from docx.shared import Pt, RGBColor

    try:
        heading = doc.add_heading(_clean(text), level=level)
        heading.paragraph_format.space_before = Pt(16)
        heading.paragraph_format.space_after = Pt(8)
        heading.paragraph_format.keep_with_next = True
        for run in heading.runs:
            _style_run(run, size=16 if level == 1 else 13, bold=True, color=RGBColor(*_NAVY))
        _add_heading_rule(heading)
    except Exception:
        doc.add_paragraph(_clean(text))


def _add_cover_page(doc: Any, tender: dict[str, Any], profile: dict[str, Any], response: dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.different_first_page_header_footer = True
    lang = _lang_of(response)
    company = _clean(profile.get("name") or ("Candidat" if lang == "fr" else "Bidder"))
    title = _clean(tender.get("title") or ("Avis" if lang == "fr" else "Notice"))
    buyer = _clean(tender.get("buyer") or "")
    ref = _clean(tender.get("reference") or tender.get("notice_id") or "")
    country = _clean(tender.get("country") or "")
    deadline = _clean(tender.get("deadline") or "")
    band = doc.add_table(rows=1, cols=1)
    cell = band.cell(0, 0)
    _shade_cell(cell, "1B365D")
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(10)
    run = para.add_run("DOSSIER DE CANDIDATURE" if lang == "fr" else "SUBMISSION DOSSIER")
    _style_run(run, size=13, bold=True, color=RGBColor(255, 255, 255))
    accent = doc.add_table(rows=1, cols=1)
    gold = accent.cell(0, 0)
    _shade_cell(gold, "B8945A")
    gold.paragraphs[0].add_run(" ")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(36)
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(company)
    _style_run(run, size=28, bold=True, color=RGBColor(*_NAVY))
    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_after = Pt(18)
    run = notice.add_run(title)
    _style_run(run, size=16, bold=False, color=RGBColor(*_SLATE))
    labels = (
        (("Acheteur", buyer), ("Reference", ref), ("Pays", country), ("Date limite", deadline))
        if lang == "fr"
        else (
            ("Contracting authority", buyer),
            ("Reference", ref),
            ("Country", country),
            ("Deadline", deadline),
        )
    )
    meta = doc.add_table(rows=len(labels), cols=2)
    for index, (label, value) in enumerate(labels):
        left = meta.cell(index, 0)
        right = meta.cell(index, 1)
        left.text = label
        right.text = value or ("non renseigne" if lang == "fr" else "not on file")
        _shade_cell(left, "1B365D")
        _shade_cell(right, _ROW if index % 2 else "FFFFFF")
        _set_cell_borders(left, "1B365D")
        _set_cell_borders(right)
        for paragraph in left.paragraphs:
            for run in paragraph.runs:
                _style_run(run, size=11, bold=True, color=RGBColor(255, 255, 255))
        for paragraph in right.paragraphs:
            for run in paragraph.runs:
                _style_run(run, size=11, color=RGBColor(*_NAVY))
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(28)
    run = note.add_run(
        "Projet de memoire technique et financier - validation et justificatifs requis avant depot."
        if lang == "fr"
        else "Draft technical and financial memorandum - approval and supporting evidence required before submission."
    )
    _style_run(run, size=10, color=RGBColor(*_SLATE))
    try:
        doc.add_page_break()
    except Exception:
        pass


def _style_data_table(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    for r_i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if r_i == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        for cell in row.cells:
            _set_cell_borders(cell)
            if r_i == 0:
                _shade_cell(cell, "1B365D")
                color = RGBColor(255, 255, 255)
                bold = True
            else:
                _shade_cell(cell, _ROW if r_i % 2 == 0 else "FFFFFF")
                color = RGBColor(*_SLATE)
                bold = False
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.keep_with_next = False
                for run in paragraph.runs:
                    _style_run(run, size=10, bold=bold, color=color)


def _add_docx_table(doc: Any, body: str) -> bool:
    rows = [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in _clean(body).splitlines()
        if line.strip().startswith("|") and not re.fullmatch(r"[| :\-]+", line.strip())
    ]
    if len(rows) < 2:
        return False
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    for r_i, row in enumerate(rows):
        for c_i in range(width):
            table.rows[r_i].cells[c_i].text = row[c_i] if c_i < len(row) else ""
    _style_data_table(table)
    doc.add_paragraph("")
    return True


def _add_docx_body(doc: Any, text: str) -> None:
    from docx.shared import Pt, RGBColor

    body = _clean(text).strip()
    if not body:
        return
    lines = body.splitlines()
    cursor = 0
    while cursor < len(lines):
        row = lines[cursor].strip()
        cursor += 1
        if not row:
            continue
        if row.startswith("|"):
            table_lines = [row]
            while cursor < len(lines) and lines[cursor].strip().startswith("|"):
                table_lines.append(lines[cursor])
                cursor += 1
            if _add_docx_table(doc, "\n".join(table_lines)):
                continue
        if re.match(r"^#{2,3}\s+", row):
            _add_docx_heading(doc, re.sub(r"^#+\s+", "", row), level=2)
            continue
        if re.match(r"^[-*]\s+", row):
            para = doc.add_paragraph(row[2:].strip(), style="List Bullet")
        else:
            # Keep buyer numbering verbatim; Word's automatic lists renumber source clauses.
            para = doc.add_paragraph(row)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.widow_control = True
        if row.endswith(":" ) and len(row) < 100:
            para.paragraph_format.keep_with_next = True
        for run in para.runs:
            _style_run(run, size=11, bold=row.endswith(":") and len(row) < 100, color=RGBColor(*_SLATE))


def _add_toc(doc: Any, response: dict[str, Any]) -> None:
    """A real editable TOC field with a readable cached list and update-on-open."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    lang = _lang_of(response)
    _add_docx_heading(doc, section_title("toc", lang))
    paragraph = doc.add_paragraph()
    for kind in ("begin", "instruction", "separate"):
        element = OxmlElement("w:instrText" if kind == "instruction" else "w:fldChar")
        if kind == "instruction":
            element.set(qn("xml:space"), "preserve")
            element.text = ' TOC \\o "1-2" \\h \\z \\u '
        else:
            element.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                element.set(qn("w:dirty"), "true")
        paragraph.add_run()._r.append(element)
    titles = [section_title(key, lang) for key in SECTION_KEYS if key not in {"cover", "toc"} and response.get(key)]
    paragraph.add_run("\n".join(titles))
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)


def _add_header_footer(doc: Any, tender: dict[str, Any], profile: dict[str, Any], response: dict[str, Any]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    lang = _lang_of(response)
    company = _clean(profile.get("name") or "")
    title = _clean(tender.get("title") or "")[:70]
    label = "Confidentiel" if lang == "fr" else "Confidential"
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "  ·  ".join(part for part in (company, title) if part)
        for run in header.runs:
            _style_run(run, size=9, color=RGBColor(*_SLATE))
        _add_heading_rule(header)
        footer = section.footer.paragraphs[0]
        footer.text = f"{label}  ·  "
        for run in footer.runs:
            _style_run(run, size=9, color=RGBColor(*_SLATE))
        try:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), "begin")
            instr = OxmlElement("w:instrText")
            instr.text = " PAGE "
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            run = footer.add_run()
            run._r.append(fld)
            run._r.append(instr)
            run._r.append(end)
        except Exception:
            pass


def _write_docx_sections(doc: Any, tender: dict[str, Any], profile: dict[str, Any], response: dict[str, Any], *, skip: set[str] | None = None, figure: bytes | None = None) -> None:
    lang = _lang_of(response)
    skip = skip or set()
    if "cover" not in skip:
        _add_cover_page(doc, tender, profile, response)
    toc = _clean(response.get("toc") or "").strip()
    if toc:
        _add_toc(doc, response)
        try:
            doc.add_page_break()
        except Exception:
            pass
    for key in SECTION_KEYS:
        if key in {"cover", "toc"} or key in skip:
            continue
        body = _clean(response.get(key) or "").strip()
        if not body:
            continue
        if key in {"architecture", "functional", "compliance_matrix", "evidence"}:
            doc.add_page_break()
        _add_docx_heading(doc, section_title(key, lang))
        if key == "architecture" and figure:
            from docx.shared import Cm

            shape = doc.add_picture(io.BytesIO(figure), width=Cm(16.5))
            caption = _clean((response.get("technical_diagram") or {}).get("caption") or "")
            shape._inline.docPr.set("descr", caption)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            paragraph = doc.add_paragraph(caption, style="Caption")
            paragraph.paragraph_format.keep_with_next = False
        _add_docx_body(doc, body)
    _add_header_footer(doc, tender, profile, response)


def render_docx(
    store: Any,
    tender: dict[str, Any],
    profile: dict[str, Any],
    response: dict[str, Any],
) -> bytes | None:
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx missing - Word pack skipped")
        return None
    mapping = _mapping(tender, profile, response)
    src = _first_template(profile, store, "word")
    doc = None
    represented: set[str] = set()
    if src is not None:
        try:
            doc = Document(str(src))
            template_text = "\n".join(node.text or "" for node in doc.element.iter() if node.tag.endswith("}t"))
            for key in SECTION_KEYS:
                if f"{{{{{key}}}}}" in template_text or f"{{{{{key.upper()}}}}}" in template_text:
                    represented.add(key)
            for alias, key in (("summary", "executive_summary"), ("matrix", "compliance_matrix")):
                if f"{{{{{alias}}}}}" in template_text or f"{{{{{alias.upper()}}}}}" in template_text:
                    represented.add(key)
            # Never substitute a plain-text contents placeholder for the TOC field.
            toc_mapping = {**mapping, "{{toc}}": "", "{{TOC}}": ""}
            _fill_docx(doc, toc_mapping)
        except Exception as exc:
            logger.warning("Word model {} could not be opened ({}) - generating a clean dossier", src.name, exc)
            doc = None
            represented = set()
    if doc is None:
        doc = Document()
        _apply_theme(doc)
    elif doc.paragraphs or doc.tables:
        doc.add_page_break()
    _write_docx_sections(doc, tender, profile, response, skip=represented, figure=_diagram_png(store, response))
    try:
        doc.core_properties.title = _clean(tender.get("title") or "")
        doc.core_properties.author = _clean(profile.get("name") or "")
        doc.core_properties.subject = "Projet a valider" if _lang_of(response) == "fr" else "Draft for review"
    except Exception:
        pass
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _fill_pptx_text(prs: Any, mapping: dict[str, str]) -> bool:
    changed = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                full = "".join(run.text or "" for run in paragraph.runs) or paragraph.text
                new, hit = _replace_text(full, mapping)
                if not hit:
                    continue
                if paragraph.runs:
                    paragraph.runs[0].text = new
                    for run in paragraph.runs[1:]:
                        run.text = ""
                else:
                    paragraph.text = new
                changed = True
    return changed


def _pptx_lines(body: str, limit: int | None = None) -> list[str]:
    rows: list[str] = []
    for raw in _clean(body).splitlines():
        line = raw.strip()
        if not line or line.startswith("| ---"):
            continue
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|") if cell.strip()]
            if cells:
                rows.append(" · ".join(cells))
            continue
        rows.append(re.sub(r"^(?:[-*]|#{1,3})\s+", "", line))
    return rows


def _pptx_pages(body: str) -> list[str]:
    """Paginate editable text instead of discarding all but the first lines."""
    pages: list[str] = []
    current: list[str] = []
    used = 0
    for row in _pptx_lines(body):
        wrapped = textwrap.wrap(row, width=105, break_long_words=False, break_on_hyphens=False) or [row]
        for start in range(0, len(wrapped), 12):
            chunk = wrapped[start:start + 12]
            cost = len(chunk) + 1
            if current and used + cost > 16:
                pages.append("\n".join(current))
                current, used = [], 0
            current.append("\n".join(chunk))
            used += cost
    if current:
        pages.append("\n".join(current))
    return pages


def _add_pptx_cover(prs: Any, company: str, title: str, buyer: str, deadline: str, lang: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    width = prs.slide_width
    height = prs.slide_height
    fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, height)
    fill.fill.solid()
    fill.fill.fore_color.rgb = RGBColor(*_NAVY)
    fill.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, height - Inches(0.18), width, Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*_GOLD)
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), width - Inches(1.4), height - Inches(2.0))
    frame = box.text_frame
    frame.word_wrap = True
    kicker = frame.paragraphs[0]
    kicker.alignment = PP_ALIGN.LEFT
    run = kicker.add_run()
    run.text = "DOSSIER DE CANDIDATURE" if lang == "fr" else "SUBMISSION DOSSIER"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*_GOLD)
    name = frame.add_paragraph()
    run = name.add_run()
    run.text = company
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    notice = frame.add_paragraph()
    run = notice.add_run()
    run.text = title
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(226, 232, 240)
    meta = frame.add_paragraph()
    run = meta.add_run()
    run.text = "  ·  ".join(part for part in (buyer, deadline) if part)
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(203, 213, 225)


def _add_pptx_slide(prs: Any, title: str, body: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    width, height = prs.slide_width, prs.slide_height
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), width - Inches(1.2), Inches(0.8))
    heading = box.text_frame.paragraphs[0]
    heading.text = _clean(title)
    heading.font.size = Pt(25)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(*_NAVY)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.18), width - Inches(1.2), Inches(0.035))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(*_GOLD)
    rule.line.fill.background()
    body_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.45), width - Inches(1.2), height - Inches(1.95))
    frame = body_box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(_clean(body).splitlines()):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(16)
        paragraph.font.color.rgb = RGBColor(*_SLATE)
        paragraph.space_after = Pt(4)
        paragraph.line_spacing = 1.05


def render_pptx(
    store: Any,
    tender: dict[str, Any],
    profile: dict[str, Any],
    response: dict[str, Any],
) -> bytes | None:
    try:
        from pptx import Presentation
    except ImportError:
        logger.warning("python-pptx missing - PowerPoint pack skipped")
        return None
    mapping = _mapping(tender, profile, response)
    src = _first_template(profile, store, "ppt")
    prs = None
    filled = False
    if src is not None:
        try:
            prs = Presentation(str(src))
            filled = _fill_pptx_text(prs, mapping)
        except Exception as exc:
            logger.warning("PPT model {} could not be opened ({}) - generating a clean deck", src.name, exc)
            prs = None
            filled = False
    if prs is None or not filled:
        prs = Presentation()
        from pptx.util import Inches

        prs.slide_width = Inches(13.333333)
        prs.slide_height = Inches(7.5)
        lang = _lang_of(response)
        company = _clean(profile.get("name") or ("Candidat" if lang == "fr" else "Bidder"))
        title = _clean(tender.get("title") or ("Avis" if lang == "fr" else "Notice"))
        _add_pptx_cover(
            prs,
            company,
            title,
            _clean(tender.get("buyer") or ""),
            _clean(tender.get("deadline") or ""),
            lang,
        )
    lang = _lang_of(response)
    _add_pptx_slide(prs, "Lecture du support" if lang == "fr" else "Reading this deck", "Synthese de proposition pour revue. Le Word contient les fiches completes de mise en oeuvre et de recette, ainsi que le registre des pieces. Les reserves restent ouvertes avant depot." if lang == "fr" else "Proposal summary for review. The Word document contains complete implementation and acceptance responses and the evidence register. Qualifications remain open before submission.")
    figure = _diagram_png(store, response)
    if figure:
        from PIL import Image
        from pptx.util import Inches, Pt

        layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), prs.slide_width - Inches(1), Inches(0.6))
        title_box.text_frame.paragraphs[0].text = "Deroulement propose et preuves" if lang == "fr" else "Proposed delivery and evidence"
        title_box.text_frame.paragraphs[0].font.size = Pt(24)
        with Image.open(io.BytesIO(figure)) as im:
            scale = min((prs.slide_width - Inches(1)) / im.width, (prs.slide_height - Inches(1.4)) / im.height)
            width, height = round(im.width * scale), round(im.height * scale)
        slide.shapes.add_picture(io.BytesIO(figure), (prs.slide_width - width) // 2, Inches(0.95), width=width, height=height)
    for key in ("executive_summary", "company", "architecture", "methodology", "followup_kpi", "planning", "staffing", "financial_schedule", "references", "compliance_matrix", "assumptions"):
        body = _clean(response.get(key) or "").strip()
        pages = _pptx_pages(body)
        for index, page in enumerate(pages):
            title = section_title(key, lang) + (f" ({index + 1}/{len(pages)})" if len(pages) > 1 else "")
            _add_pptx_slide(prs, title, page)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def attach_exports(
    store: Any,
    tender: dict[str, Any],
    profile: dict[str, Any],
    response: dict[str, Any],
) -> dict[str, Any]:
    """Write Word/PPT next to the notice and return export refs on the response."""
    out = dict(response)
    from navin.agent.skill_routing import build_action_skill_context

    export_skills = build_action_skill_context("tenders", "export")
    out["export_generation"] = {"mode": "deterministic", "status": "needs_review", "skills": export_skills.metadata, "warnings": [], "pptx_scope": "summary_with_requirement_matrix", "docx_scope": "complete_editable_dossier"}
    tid = _safe_name(str(tender.get("id") or "notice"), "notice")
    exports: dict[str, Any] = {}
    _attach_diagram_exports(store, tender, out, exports)
    out["exports"] = exports
    if out["export_generation"]["warnings"]:
        generation = dict(out.get("generation") or {})
        generation["warnings"] = list(dict.fromkeys([*(generation.get("warnings") or []), *out["export_generation"]["warnings"]]))
        out["generation"] = generation
    docx = render_docx(store, tender, profile, out)
    if docx:
        name = f"{tid}_response.docx"
        record = store.save_bytes(name, docx, file_id=f"out_docx_{tid}")
        exports["docx"] = {
            **record,
            "kind": "docx",
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }
    pptx = render_pptx(store, tender, profile, out)
    if pptx:
        name = f"{tid}_response.pptx"
        record = store.save_bytes(name, pptx, file_id=f"out_pptx_{tid}")
        exports["pptx"] = {
            **record,
            "kind": "pptx",
            "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
    out["exports"] = exports
    out["pack_ready"] = bool(exports.get("docx") or exports.get("pptx"))
    if not out["pack_ready"]:
        warning = "Aucun export Word/PPT produit; les diagrammes seuls ne constituent pas le dossier." if _lang_of(out) == "fr" else "No Word/PPT export produced; diagrams alone do not constitute the dossier."
        out["export_generation"]["warnings"].append(warning)
        generation = dict(out.get("generation") or {})
        generation["warnings"] = list(dict.fromkeys([*(generation.get("warnings") or []), warning]))
        out["generation"] = generation
    return out
