#!/usr/bin/env python3
# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only
"""Apply the Navin table style to a worksheet range.

Usage:
    xlsx_style.py <file.xlsx> [--sheet NAME] [--range A1:F20] [--out OUT.xlsx]
                  [--no-zebra] [--no-freeze] [--no-filter] [--no-widths]
                  [--number-format COL=FORMAT ...] [--model-colors]

Header row: Ink fill (#17171a), Paper text (#f4f3ee), bold, centered.
Body: thin Paper borders, zebra Cream (#faf9f6) on even rows, numbers right
aligned. Freeze panes under the header, autofilter on the range, column
widths fitted to content (capped at 60). --model-colors recolors fonts with
the financial model convention: inputs Teal, formulas Ink, cross-sheet links
Leaf, cells starting with "check" Coral.

Exit codes: 0 ok, 2 usage or I/O error.
"""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

# Errors a malformed input or a bad argument can raise; reported as JSON with exit 2.
CLI_ERRORS = (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile, InvalidFileException)

PALETTE = {
    "ink": "17171A",
    "graphite": "121214",
    "paper": "F4F3EE",
    "cream": "FAF9F6",
    "lime": "C6E88A",
    "teal": "0E7490",
    "sky": "0EA5E9",
    "leaf": "6FBF6A",
    "violet": "7C3AED",
    "coral": "FF5F57",
    "amber": "FEBC2E",
}
TEXT_FONT = "Inter"
DEFAULT_NUMBER_FORMAT = "#,##0.00"
DEFAULT_INT_FORMAT = "#,##0"


def _used_range(ws) -> str:
    return f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _parse_formats(items: list[str]) -> dict[str, str]:
    out = {}
    for item in items:
        if "=" not in item:
            raise ValueError(f"--number-format expects COL=FORMAT, got {item!r}")
        col, fmt = item.split("=", 1)
        out[col.strip().upper()] = fmt
    return out


def style_range(ws, ref: str, zebra: bool, freeze: bool, autofilter: bool, widths: bool, number_formats: dict[str, str], model_colors: bool) -> dict:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    thin = Side(style="thin", color=PALETTE["paper"])
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=PALETTE["ink"])
    zebra_fill = PatternFill("solid", fgColor=PALETTE["cream"])
    header_font = Font(name=TEXT_FONT, bold=True, color=PALETTE["paper"], size=11)
    body_font = Font(name=TEXT_FONT, color=PALETTE["ink"], size=10)

    for col in range(min_col, max_col + 1):
        cell = ws.cell(row=min_row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[min_row].height = 22

    styled = 0
    for row in range(min_row + 1, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.font = body_font
            if zebra and (row - min_row) % 2 == 0:
                cell.fill = zebra_fill
            letter = get_column_letter(col)
            if letter in number_formats:
                cell.number_format = number_formats[letter]
            elif isinstance(cell.value, float):
                cell.number_format = DEFAULT_NUMBER_FORMAT
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                cell.number_format = DEFAULT_INT_FORMAT
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.alignment = Alignment(horizontal="right")
            if model_colors:
                cell.font = Font(name=TEXT_FONT, size=10, color=_model_color(cell.value))
            styled += 1

    if freeze:
        ws.freeze_panes = ws.cell(row=min_row + 1, column=min_col)
    if autofilter:
        ws.auto_filter.ref = ref
    if widths:
        for col in range(min_col, max_col + 1):
            letter = get_column_letter(col)
            longest = 0
            for row in range(min_row, max_row + 1):
                value = ws.cell(row=row, column=col).value
                if value is None:
                    continue
                text = str(value) if not (isinstance(value, str) and value.startswith("=")) else "0000000000"
                longest = max(longest, len(text))
            ws.column_dimensions[letter].width = min(60, max(8, longest + 2))
    return {"range": ref, "header_row": min_row, "cells_styled": styled}


def _model_color(value) -> str:
    if isinstance(value, str):
        if value.lower().startswith("check"):
            return PALETTE["coral"]
        if value.startswith("="):
            return PALETTE["leaf"] if "!" in value else PALETTE["ink"]
        return PALETTE["ink"]
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return PALETTE["teal"]
    return PALETTE["ink"]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="xlsx_style.py", description="Apply the Navin table style.")
    parser.add_argument("file", type=Path)
    parser.add_argument("--sheet", default=None, help="sheet name (default: active)")
    parser.add_argument("--range", dest="ref", default=None, help="A1 range (default: used range)")
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument("--no-zebra", action="store_true")
    parser.add_argument("--no-freeze", action="store_true")
    parser.add_argument("--no-filter", action="store_true")
    parser.add_argument("--no-widths", action="store_true")
    parser.add_argument("--number-format", action="append", default=[], metavar="COL=FORMAT")
    parser.add_argument("--model-colors", action="store_true", help="financial model font colors")
    args = parser.parse_args(argv)
    try:
        wb = load_workbook(str(args.file))
        ws = wb[args.sheet] if args.sheet else wb.active
        ref = args.ref or _used_range(ws)
        result = style_range(
            ws,
            ref,
            zebra=not args.no_zebra,
            freeze=not args.no_freeze,
            autofilter=not args.no_filter,
            widths=not args.no_widths,
            number_formats=_parse_formats(args.number_format),
            model_colors=args.model_colors,
        )
        target = args.out or args.file
        wb.save(str(target))
        result.update({"ok": True, "sheet": ws.title, "output": str(target)})
        print(json.dumps(result))
        return 0
    except CLI_ERRORS as exc:
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
