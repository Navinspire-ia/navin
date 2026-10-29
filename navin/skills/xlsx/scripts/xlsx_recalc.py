#!/usr/bin/env python3
# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only
"""Recalculate a workbook and scan it for error values.

Usage:
    xlsx_recalc.py <file.xlsx> [--out OUT.xlsx] [--static] [--timeout SEC]

With LibreOffice available, the workbook is re-saved headless so every formula
is recomputed, then the cached values are scanned for #REF!, #DIV/0!, #VALUE!,
#NAME?, #N/A, #NUM!, #NULL!. Without LibreOffice (or with --static), a static
audit of the formulas is run instead: references to missing sheets, literal
error values, out-of-range cell references, obvious division by a zero
literal. The JSON report always says whether a recalculation happened
("recalculated": true/false).

Exit codes: 0 no error found, 1 errors found, 2 usage or I/O error.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException

# Errors a malformed input, a bad argument or a failed LibreOffice run can raise.
CLI_ERRORS = (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile, InvalidFileException, subprocess.SubprocessError)

ERROR_VALUES = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
MAX_ROWS = 1_048_576
MAX_COLS = 16_384

# 'Sheet Name'!A1 or Sheet1!A1:B2 (the sheet part is captured without quotes)
SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!\$?[A-Z]{1,3}\$?\d+")
CELL_RE = re.compile(r"(?<![A-Za-z0-9_!.])\$?([A-Z]{1,3})\$?(\d+)(?![A-Za-z0-9_(])")
DIV_ZERO_RE = re.compile(r"/\s*0(?![.\d])")


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    return None


# --------------------------------------------------------------------------- #
# recalculation through LibreOffice
# --------------------------------------------------------------------------- #


def _no_window() -> dict:
    """No console window for soffice on Windows, through navin's helper when navin is importable."""
    try:
        from navin.utils.proc import no_window_kwargs

        return dict(no_window_kwargs())
    except ImportError:
        return {"creationflags": 0x08000000} if os.name == "nt" else {}


def recalc_with_soffice(source: Path, out: Path, timeout: int) -> None:
    binary = find_soffice()
    if not binary:
        raise RuntimeError("soffice not found")
    with tempfile.TemporaryDirectory(prefix="navin-recalc-") as tmp:
        tmp_dir = Path(tmp)
        cmd = [
            binary,
            "--headless",
            "--norestore",
            f"-env:UserInstallation=file://{tmp_dir / 'profile'}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(tmp_dir),
            str(source.resolve()),
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, check=False, **_no_window())
        produced = tmp_dir / f"{source.stem}.xlsx"
        if proc.returncode != 0 or not produced.is_file():
            raise RuntimeError(f"soffice failed: {(proc.stderr or proc.stdout).strip()[-500:]}")
        shutil.copyfile(produced, out)


def scan_cached_errors(path: Path) -> list[dict]:
    wb_v = load_workbook(str(path), data_only=True)
    wb_f = load_workbook(str(path), data_only=False)
    errors = []
    for name in wb_v.sheetnames:
        ws_v, ws_f = wb_v[name], wb_f[name]
        for row in ws_v.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.upper() in ERROR_VALUES:
                    errors.append(
                        {"sheet": name, "cell": cell.coordinate, "error": cell.value, "formula": ws_f[cell.coordinate].value}
                    )
    wb_v.close()
    wb_f.close()
    return errors


# --------------------------------------------------------------------------- #
# static audit
# --------------------------------------------------------------------------- #


def static_audit(path: Path) -> dict:
    wb = load_workbook(str(path), data_only=False)
    sheet_names = set(wb.sheetnames)
    defined = {dn.name for dn in wb.defined_names.values()}
    issues = []
    formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                if value.upper() in ERROR_VALUES:
                    issues.append({"sheet": ws.title, "cell": cell.coordinate, "kind": "literal_error", "detail": value})
                    continue
                if not value.startswith("="):
                    continue
                formulas += 1
                body = value[1:]
                if "#REF!" in body.upper():
                    issues.append({"sheet": ws.title, "cell": cell.coordinate, "kind": "ref_error", "detail": value})
                for quoted, plain in SHEET_REF_RE.findall(body):
                    target = (quoted or plain).replace("''", "'")
                    if target not in sheet_names and target not in defined:
                        issues.append(
                            {"sheet": ws.title, "cell": cell.coordinate, "kind": "missing_sheet", "detail": f"{target} in {value}"}
                        )
                # strip string literals before scanning bare references
                stripped = re.sub(r'"[^"]*"', '""', body)
                for col, row_s in CELL_RE.findall(stripped):
                    if int(row_s) > MAX_ROWS or int(row_s) == 0 or column_index_from_string(col) > MAX_COLS:
                        issues.append(
                            {"sheet": ws.title, "cell": cell.coordinate, "kind": "out_of_range", "detail": f"{col}{row_s}"}
                        )
                if DIV_ZERO_RE.search(stripped):
                    issues.append({"sheet": ws.title, "cell": cell.coordinate, "kind": "div_by_zero_literal", "detail": value})
    wb.close()
    return {"formulas": formulas, "issues": issues}


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #


def run(path: Path, out: Path | None, static: bool, timeout: int) -> dict:
    target = out or path
    report: dict = {"file": str(path), "output": str(target), "recalculated": False, "soffice": find_soffice()}
    if not static and report["soffice"]:
        try:
            recalc_with_soffice(path, target, timeout)
            report["recalculated"] = True
            report["error_cells"] = scan_cached_errors(target)
        except (RuntimeError, subprocess.TimeoutExpired) as exc:
            report["recalc_error"] = str(exc)
    if not report["recalculated"]:
        report["note"] = "not recalculated: LibreOffice unavailable or --static; static formula audit only"
        audit = static_audit(path)
        report["formulas"] = audit["formulas"]
        report["error_cells"] = audit["issues"]
        if out and out != path:
            shutil.copyfile(path, out)
    report["ok"] = not report["error_cells"]
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="xlsx_recalc.py", description="Recalculate and scan for formula errors.")
    parser.add_argument("file", type=Path)
    parser.add_argument("--out", type=Path, default=None, help="write the recalculated copy here (default: in place)")
    parser.add_argument("--static", action="store_true", help="skip LibreOffice, static audit only")
    parser.add_argument("--timeout", type=int, default=120)
    args = parser.parse_args(argv)
    try:
        report = run(args.file, args.out, args.static, args.timeout)
    except CLI_ERRORS as exc:
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}), file=sys.stderr)
        return 2
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
