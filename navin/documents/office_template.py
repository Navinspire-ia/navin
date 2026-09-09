# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""Work from the user's own PowerPoint or Word template.

A company deck or a letterhead arrives as a ``.pptx`` or a ``.docx`` the user
already owns. That file is the design: its slide layouts, its theme fonts and
colours, its heading styles, its header and footer. This module lets the agent
use it the way a colleague would, without redrawing anything:

``inspect``
    Lists what the template offers: slide layouts and their placeholders (with
    the box each one occupies), theme colours and fonts, Word paragraph styles,
    the ``{{placeholders}}`` found in the text, headers and footers, tables.

``fill``
    Replaces ``{{key}}`` placeholders everywhere text lives (body, tables,
    headers, footers, notes, grouped shapes), keeping the run that carried the
    placeholder so bold, colour and size survive. A list value becomes one
    paragraph per item. ``{{image:key}}`` becomes a picture. A table row
    holding ``{{rows.field}}`` is repeated for every entry of ``rows``.

``build`` (PPTX)
    Adds slides from a deck JSON, choosing the template's own layouts by name
    and filling their placeholders: title, subtitle, bulleted body with levels,
    picture, table and speaker notes. The sample slides shipped in the template
    are dropped unless ``--keep-slides`` is passed.

Usage::

    python3 -m navin.documents.office_template inspect brand.pptx
    python3 -m navin.documents.office_template fill letter.docx --data values.json -o letter_filled.docx
    python3 -m navin.documents.office_template build brand.pptx --slides deck.json -o deck.pptx
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable, Sequence

_PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.:-]+)\s*\}\}")
_ROW_FIELD_RE = re.compile(r"\{\{\s*([A-Za-z0-9_-]+)\.([A-Za-z0-9_-]+)\s*\}\}")
# What inspect lists: every marker, including "{{image:logo|3cm}}" with its size.
_ANY_PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.:|-]+)\s*\}\}")
_EMU_PER_INCH = 914400


class TemplateError(RuntimeError):
    """Raised when the template or the data cannot be used."""


def _inches(value: int | None) -> float:
    return round((value or 0) / _EMU_PER_INCH, 2)


def _load_json(path: str | Path) -> Any:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise TemplateError(f"cannot read {path}: {exc}") from exc


def _lookup(data: dict[str, Any], dotted: str) -> Any:
    node: Any = data
    for part in dotted.split("."):
        if isinstance(node, dict) and part in node:
            node = node[part]
        else:
            return None
    return node


_SEARCH_DIRS: list[Path] = []


def _resolve_file(value: Any) -> Path | None:
    """A picture named in the data, looked up where the user would expect.

    Relative names resolve against the data file, then the template, then the
    current folder: "logo.png" beside values.json is what people write.
    """
    if not value:
        return None
    raw = Path(str(value)).expanduser()
    candidates = [raw] if raw.is_absolute() else [base / raw for base in _SEARCH_DIRS] + [raw]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


# --------------------------------------------------------------------- run-level replace


def _replace_in_runs(runs: list[Any], data: dict[str, Any], report: dict[str, Any]) -> list[Any]:
    """Replace placeholders across a list of runs; return list-valued leftovers.

    A placeholder often spans several runs ("{{", "name", "}}" each carrying
    its own formatting after an edit). The joined text is searched and the
    replacement lands in the run where the placeholder starts; the runs it
    spilled into lose the characters that belonged to it.
    """
    if not runs:
        return []
    texts = [run.text or "" for run in runs]
    joined = "".join(texts)
    if "{{" not in joined:
        return []
    deferred: list[Any] = []
    # Positions of each run in the joined string, recomputed after every edit.
    for match in list(_PLACEHOLDER_RE.finditer(joined))[::-1]:
        key = match.group(1)
        if key.startswith("image:"):
            continue
        # "{{rows.total}}" inside a repeated row belongs to the row expansion;
        # "{{client.name}}" is an ordinary nested value.
        if "." in key and isinstance(_lookup(data, key.split(".")[0]), list):
            continue
        value = _lookup(data, key)
        if value is None:
            report.setdefault("missing", set()).add(key)
            continue
        if isinstance(value, (list, tuple)):
            deferred.append((key, list(value)))
            continue
        replacement = str(value)
        start, end = match.start(), match.end()
        cursor = 0
        first = True
        for run, text in zip(runs, texts, strict=False):
            run_start, run_end = cursor, cursor + len(text)
            cursor = run_end
            if run_end <= start or run_start >= end:
                continue
            local_start = max(start - run_start, 0)
            local_end = min(end - run_start, len(text))
            new_text = text[:local_start] + (replacement if first else "") + text[local_end:]
            first = False
            run.text = new_text
        texts = [run.text or "" for run in runs]
        joined = "".join(texts)
        report["replaced"] = report.get("replaced", 0) + 1
    return deferred


def _paragraph_runs(paragraph) -> list[Any]:
    return list(paragraph.runs)


# ---------------------------------------------------------------------------- DOCX


def _docx_iter_paragraphs(document) -> Iterable[tuple[Any, Any]]:
    """(paragraph, container) for the body, tables, headers and footers."""
    for paragraph in document.paragraphs:
        yield paragraph, document
    for table in document.tables:
        yield from _docx_table_paragraphs(table)
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            try:
                if part.is_linked_to_previous:
                    continue
            except Exception:  # pragma: no cover - part without the flag
                pass
            for paragraph in part.paragraphs:
                yield paragraph, part
            for table in part.tables:
                yield from _docx_table_paragraphs(table)


def _docx_table_paragraphs(table) -> Iterable[tuple[Any, Any]]:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph, cell
            for nested in cell.tables:
                yield from _docx_table_paragraphs(nested)


def _docx_expand_rows(document, data: dict[str, Any], report: dict[str, Any]) -> None:
    """Repeat a table row carrying ``{{rows.field}}`` once per entry."""
    tables = list(document.tables)
    for section in document.sections:
        tables.extend(section.header.tables)
        tables.extend(section.footer.tables)
    for table in tables:
        for row in list(table.rows):
            text = " ".join(cell.text for cell in row.cells)
            names = {m.group(1) for m in _ROW_FIELD_RE.finditer(text)}
            entries = None
            name = None
            for candidate in names:
                value = _lookup(data, candidate)
                if isinstance(value, list):
                    entries, name = value, candidate
                    break
            if entries is None or name is None:
                continue
            template_tr = row._tr
            parent = template_tr.getparent()
            anchor = template_tr
            for entry in entries:
                clone = copy.deepcopy(template_tr)
                anchor.addnext(clone)
                anchor = clone
                item = entry if isinstance(entry, dict) else {"value": entry}
                scoped = {**data, name: item}
                from docx.table import _Row

                new_row = _Row(clone, table)
                for cell in new_row.cells:
                    for paragraph in cell.paragraphs:
                        _replace_row_fields(paragraph, name, scoped, report)
            parent.remove(template_tr)
            report["rows_expanded"] = report.get("rows_expanded", 0) + len(entries)


def _replace_row_fields(paragraph, name: str, scoped: dict[str, Any], report: dict[str, Any]) -> None:
    runs = _paragraph_runs(paragraph)
    texts = [run.text or "" for run in runs]
    joined = "".join(texts)
    for match in list(_ROW_FIELD_RE.finditer(joined))[::-1]:
        if match.group(1) != name:
            continue
        value = _lookup(scoped, f"{name}.{match.group(2)}")
        replacement = "" if value is None else str(value)
        start, end = match.start(), match.end()
        cursor = 0
        first = True
        for run, text in zip(runs, texts, strict=False):
            run_start, run_end = cursor, cursor + len(text)
            cursor = run_end
            if run_end <= start or run_start >= end:
                continue
            local_start = max(start - run_start, 0)
            local_end = min(end - run_start, len(text))
            run.text = text[:local_start] + (replacement if first else "") + text[local_end:]
            first = False
        texts = [run.text or "" for run in runs]
        joined = "".join(texts)
        report["replaced"] = report.get("replaced", 0) + 1


def _docx_list_paragraph(paragraph, key: str, items: list[Any]) -> None:
    """Turn one placeholder paragraph into one paragraph per item, same style."""
    template = paragraph._p
    runs = paragraph.runs
    if not runs:
        return
    pattern = re.compile(r"\{\{\s*" + re.escape(key) + r"\s*\}\}")
    anchor = template
    for item in items:
        clone = copy.deepcopy(template)
        anchor.addnext(clone)
        anchor = clone
        from docx.text.paragraph import Paragraph

        new_paragraph = Paragraph(clone, paragraph._parent)
        joined = "".join(run.text or "" for run in new_paragraph.runs)
        replaced = pattern.sub(str(item), joined, count=1)
        new_runs = new_paragraph.runs
        new_runs[0].text = replaced
        for extra in new_runs[1:]:
            extra.text = ""
    template.getparent().remove(template)


def _docx_image_placeholder(paragraph, data: dict[str, Any], report: dict[str, Any]) -> None:
    joined = "".join(run.text or "" for run in paragraph.runs)
    match = re.search(r"\{\{\s*image:([A-Za-z0-9_.-]+)(?:\|(\d+(?:\.\d+)?)(in|cm|mm))?\s*\}\}", joined)
    if not match:
        return
    key = match.group(1)
    value = _lookup(data, key)
    path = _resolve_file(value)
    if path is None:
        report.setdefault("missing", set()).add(f"image:{key}" + (f" ({value})" if value else ""))
        return
    from docx.shared import Cm, Inches, Mm

    width = None
    if match.group(2):
        amount = float(match.group(2))
        width = {"in": Inches, "cm": Cm, "mm": Mm}[match.group(3)](amount)
    runs = paragraph.runs
    for run in runs:
        run.text = ""
    target = runs[0] if runs else paragraph.add_run()
    if width is None:
        target.add_picture(str(path), width=Inches(5.5))
    else:
        target.add_picture(str(path), width=width)
    report["images"] = report.get("images", 0) + 1


def fill_docx(template: Path, data: dict[str, Any], output: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(template))
    report: dict[str, Any] = {}
    _docx_expand_rows(document, data, report)
    for paragraph, _container in list(_docx_iter_paragraphs(document)):
        if "{{" not in paragraph.text:
            continue
        _docx_image_placeholder(paragraph, data, report)
        deferred = _replace_in_runs(_paragraph_runs(paragraph), data, report)
        for key, items in deferred:
            _docx_list_paragraph(paragraph, key, items)
            report["replaced"] = report.get("replaced", 0) + 1
    title = data.get("title") or data.get("document_title")
    if isinstance(title, str) and title.strip():
        document.core_properties.title = title.strip()[:255]
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    report["output"] = str(output)
    return _finish_report(report)


def inspect_docx(template: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(template))
    placeholders: dict[str, int] = {}
    for paragraph, _ in _docx_iter_paragraphs(document):
        for match in _ANY_PLACEHOLDER_RE.finditer(paragraph.text):
            placeholders[match.group(1)] = placeholders.get(match.group(1), 0) + 1
    used_styles: dict[str, int] = {}
    for paragraph in document.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else "Normal"
        used_styles[name] = used_styles.get(name, 0) + 1
    available = sorted(
        style.name for style in document.styles if style.type == 1 and not style.hidden
    )[:60]
    sections = [
        {
            "orientation": "landscape" if section.page_width > section.page_height else "portrait",
            "page_in": (_inches(section.page_width), _inches(section.page_height)),
            "margins_in": {
                "top": _inches(section.top_margin),
                "bottom": _inches(section.bottom_margin),
                "left": _inches(section.left_margin),
                "right": _inches(section.right_margin),
            },
            "header": " ".join(p.text for p in section.header.paragraphs if p.text.strip())[:120],
            "footer": " ".join(p.text for p in section.footer.paragraphs if p.text.strip())[:120],
        }
        for section in document.sections
    ]
    tables = [
        {"rows": len(table.rows), "columns": len(table.columns), "first_row": [c.text[:24] for c in table.rows[0].cells]}
        for table in document.tables
        if table.rows
    ]
    return {
        "kind": "docx",
        "path": str(template),
        "title": document.core_properties.title or None,
        "paragraphs": len(document.paragraphs),
        "words": sum(len(p.text.split()) for p in document.paragraphs),
        "placeholders": placeholders,
        "styles_used": used_styles,
        "styles_available": available,
        "sections": sections,
        "tables": tables,
        "images": len(document.inline_shapes),
    }


# ---------------------------------------------------------------------------- PPTX


def _pptx_iter_text_frames(shapes) -> Iterable[Any]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_iter_text_frames(shape.shapes)
            continue
        if shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def _pptx_expand_rows(slide, data: dict[str, Any], report: dict[str, Any]) -> None:
    for shape in slide.shapes:
        if not (getattr(shape, "has_table", False) and shape.has_table):
            continue
        table = shape.table
        for row in list(table.rows):
            text = " ".join(cell.text for cell in row.cells)
            names = {m.group(1) for m in _ROW_FIELD_RE.finditer(text)}
            entries = name = None
            for candidate in names:
                value = _lookup(data, candidate)
                if isinstance(value, list):
                    entries, name = value, candidate
                    break
            if entries is None or name is None:
                continue
            template_tr = row._tr
            parent = template_tr.getparent()
            anchor = template_tr
            from pptx.table import _Row

            for entry in entries:
                clone = copy.deepcopy(template_tr)
                anchor.addnext(clone)
                anchor = clone
                item = entry if isinstance(entry, dict) else {"value": entry}
                scoped = {**data, name: item}
                for cell in _Row(clone, table).cells:
                    for paragraph in cell.text_frame.paragraphs:
                        _replace_row_fields(paragraph, name, scoped, report)
            parent.remove(template_tr)
            report["rows_expanded"] = report.get("rows_expanded", 0) + len(entries)


def _pptx_list_paragraph(text_frame, paragraph, key: str, items: list[Any]) -> None:
    template = paragraph._p
    pattern = re.compile(r"\{\{\s*" + re.escape(key) + r"\s*\}\}")
    anchor = template
    from pptx.text.text import _Paragraph

    for item in items:
        clone = copy.deepcopy(template)
        anchor.addnext(clone)
        anchor = clone
        new_paragraph = _Paragraph(clone, text_frame)
        runs = new_paragraph.runs
        if not runs:
            continue
        joined = "".join(run.text for run in runs)
        runs[0].text = pattern.sub(str(item), joined, count=1)
        for extra in runs[1:]:
            extra.text = ""
    template.getparent().remove(template)


def _pptx_image_placeholders(slide, data: dict[str, Any], report: dict[str, Any]) -> None:
    """A shape whose text is ``{{image:key}}`` becomes a picture in its box."""
    for shape in list(slide.shapes):
        if not shape.has_text_frame:
            continue
        match = re.fullmatch(r"\s*\{\{\s*image:([A-Za-z0-9_.-]+)\s*\}\}\s*", shape.text_frame.text)
        if not match:
            continue
        value = _lookup(data, match.group(1))
        path = _resolve_file(value)
        if path is None:
            report.setdefault("missing", set()).add(f"image:{match.group(1)}")
            continue
        if shape.is_placeholder and getattr(shape.placeholder_format, "type", None) is not None:
            try:
                shape.insert_picture(str(path))
                report["images"] = report.get("images", 0) + 1
                continue
            except Exception:
                pass
        slide.shapes.add_picture(str(path), shape.left, shape.top, shape.width, shape.height)
        shape._element.getparent().remove(shape._element)
        report["images"] = report.get("images", 0) + 1


def fill_pptx(template: Path, data: dict[str, Any], output: Path) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(str(template))
    report: dict[str, Any] = {}
    for slide in presentation.slides:
        _pptx_expand_rows(slide, data, report)
        _pptx_image_placeholders(slide, data, report)
        frames = list(_pptx_iter_text_frames(slide.shapes))
        if slide.has_notes_slide:
            frames.append(slide.notes_slide.notes_text_frame)
        for frame in frames:
            for paragraph in list(frame.paragraphs):
                deferred = _replace_in_runs(list(paragraph.runs), data, report)
                for key, items in deferred:
                    _pptx_list_paragraph(frame, paragraph, key, items)
                    report["replaced"] = report.get("replaced", 0) + 1
    title = data.get("title") or data.get("deck_title")
    if isinstance(title, str) and title.strip():
        presentation.core_properties.title = title.strip()[:255]
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    report["output"] = str(output)
    return _finish_report(report)


def _theme_summary(presentation) -> dict[str, Any]:
    """Colours and fonts of the first master's theme, straight from theme1.xml."""
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT

        master = presentation.slide_masters[0]
        theme_part = master.part.part_related_by(RT.THEME)
        from lxml import etree

        root = etree.fromstring(theme_part.blob)
    except Exception:
        return {}
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    colors: dict[str, str] = {}
    scheme = root.find(".//a:clrScheme", ns)
    if scheme is not None:
        for node in scheme:
            tag = etree.QName(node).localname
            value = node.find("a:srgbClr", ns)
            system = node.find("a:sysClr", ns)
            if value is not None:
                colors[tag] = "#" + value.get("val", "")
            elif system is not None:
                colors[tag] = "#" + system.get("lastClr", "")
    fonts: dict[str, str] = {}
    for role, xpath in (("heading", ".//a:majorFont/a:latin"), ("body", ".//a:minorFont/a:latin")):
        node = root.find(xpath, ns)
        if node is not None:
            fonts[role] = node.get("typeface", "")
    return {"colors": colors, "fonts": fonts}


def inspect_pptx(template: Path) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(str(template))
    layouts = []
    for index, layout in enumerate(presentation.slide_layouts):
        placeholders = []
        for shape in layout.placeholders:
            fmt = shape.placeholder_format
            placeholders.append(
                {
                    "idx": fmt.idx,
                    "type": str(fmt.type).split(".")[-1].split(" ")[0],
                    "name": shape.name,
                    "box_in": [_inches(shape.left), _inches(shape.top), _inches(shape.width), _inches(shape.height)],
                }
            )
        layouts.append({"index": index, "name": layout.name, "placeholders": placeholders})
    slides = []
    placeholders_found: dict[str, int] = {}
    for number, slide in enumerate(presentation.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        slides.append({"number": number, "layout": slide.slide_layout.name, "title": title[:80]})
        for frame in _pptx_iter_text_frames(slide.shapes):
            for match in _ANY_PLACEHOLDER_RE.finditer(frame.text):
                placeholders_found[match.group(1)] = placeholders_found.get(match.group(1), 0) + 1
    return {
        "kind": "pptx",
        "path": str(template),
        "title": presentation.core_properties.title or None,
        "slide_size_in": (_inches(presentation.slide_width), _inches(presentation.slide_height)),
        "theme": _theme_summary(presentation),
        "layouts": layouts,
        "slides": slides,
        "placeholders": placeholders_found,
    }


# ------------------------------------------------------------------- PPTX build


def _remove_slides(presentation) -> int:
    """Drop every slide, keeping masters and layouts (the design)."""
    id_list = presentation.slides._sldIdLst
    count = 0
    for sld_id in list(id_list):
        rel_id = sld_id.rId
        presentation.part.drop_rel(rel_id)
        id_list.remove(sld_id)
        count += 1
    return count


def _pick_layout(presentation, wanted: str | None, slide: dict[str, Any]):
    layouts = list(presentation.slide_layouts)
    if wanted:
        if wanted.isdigit() and int(wanted) < len(layouts):
            return layouts[int(wanted)]
        lowered = wanted.lower()
        for layout in layouts:
            if layout.name.lower() == lowered:
                return layout
        for layout in layouts:
            if lowered in layout.name.lower():
                return layout
    # No name, or a name the template does not have: choose by content.
    has_body = bool(slide.get("body") or slide.get("items") or slide.get("bullets"))
    has_picture = bool(slide.get("picture") or slide.get("image"))
    has_table = bool(slide.get("table"))
    preferences: list[str]
    if slide.get("subtitle") and not has_body and not has_picture:
        preferences = ["title slide", "titre", "cover"]
    elif has_picture and has_body:
        preferences = ["two content", "picture with caption", "deux contenus", "content with caption"]
    elif has_picture:
        preferences = ["picture with caption", "image", "title only", "titre seul", "blank"]
    elif has_body or has_table:
        preferences = ["title and content", "titre et contenu", "content", "title only"]
    else:
        preferences = ["section header", "title only", "titre seul", "blank"]
    for wanted_name in preferences:
        for layout in layouts:
            if wanted_name in layout.name.lower():
                return layout
    return layouts[min(1, len(layouts) - 1)]


def _placeholder_by_types(slide, types: Sequence[str], skip: set[int]) -> Any:
    for shape in slide.placeholders:
        kind = str(shape.placeholder_format.type).split(".")[-1].split(" ")[0]
        if kind in types and shape.placeholder_format.idx not in skip:
            return shape
    return None


def _write_bullets(text_frame, items: Sequence[Any], level: int = 0, first: bool = True) -> bool:
    for item in items:
        if isinstance(item, (list, tuple)):
            first = _write_bullets(text_frame, item, level + 1, first)
            continue
        if isinstance(item, dict):
            text = str(item.get("text") or item.get("label") or "")
            children = item.get("items") or item.get("children") or []
        else:
            text, children = str(item), []
        paragraph = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
        first = False
        paragraph.text = text
        paragraph.level = min(level, 4)
        if children:
            first = _write_bullets(text_frame, children, level + 1, first)
    return first


def build_pptx(
    template: Path, deck: dict[str, Any], output: Path, *, keep_slides: bool = False
) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.util import Pt

    presentation = Presentation(str(template))
    report: dict[str, Any] = {"layouts_used": {}, "warnings": []}
    if not keep_slides:
        report["template_slides_removed"] = _remove_slides(presentation)
    slides = deck.get("slides") if isinstance(deck, dict) else deck
    if not isinstance(slides, list) or not slides:
        raise TemplateError("the deck JSON needs a non-empty \"slides\" list")
    for number, spec in enumerate(slides, start=1):
        if not isinstance(spec, dict):
            continue
        layout = _pick_layout(presentation, str(spec.get("layout") or ""), spec)
        slide = presentation.slides.add_slide(layout)
        report["layouts_used"][layout.name] = report["layouts_used"].get(layout.name, 0) + 1
        used: set[int] = set()
        title = spec.get("title")
        if title and slide.shapes.title is not None:
            slide.shapes.title.text = str(title)
            used.add(slide.shapes.title.placeholder_format.idx)
        subtitle = spec.get("subtitle")
        if subtitle:
            box = _placeholder_by_types(slide, ("SUBTITLE", "BODY"), used)
            if box is not None:
                box.text_frame.text = str(subtitle)
                used.add(box.placeholder_format.idx)
        body = spec.get("body") or spec.get("items") or spec.get("bullets")
        if body:
            box = _placeholder_by_types(slide, ("BODY", "OBJECT"), used)
            if box is None:
                report["warnings"].append(f"slide {number}: layout '{layout.name}' has no body placeholder; text box added")
                box = slide.shapes.add_textbox(
                    presentation.slide_width // 12,
                    presentation.slide_height // 4,
                    presentation.slide_width * 10 // 12,
                    presentation.slide_height // 2,
                )
                box.text_frame.word_wrap = True
            else:
                used.add(box.placeholder_format.idx)
            items = body if isinstance(body, list) else [str(body)]
            _write_bullets(box.text_frame, items)
        picture = spec.get("picture") or spec.get("image")
        if picture:
            path = _resolve_file(picture)
            if path is not None:
                box = _placeholder_by_types(slide, ("PICTURE", "OBJECT", "BODY"), used)
                if box is not None and hasattr(box, "insert_picture"):
                    box.insert_picture(str(path))
                    used.add(box.placeholder_format.idx)
                elif box is not None:
                    slide.shapes.add_picture(str(path), box.left, box.top, box.width, box.height)
                    used.add(box.placeholder_format.idx)
                    box._element.getparent().remove(box._element)
                else:
                    width = presentation.slide_width // 2
                    slide.shapes.add_picture(str(path), presentation.slide_width - width - presentation.slide_width // 20, presentation.slide_height // 4, width=width)
            else:
                report["warnings"].append(f"slide {number}: picture not found: {picture}")
        table = spec.get("table")
        if isinstance(table, dict) and (table.get("rows") or table.get("headers")):
            headers = [str(h) for h in (table.get("headers") or [])]
            rows = [[str(c) for c in row] for row in (table.get("rows") or []) if isinstance(row, (list, tuple))]
            columns = max([len(headers)] + [len(r) for r in rows]) if (headers or rows) else 0
            if columns:
                box = _placeholder_by_types(slide, ("BODY", "OBJECT", "TABLE"), used)
                if box is not None:
                    left, top, width, height = box.left, box.top, box.width, box.height
                    used.add(box.placeholder_format.idx)
                    box._element.getparent().remove(box._element)
                else:
                    left = presentation.slide_width // 12
                    top = presentation.slide_height // 4
                    width = presentation.slide_width * 10 // 12
                    height = presentation.slide_height // 2
                total_rows = len(rows) + (1 if headers else 0)
                shape = slide.shapes.add_table(total_rows, columns, left, top, width, height)
                grid = shape.table
                row_index = 0
                if headers:
                    for col, header in enumerate(headers):
                        grid.cell(0, col).text = header
                    row_index = 1
                for row in rows:
                    for col in range(columns):
                        grid.cell(row_index, col).text = row[col] if col < len(row) else ""
                    row_index += 1
                for row in grid.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(14 if total_rows <= 8 else 11)
        notes = spec.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)
        # An unfilled placeholder shows "Click to add text" in edit mode and
        # prints as an empty box: it leaves with the sample slides.
        for shape in list(slide.placeholders):
            if shape.placeholder_format.idx in used:
                continue
            if shape.has_text_frame and not shape.text_frame.text.strip():
                shape._element.getparent().remove(shape._element)
    deck_title = deck.get("title") if isinstance(deck, dict) else None
    first_title = next((str(s.get("title")) for s in slides if isinstance(s, dict) and s.get("title")), "")
    presentation.core_properties.title = str(deck_title or first_title or output.stem)[:255]
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    report["output"] = str(output)
    report["slides"] = len(presentation.slides)
    return report


# ------------------------------------------------------------------------ CLI


def _finish_report(report: dict[str, Any]) -> dict[str, Any]:
    missing = report.get("missing")
    report["missing"] = sorted(missing) if missing else []
    report.setdefault("replaced", 0)
    return report


def inspect(template: str | Path) -> dict[str, Any]:
    path = Path(template).expanduser()
    if not path.is_file():
        raise TemplateError(f"not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return inspect_pptx(path)
    if suffix == ".docx":
        return inspect_docx(path)
    raise TemplateError(f"{path.name}: use a .pptx or a .docx template")


def _set_search_dirs(*paths: str | Path | None) -> None:
    _SEARCH_DIRS.clear()
    for raw in paths:
        if not raw:
            continue
        folder = Path(raw).expanduser()
        folder = folder if folder.is_dir() else folder.parent
        if folder.is_dir() and folder not in _SEARCH_DIRS:
            _SEARCH_DIRS.append(folder)


def fill(
    template: str | Path,
    data: dict[str, Any],
    output: str | Path,
    *,
    data_path: str | Path | None = None,
) -> dict[str, Any]:
    path = Path(template).expanduser()
    if not path.is_file():
        raise TemplateError(f"not found: {path}")
    if not isinstance(data, dict):
        raise TemplateError("the data file must hold a JSON object of key: value")
    _set_search_dirs(data_path, path)
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return fill_pptx(path, data, Path(output).expanduser())
    if suffix == ".docx":
        return fill_docx(path, data, Path(output).expanduser())
    raise TemplateError(f"{path.name}: use a .pptx or a .docx template")


def build(
    template: str | Path,
    deck: dict[str, Any],
    output: str | Path,
    *,
    keep_slides: bool = False,
    deck_path: str | Path | None = None,
) -> dict[str, Any]:
    path = Path(template).expanduser()
    if not path.is_file():
        raise TemplateError(f"not found: {path}")
    if path.suffix.lower() != ".pptx":
        raise TemplateError("build works on a .pptx template; use fill for Word documents")
    _set_search_dirs(deck_path, path)
    return build_pptx(path, deck, Path(output).expanduser(), keep_slides=keep_slides)


def _print_inspection(info: dict[str, Any]) -> None:
    print(f"{info['path']} ({info['kind']})" + (f' - "{info["title"]}"' if info.get("title") else ""))
    if info["kind"] == "pptx":
        size = info["slide_size_in"]
        print(f"  slide size: {size[0]} x {size[1]} in")
        theme = info.get("theme") or {}
        if theme.get("fonts"):
            print(f"  fonts: heading {theme['fonts'].get('heading')}, body {theme['fonts'].get('body')}")
        if theme.get("colors"):
            accents = ", ".join(f"{k} {v}" for k, v in theme["colors"].items() if k.startswith("accent"))
            print(f"  accents: {accents}")
        print(f"  layouts ({len(info['layouts'])}):")
        for layout in info["layouts"]:
            kinds = ", ".join(f"{p['type'].lower()}#{p['idx']}" for p in layout["placeholders"])
            print(f"    [{layout['index']}] {layout['name']}: {kinds or 'no placeholder'}")
        print(f"  slides: {len(info['slides'])}" + (" (sample content; build replaces them)" if info["slides"] else ""))
    else:
        for section in info["sections"]:
            print(f"  page {section['page_in'][0]} x {section['page_in'][1]} in, {section['orientation']}, margins {section['margins_in']}")
            if section["header"]:
                print(f"  header: {section['header']}")
            if section["footer"]:
                print(f"  footer: {section['footer']}")
        print(f"  paragraphs: {info['paragraphs']}, words: {info['words']}, tables: {len(info['tables'])}, images: {info['images']}")
        print("  styles used: " + ", ".join(f"{k} x{v}" for k, v in sorted(info["styles_used"].items())))
    if info.get("placeholders"):
        print("  placeholders: " + ", ".join(f"{{{{{k}}}}} x{v}" for k, v in sorted(info["placeholders"].items())))
    else:
        print("  placeholders: none ({{key}} markers would be filled by `fill`)")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="office_template",
        description="Inspect, fill or build from a user-provided .pptx / .docx template.",
    )
    sub = parser.add_subparsers(dest="command", required=True)
    p_inspect = sub.add_parser("inspect", help="List layouts, placeholders, styles, theme")
    p_inspect.add_argument("template")
    p_inspect.add_argument("--json", action="store_true")
    p_fill = sub.add_parser("fill", help="Replace {{placeholders}} with values from a JSON file")
    p_fill.add_argument("template")
    p_fill.add_argument("--data", required=True, help="JSON object of key: value")
    p_fill.add_argument("-o", "--output", required=True)
    p_fill.add_argument("--strict", action="store_true", help="Fail when a placeholder has no value")
    p_fill.add_argument("--json", action="store_true")
    p_build = sub.add_parser("build", help="Add slides from a deck JSON using the template's layouts")
    p_build.add_argument("template")
    p_build.add_argument("--slides", required=True, help="deck.json with a slides list")
    p_build.add_argument("-o", "--output", required=True)
    p_build.add_argument("--keep-slides", action="store_true", help="Keep the template's own slides")
    p_build.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    try:
        if args.command == "inspect":
            info = inspect(args.template)
            if args.json:
                print(json.dumps(info, ensure_ascii=False, indent=2))
            else:
                _print_inspection(info)
            return 0
        if args.command == "fill":
            report = fill(args.template, _load_json(args.data), args.output, data_path=args.data)
            if args.json:
                print(json.dumps(report, ensure_ascii=False, indent=2))
            else:
                print(f"Wrote {report['output']} ({report['replaced']} replacement(s)"
                      + (f", {report['rows_expanded']} row(s) expanded" if report.get("rows_expanded") else "")
                      + (f", {report['images']} picture(s)" if report.get("images") else "")
                      + ")")
                if report["missing"]:
                    print("office_template: no value for " + ", ".join(report["missing"]), file=sys.stderr)
            return 1 if (args.strict and report["missing"]) else 0
        report = build(
            args.template,
            _load_json(args.slides),
            args.output,
            keep_slides=args.keep_slides,
            deck_path=args.slides,
        )
        if args.json:
            print(json.dumps(report, ensure_ascii=False, indent=2))
        else:
            layouts = ", ".join(f"{k} x{v}" for k, v in report["layouts_used"].items())
            print(f"Wrote {report['output']} ({report['slides']} slides; layouts: {layouts})")
            for warning in report["warnings"]:
                print(f"office_template: {warning}", file=sys.stderr)
        return 0
    except TemplateError as exc:
        print(f"office_template: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
