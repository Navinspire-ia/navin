# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""Check a finished DOCX, PPTX, XLSX or PDF before it is handed over.

The HTML quality passes (``word_qa``, ``ppt_qa``) look at the page before it
becomes a file. This one opens the file that will actually be sent, the way the
reader will, and reports what would embarrass the sender: sample copy left in
place, a placeholder never filled, a slide that is a picture of a slide, an
empty section, a broken formula, a PDF with no selectable text.

Findings come with a severity. ``block`` means the file must not ship as is
(the command exits 1); ``warn`` is worth a look. ``--json`` prints the whole
report for a script to read.

Usage::

    python3 -m navin.documents.doc_check report.docx
    python3 -m navin.documents.doc_check deck.pptx --json
    python3 -m navin.documents.doc_check workbook.xlsx report.pdf
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

if __package__:
    from . import _pdf
else:  # copied as a standalone folder into a user workspace
    import _pdf

# Phrases that only ever come from a lookbook or a generator's placeholder.
_SAMPLE_RE = re.compile(
    r"lorem ipsum|dolor sit amet|\bipsum\b|your (?:text|title|name|company|logo) here|"
    r"\binsert (?:text|title|your|image|logo|date|name)\b|\[insert\b|"
    r"\bjohn doe\b|\bjane doe\b|click to (?:add|edit)|the title is the message",
    re.IGNORECASE,
)
# Phrases that usually are leftovers but have legitimate uses (a form label,
# a roadmap item): reported, not blocking.
_SUSPECT_RE = re.compile(
    r"\bplaceholder\b|\bsample (?:text|title|copy|content|data)\b|\bacme corp\b|"
    r"\bcompany name\b|\btbd\b|\btodo\b|\bfixme\b|\bxxx+\b|\bslide title\b",
    re.IGNORECASE,
)
_PLACEHOLDER_RE = re.compile(r"\{\{[^{}]{1,80}\}\}|\{%[^%]{1,80}%\}|\[\[[^\[\]]{1,80}\]\]|<<[^<>]{1,80}>>|\$\{[^}]{1,80}\}")
_DASH_RE = re.compile("[\u2013\u2014]")
_EMOJI_RE = re.compile(
    "[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F2FF\U0001F900-\U0001F9FF"
    "\U00002B50\U00002B55\U00002705\U0000274C\U0000203C\U00002049]"
)
_FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
_NUMBER_TEXT_RE = re.compile(r"^\s*[-+]?\d[\d\s.,]*%?\s*$")

SEVERITIES = ("block", "warn", "info")


@dataclass
class Finding:
    code: str
    severity: str
    where: str
    message: str


@dataclass
class Report:
    path: str
    kind: str
    stats: dict[str, Any] = field(default_factory=dict)
    findings: list[Finding] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return not any(f.severity == "block" for f in self.findings)

    def add(self, code: str, severity: str, where: str, message: str) -> None:
        self.findings.append(Finding(code, severity, where, message))

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "kind": self.kind,
            "pass": self.passed,
            "stats": self.stats,
            "findings": [asdict(f) for f in self.findings],
        }


def _excerpt(text: str, match: re.Match[str], width: int = 40) -> str:
    start = max(0, match.start() - width // 2)
    end = min(len(text), match.end() + width // 2)
    snippet = " ".join(text[start:end].split())
    return f'"{snippet}"'


def check_text(report: Report, where: str, text: str, *, seen: set[str]) -> None:
    """Text rules shared by every format; one finding per code and location."""
    if not text or not text.strip():
        return
    for code, pattern, severity, advice in (
        ("leftover-sample", _SAMPLE_RE, "block", "sample copy from the template is still in the file: replace it with real content"),
        ("suspect-copy", _SUSPECT_RE, "warn", "this reads like a leftover or a stub: make sure it is meant to be there"),
        ("unresolved-placeholder", _PLACEHOLDER_RE, "block", "a placeholder was never filled"),
        ("dash", _DASH_RE, "block", "em or en dash (U+2014 / U+2013): use a comma, a colon, a period or a plain hyphen"),
        ("emoji", _EMOJI_RE, "block", "emoji in a deliverable: use a number, a plain glyph or a shape from the design"),
    ):
        match = pattern.search(text)
        if match is None:
            continue
        key = f"{code}|{where}"
        if key in seen:
            continue
        seen.add(key)
        report.add(code, severity, where, f"{advice} ({_excerpt(text, match)})")


def _words(text: str) -> int:
    return len(re.findall(r"\w+", text or ""))


# --------------------------------------------------------------------------- DOCX


def _heading_level(paragraph) -> int:
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    match = re.match(r"heading\s*(\d)", name, re.I)
    if match:
        return int(match.group(1))
    if name.lower() == "title":
        return 1
    return 0


def check_docx(path: Path) -> Report:
    from docx import Document

    report = Report(str(path), "docx")
    document = Document(str(path))
    body = document.paragraphs
    seen: set[str] = set()
    words = 0
    headings: list[tuple[int, str]] = []
    fonts: set[str] = set()
    tiny: list[str] = []
    previous_heading: tuple[int, str] | None = None
    body_since_heading = 0
    for index, paragraph in enumerate(body, start=1):
        text = paragraph.text
        words += _words(text)
        level = _heading_level(paragraph)
        where = f"paragraph {index}"
        check_text(report, where, text, seen=seen)
        if level:
            if previous_heading is not None and body_since_heading == 0 and previous_heading[0] >= level:
                report.add(
                    "empty-section",
                    "warn",
                    f"heading \"{previous_heading[1][:48]}\"",
                    "a heading with no content under it before the next heading of the same level: write the section or remove the title",
                )
            if headings and level > headings[-1][0] + 1:
                report.add(
                    "level-skipped",
                    "warn",
                    where,
                    f"heading jumps from level {headings[-1][0]} to {level}: Word builds the outline and the TOC from these",
                )
            headings.append((level, text.strip()))
            previous_heading = (level, text.strip())
            body_since_heading = 0
        elif text.strip():
            body_since_heading += 1
        for run in paragraph.runs:
            if run.font.name:
                fonts.add(run.font.name)
            if run.font.size is not None and run.font.size.pt < 7 and run.text.strip():
                tiny.append(where)
    tables = document.tables
    empty_cells = total_cells = 0
    for t_index, table in enumerate(tables, start=1):
        for row in table.rows:
            for cell in row.cells:
                total_cells += 1
                if not cell.text.strip():
                    empty_cells += 1
                check_text(report, f"table {t_index}", cell.text, seen=seen)
    for section in document.sections:
        for part_name, part in (("header", section.header), ("footer", section.footer)):
            for paragraph in part.paragraphs:
                check_text(report, part_name, paragraph.text, seen=seen)
    images = len(document.inline_shapes)
    xml = document.element.xml
    has_toc = "TOC \\" in xml or 'w:instr="TOC' in xml or "TOC \\o" in xml
    report.stats = {
        "paragraphs": len(body),
        "words": words,
        "headings": len(headings),
        "tables": len(tables),
        "images": images,
        "sections": len(document.sections),
        "toc": has_toc,
        "title": document.core_properties.title or None,
        "language": document.core_properties.language or None,
        "fonts": sorted(fonts),
    }
    if words < 30 and not tables and not images:
        report.add("empty-document", "block", "document", f"only {words} words: the document is empty")
    if words > 400 and not headings:
        report.add("no-heading", "warn", "document", "a long document without a single heading: give the sections titles so Word can build an outline")
    if len(headings) >= 6 and words > 1200 and not has_toc:
        report.add("no-toc", "warn", "document", "six headings and more than a thousand words but no table of contents: add data-doc-toc after the cover")
    if total_cells and empty_cells / total_cells > 0.4:
        report.add("hollow-table", "warn", "tables", f"{empty_cells} of {total_cells} table cells are empty: fill them or drop the empty columns")
    if len(fonts) > 4:
        report.add("font-zoo", "warn", "document", f"{len(fonts)} font families ({', '.join(sorted(fonts)[:5])}): two are enough, heading and body")
    if tiny:
        report.add("tiny-text", "warn", tiny[0], f"text under 7pt in {len(tiny)} place(s): unreadable in print")
    if not document.core_properties.title:
        report.add("no-title", "warn", "properties", "the document has no title property: Word shows it in the file list and the recent files")
    return report


# --------------------------------------------------------------------------- PPTX


def _shape_texts(shape) -> Iterable[str]:
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            yield "".join(run.text for run in paragraph.runs)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text
    if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
        for child in shape.shapes:
            yield from _shape_texts(child)


def check_pptx(path: Path) -> Report:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    report = Report(str(path), "pptx")
    presentation = Presentation(str(path))
    width, height = presentation.slide_width, presentation.slide_height
    seen: set[str] = set()
    titles: dict[str, list[int]] = {}
    picture_uses: dict[str, list[int]] = {}
    notes = 0
    charts = tables = pictures = 0
    tiny: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        where = f"slide {number}"
        texts: list[str] = []
        slide_pictures: list[tuple[str, float]] = []
        slide_charts = slide_tables = 0
        off_canvas = False
        for shape in slide.shapes:
            for text in _shape_texts(shape):
                if text.strip():
                    texts.append(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    digest = hashlib.sha1(shape.image.blob).hexdigest()
                except Exception:  # pragma: no cover - unreadable picture part
                    digest = f"?{number}"
                area = (shape.width or 0) * (shape.height or 0) / max(width * height, 1)
                slide_pictures.append((digest, area))
            if getattr(shape, "has_chart", False) and shape.has_chart:
                slide_charts += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                slide_tables += 1
            left, top = shape.left or 0, shape.top or 0
            right, bottom = left + (shape.width or 0), top + (shape.height or 0)
            # Slide furniture (a footer credit, a page number, a top kicker)
            # is allowed small type; body copy is not.
            furniture = top >= 0.9 * height or bottom <= 0.07 * height
            if shape.has_text_frame and not furniture:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size is not None and run.font.size.pt < 9 and run.text.strip():
                            tiny.append(where)
            tolerance = int(0.02 * width)
            if left < -tolerance or top < -tolerance or right > width + tolerance or bottom > height + tolerance:
                off_canvas = True
        charts += slide_charts
        tables += slide_tables
        pictures += len(slide_pictures)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes += 1
        for text in texts:
            check_text(report, where, text, seen=seen)
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        if title:
            titles.setdefault(title.lower(), []).append(number)
        for digest, area in slide_pictures:
            # A logo repeats by design; a hero picture reused is a lazy deck.
            if area >= 0.08:
                picture_uses.setdefault(digest, []).append(number)
        word_count = sum(_words(text) for text in texts)
        full_bleed = any(area >= 0.8 for _, area in slide_pictures)
        if not texts and not slide_charts and not slide_tables:
            if slide_pictures and full_bleed:
                report.add("rasterized-slide", "block", where, "the slide is a single picture with no text box: nothing can be edited or translated. Rebuild it from HTML with html2pptx")
            elif not slide_pictures:
                report.add("empty-slide", "block", where, "the slide holds nothing")
        elif word_count < 12 and not slide_pictures and not slide_charts and not slide_tables and number > 1:
            report.add("thin-slide", "warn", where, f"{word_count} words and no visual: a title and one line is a rejected slide. Add items, a KPI, a chart or a photo")
        if off_canvas:
            report.add("off-canvas", "warn", where, "a shape sits partly outside the slide: it gets clipped in PowerPoint")
    for title, numbers in titles.items():
        if len(numbers) > 1:
            report.add("duplicate-title", "warn", f"slides {', '.join(map(str, numbers))}", f'the same title "{title[:40]}" appears on several slides: each slide states one idea')
    for digest, numbers in picture_uses.items():
        if len(numbers) >= 3:
            report.add("repeated-picture", "warn", f"slides {', '.join(map(str, numbers[:5]))}", "the same large picture is reused across slides: use distinct visuals")
    if tiny:
        report.add("tiny-text", "warn", tiny[0], f"text under 9pt in {len(tiny)} place(s): unreadable from the back of the room")
    count = len(presentation.slides)
    report.stats = {
        "slides": count,
        "pictures": pictures,
        "charts": charts,
        "tables": tables,
        "slides_with_notes": notes,
        "title": presentation.core_properties.title or None,
        "size_in": (round(width / 914400, 2), round(height / 914400, 2)),
    }
    if count == 0:
        report.add("empty-deck", "block", "deck", "the deck has no slide")
    if count >= 4 and notes == 0:
        report.add("no-notes", "info", "deck", "no speaker notes: add a notes field per slide in the deck JSON (html2pptx copies them into the file)")
    return report


# --------------------------------------------------------------------------- XLSX


def check_xlsx(path: Path) -> Report:
    from openpyxl import load_workbook

    report = Report(str(path), "xlsx")
    workbook = load_workbook(str(path), data_only=False)
    seen: set[str] = set()
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        cells = formulas = numbers_as_text = 0
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None or value == "":
                    continue
                cells += 1
                if isinstance(value, str):
                    if value.startswith("="):
                        formulas += 1
                        continue
                    if value.strip() in _FORMULA_ERRORS:
                        report.add("formula-error", "block", f"{sheet.title}!{cell.coordinate}", f"the cell shows {value.strip()}: fix the formula or the data it points at")
                    elif _NUMBER_TEXT_RE.match(value) and len(value.strip()) < 20:
                        numbers_as_text += 1
                    check_text(report, f"{sheet.title}!{cell.coordinate}", value, seen=seen)
        first_row_empty = all(
            (cell.value is None or str(cell.value).strip() == "") for cell in next(sheet.iter_rows(min_row=1, max_row=1), [])
        )
        sheets.append({"name": sheet.title, "cells": cells, "formulas": formulas, "rows": sheet.max_row, "columns": sheet.max_column})
        if cells == 0:
            report.add("empty-sheet", "warn", sheet.title, "the sheet is empty: fill it or remove it")
            continue
        if first_row_empty:
            report.add("no-header-row", "warn", sheet.title, "row 1 is empty: a table starts with its header row (filters and freeze panes depend on it)")
        if cells and numbers_as_text / cells > 0.3:
            report.add("numbers-as-text", "warn", sheet.title, f"{numbers_as_text} numeric values are stored as text: they cannot be summed or charted. Write real numbers with a number format")
    report.stats = {"sheets": sheets, "title": workbook.properties.title or None}
    if not sheets:
        report.add("empty-workbook", "block", "workbook", "the workbook has no sheet")
    return report


# --------------------------------------------------------------------------- PDF


def check_pdf(path: Path) -> Report:
    report = Report(str(path), "pdf")
    info = _pdf.metadata(path)
    texts = _pdf.page_texts(path)
    seen: set[str] = set()
    for number, text in enumerate(texts, start=1):
        check_text(report, f"page {number}", text, seen=seen)
    thin = _pdf.looks_rasterized(texts)
    pages = len(texts)
    report.stats = {**info, "words": sum(_words(text) for text in texts)}
    if pages == 0:
        report.add("empty-pdf", "block", "document", "the PDF has no page")
    elif thin and len(thin) == pages:
        report.add("no-text-layer", "block", "document", "no page carries selectable text: the PDF is a stack of pictures. Print the HTML with html2pdf so text stays text")
    elif thin:
        report.add("thin-page", "warn", f"page(s) {', '.join(map(str, thin[:6]))}", "almost no selectable text: a blank page, or text rendered as an image")
    if not info.get("title"):
        report.add("no-title", "warn", "properties", "the PDF has no title: set <title> in the HTML before printing")
    return report


# --------------------------------------------------------------------------- CLI

CHECKERS = {
    ".docx": check_docx,
    ".pptx": check_pptx,
    ".xlsx": check_xlsx,
    ".xlsm": check_xlsx,
    ".pdf": check_pdf,
}


def check(path: str | Path) -> Report:
    """Check one file; the format follows the extension."""
    target = Path(path).expanduser()
    if not target.is_file():
        raise FileNotFoundError(f"not found: {target}")
    checker = CHECKERS.get(target.suffix.lower())
    if checker is None:
        raise ValueError(f"{target.name}: unsupported format (docx, pptx, xlsx, pdf)")
    return checker(target)


def _print_report(report: Report) -> None:
    verdict = "PASS" if report.passed else "FAIL"
    summary = ", ".join(
        f"{key} {value}" for key, value in report.stats.items()
        if key in {"pages", "slides", "words", "tables", "charts", "pictures", "images", "headings"} and value
    )
    print(f"{report.path}: {verdict}" + (f" ({summary})" if summary else ""))
    order = {severity: index for index, severity in enumerate(SEVERITIES)}
    for finding in sorted(report.findings, key=lambda f: order.get(f.severity, 9)):
        print(f"  [{finding.severity}] {finding.code} @ {finding.where}: {finding.message}")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="doc_check",
        description="Check finished DOCX / PPTX / XLSX / PDF files for what a reader would notice.",
    )
    parser.add_argument("files", nargs="+", help="Files to check")
    parser.add_argument("--json", action="store_true", help="Print the reports as JSON")
    args = parser.parse_args(argv)
    reports: list[Report] = []
    status = 0
    for raw in args.files:
        try:
            report = check(raw)
        except (FileNotFoundError, ValueError) as exc:
            print(f"doc_check: {exc}", file=sys.stderr)
            status = 2
            continue
        except Exception as exc:  # a file the library cannot open is a failed deliverable
            print(f"doc_check: {raw}: cannot be opened ({exc})", file=sys.stderr)
            status = 1
            continue
        reports.append(report)
        if not report.passed:
            status = max(status, 1)
    if args.json:
        payload = [report.as_dict() for report in reports]
        print(json.dumps(payload[0] if len(payload) == 1 else payload, ensure_ascii=False, indent=2))
    else:
        for report in reports:
            _print_report(report)
    return status


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
