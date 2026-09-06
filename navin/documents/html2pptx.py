"""Convert HTML slide masters into a native, editable PPTX deck.

The document template library ships 1920x1080 HTML slides. Rasterizing each
one into a full-bleed picture keeps the design pixel-perfect but delivers a
deck nobody can edit: no text to fix, no image to swap, no translation
possible. This module keeps the design *and* the editability.

Every slide goes through headless Chromium twice:

1. an injected script measures the laid-out page (boxes, typography, colors,
   images) and writes the result as JSON inside the DOM, which is read back
   with ``--dump-dom``. The same script then neutralizes everything it
   measured (text turns transparent, converted pictures and solid blocks turn
   invisible);
2. a screenshot of that neutralized page captures whatever could not be
   translated into Office primitives: gradients, SVG, shadows, pseudo
   elements. It becomes the slide backdrop, and is skipped entirely when the
   remaining decor is a flat color.

The deck is then assembled with python-pptx: backdrop, solid shapes, pictures,
then real text boxes on top. Chromium does the layout, PowerPoint owns the
content.

Usage::

    python3 -m navin.documents.html2pptx slides/ -o deck.pptx
    python3 -m navin.documents.html2pptx slide_01.html slide_02.html -o deck.pptx
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Sequence

if __package__:
    from . import _dom
    from ._assets import flat_color, resolve_asset
    from ._audit import repeated_pictures, unreadable_texts
    from ._chromium import ConversionError, find_chromium, measure
    from ._fonts import office_font
else:  # copied as a standalone folder into a user workspace
    import _dom
    from _assets import flat_color, resolve_asset
    from _audit import repeated_pictures, unreadable_texts
    from _chromium import ConversionError, find_chromium, measure
    from _fonts import office_font

SLIDE_WIDTH_PX = 1920
SLIDE_HEIGHT_PX = 1080
# A 1920px-wide slide maps onto the 13.333in widescreen canvas, so one CSS
# pixel is exactly 6350 EMU and 0.5pt. Every geometry conversion uses these.
EMU_PER_PX = 6350
PT_PER_PX = 0.5

# Measures the slide, then hides what it measured so the screenshot of the
# same page yields the decor layer alone.
_MEASURE_BODY = r"""
  const out = { texts: [], shapes: [], images: [], charts: [], overflow: 0 };
  let uid = 0;

  // A theme token or a CSS colour as 6-digit hex, whatever notation it uses.
  const hexOf = (raw) => {
    const value = (raw || "").trim();
    if (!value) return null;
    const short = /^#([0-9a-f])([0-9a-f])([0-9a-f])$/i.exec(value);
    if (short) return (short[1] + short[1] + short[2] + short[2] + short[3] + short[3]).toUpperCase();
    const long = /^#([0-9a-f]{6})/i.exec(value);
    if (long) return long[1].toUpperCase();
    const parsed = color(value);
    return parsed ? parsed.hex : null;
  };

  // A block carrying data-chart is data, not decor: it leaves the page as a
  // native chart the reader can edit, and its CSS drawing is dropped.
  const chartSpec = (el, cs, rect) => {
    const raw = el.getAttribute("data-chart");
    if (!raw) return null;
    let spec = null;
    try { spec = JSON.parse(raw); } catch (e) { return null; }
    if (!spec || typeof spec !== "object") return null;
    const rootStyle = getComputedStyle(document.documentElement);
    const token = (name) => hexOf(rootStyle.getPropertyValue(name)) || hexOf(cs.getPropertyValue(name));
    const palette = ["--nv-accent", "--nv-accent-2", "--nv-accent2", "--nv-text", "--nv-muted"]
      .map(token)
      .filter(Boolean)
      .filter((hex, index, all) => all.indexOf(hex) === index);
    const ink = hexOf(cs.color) || "333333";
    return {
      x: round(rect.left),
      y: round(rect.top),
      w: round(rect.width),
      h: round(rect.height),
      spec,
      palette,
      ink,
      muted: token("--nv-muted") || ink,
      grid: token("--nv-border") || null,
      font: fontFamily(cs.fontFamily),
      families: fontStack(cs.fontFamily),
      size: round(parseFloat(cs.fontSize) || 20),
    };
  };
  // The frame is pinned to the slide size, so anything reaching past it is
  // content the deck will clip: worth reporting rather than silently cutting.
  const noteOverflow = (rect) => {
    out.overflow = Math.max(out.overflow, round(rect.bottom - H), round(rect.right - W));
  };

  const solidShape = (el, cs, rect) => {
    // The decor screenshot is a single flat layer, so a full-bleed block
    // converted into a shape would sit on top of it and hide the artwork.
    // Page-sized backgrounds stay in the screenshot.
    if (rect.width * rect.height > 0.8 * W * H) return null;
    if (cs.backgroundImage && cs.backgroundImage !== "none") return null;
    if (cs.boxShadow && cs.boxShadow !== "none") return null;
    if (cs.filter && cs.filter !== "none") return null;
    if (cs.clipPath && cs.clipPath !== "none") return null;
    if (cs.mixBlendMode && cs.mixBlendMode !== "normal") return null;
    const transform = cs.transform;
    if (transform && transform !== "none" && !/^matrix\(1, 0, 0, 1,/.test(transform)) return null;
    const fill = color(cs.backgroundColor);
    const widths = [
      parseFloat(cs.borderTopWidth) || 0,
      parseFloat(cs.borderRightWidth) || 0,
      parseFloat(cs.borderBottomWidth) || 0,
      parseFloat(cs.borderLeftWidth) || 0,
    ];
    const uniformBorder = widths.every((w) => Math.abs(w - widths[0]) < 0.5);
    const borderColor = color(cs.borderTopColor);
    const hasBorder = uniformBorder && widths[0] > 0.5 && !!borderColor;
    if (!fill && !hasBorder) return null;
    if (!uniformBorder) return null;
    if (fill && fill.alpha < 0.98) return null;
    const radii = [
      parseFloat(cs.borderTopLeftRadius) || 0,
      parseFloat(cs.borderTopRightRadius) || 0,
      parseFloat(cs.borderBottomRightRadius) || 0,
      parseFloat(cs.borderBottomLeftRadius) || 0,
    ];
    if (!radii.every((r) => Math.abs(r - radii[0]) < 0.5)) return null;
    return {
      x: round(rect.left),
      y: round(rect.top),
      w: round(rect.width),
      h: round(rect.height),
      fill: fill ? fill.hex : null,
      radius: round(radii[0]),
      border: hasBorder ? { width: round(widths[0]), color: borderColor.hex } : null,
    };
  };

  // The slide title, so it can land in PowerPoint's title placeholder instead
  // of an anonymous textbox: that placeholder is what fills the Outline view,
  // names the slide in the thumbnail panel, and lets a reader navigate a deck
  // of forty slides.
  const titleEl = document.querySelector(".nv-title") ||
    document.querySelector(".nv-slide h1, .nv-slide h2") ||
    document.querySelector("h1, h2");

  // A photo under a scrim (a gradient or translucent block painted above it)
  // must stay in the decor screenshot: extracting it would re-add it on top
  // of the flattened scrim and undo the darkening that keeps the cover text
  // readable.
  const overlays = [];
  for (const el of document.querySelectorAll("body *")) {
    const cs = getComputedStyle(el);
    const rect = el.getBoundingClientRect();
    if (!visible(cs, rect)) continue;
    const gradient = (cs.backgroundImage || "").includes("gradient(");
    const bg = color(cs.backgroundColor);
    const translucent = !!bg && bg.alpha < 0.97;
    const blurred = (cs.backdropFilter || "none") !== "none";
    if (gradient || translucent || blurred) overlays.push({ el, rect });
  }
  const underOverlay = (el, rect) => {
    const area = rect.width * rect.height;
    if (area <= 0) return false;
    for (const o of overlays) {
      if (o.el === el) continue;
      const above = el.contains(o.el) ||
        (el.compareDocumentPosition(o.el) & Node.DOCUMENT_POSITION_FOLLOWING);
      if (!above) continue;
      const ix = Math.max(0, Math.min(rect.right, o.rect.right) - Math.max(rect.left, o.rect.left));
      const iy = Math.max(0, Math.min(rect.bottom, o.rect.bottom) - Math.max(rect.top, o.rect.top));
      if (ix * iy >= 0.5 * area) return true;
    }
    return false;
  };

  const walk = (el) => {
    const cs = getComputedStyle(el);
    const rect = el.getBoundingClientRect();
    if (SKIP_TAGS.has(el.tagName)) return;
    if (!visible(cs, rect)) return;

    const chart = chartSpec(el, cs, rect);
    if (chart) {
      const id = `c${uid++}`;
      el.setAttribute("data-h2x-chart", id);
      chart.id = id;
      out.charts.push(chart);
      noteOverflow(rect);
      return;
    }

    const picture = pictureSource(el, cs);
    if (picture && embeddable(picture.src) && !underOverlay(el, rect)) {
      const id = `p${uid++}`;
      el.setAttribute("data-h2x-picture", id);
      out.images.push({
        id,
        src: picture.src,
        fit: picture.fit,
        x: round(rect.left),
        y: round(rect.top),
        w: round(rect.width),
        h: round(rect.height),
      });
      if (el.tagName === "IMG") return;
    } else if (el.tagName === "IMG") {
      return;
    } else {
      const shape = solidShape(el, cs, rect);
      if (shape) {
        const id = `s${uid++}`;
        // Fill and border are neutralized independently: a block can keep a
        // translucent background in the decor while its border becomes native.
        if (shape.fill) el.setAttribute("data-h2x-fill", id);
        if (shape.border) el.setAttribute("data-h2x-line", id);
        shape.id = id;
        out.shapes.push(shape);
      }
    }

    const blockKids = Array.from(el.children).filter(
      (kid) => !SKIP_TAGS.has(kid.tagName) && isBlockish(kid) && hasText(kid),
    );
    if (hasText(el) && blockKids.length === 0) {
      const runs = runsOf(el);
      if (runs.length) {
        noteOverflow(rect);
        const box = contentBox(el, cs, rect);
        const lineHeight = parseFloat(cs.lineHeight);
        const id = `t${uid++}`;
        el.setAttribute("data-h2x-text", id);
        // A substituted font is often wider than the web font it replaces, so
        // a single-line label could wrap and break the layout. Counting the
        // line boxes here lets the deck disable wrapping for those.
        const range = document.createRange();
        range.selectNodeContents(el);
        const tops = new Set(
          Array.from(range.getClientRects()).map((r) => Math.round(r.top)),
        );
        out.texts.push({
          id,
          title: el === titleEl,
          x: round(box.x),
          y: round(box.y),
          w: round(box.w),
          h: round(box.h),
          align: cs.textAlign,
          rtl: cs.direction === "rtl",
          lineHeight: Number.isFinite(lineHeight) ? round(lineHeight) : null,
          singleLine: tops.size <= 1,
          runs,
        });
      }
      return;
    }
    for (const kid of Array.from(el.children)) walk(kid);
  };

  walk(document.body);

  // Neutralize everything already translated into Office primitives so the
  // screenshot only carries the decor that has no native equivalent.
  const style = document.createElement("style");
  style.textContent = `
    [data-h2x-text], [data-h2x-text] * {
      color: transparent !important;
      -webkit-text-fill-color: transparent !important;
      text-shadow: none !important;
      text-decoration-color: transparent !important;
      caret-color: transparent !important;
    }
    [data-h2x-gradient-text] {
      background-image: none !important;
      -webkit-background-clip: border-box !important;
      background-clip: border-box !important;
    }
    [data-h2x-picture] { visibility: hidden !important; }
    [data-h2x-chart] { visibility: hidden !important; }
    [data-h2x-fill] { background-color: transparent !important; }
    [data-h2x-line] { border-color: transparent !important; }
  `;
  document.head.appendChild(style);
  publish(out);
"""


@dataclass
class SlideData:
    """One measured slide: decor screenshot plus native element descriptions."""

    source: Path
    texts: list[dict[str, Any]] = field(default_factory=list)
    shapes: list[dict[str, Any]] = field(default_factory=list)
    images: list[dict[str, Any]] = field(default_factory=list)
    charts: list[dict[str, Any]] = field(default_factory=list)
    backdrop: Path | None = None
    backdrop_color: str | None = None
    overflow: float = 0.0


def measure_slide(slide: Path, chromium: str, workdir: Path, timeout: int) -> SlideData:
    """Measure one slide and capture its decor layer."""
    backdrop = workdir / f"__h2x_bg_{slide.stem}.png"
    payload = measure(
        slide,
        _dom.script(_MEASURE_BODY, SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX),
        chromium=chromium,
        workdir=workdir,
        timeout=timeout,
        viewport=(SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX),
        prefix="slide",
        screenshot=backdrop,
        frame=(SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX),
    )
    data = SlideData(
        source=slide,
        texts=payload.get("texts", []),
        shapes=payload.get("shapes", []),
        images=payload.get("images", []),
        charts=payload.get("charts", []),
        overflow=float(payload.get("overflow") or 0),
    )
    if backdrop.is_file():
        flat = flat_color(backdrop)
        if flat:
            data.backdrop_color = flat
        else:
            data.backdrop = backdrop
    for index, image in enumerate(data.images):
        image["path"] = resolve_asset(image.get("src", ""), slide, workdir, index)
    return data


def _emu(value: float) -> int:
    return int(round(value * EMU_PER_PX))


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_color)


def _add_backdrop(slide, data: SlideData, width: int, height: int) -> None:
    if data.backdrop_color:
        background = slide.background
        background.fill.solid()
        background.fill.fore_color.rgb = _rgb(data.backdrop_color)
        return
    if data.backdrop:
        slide.shapes.add_picture(str(data.backdrop), 0, 0, width=width, height=height)


def _add_shapes(slide, shapes: list[dict[str, Any]]) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    for item in shapes:
        radius = float(item.get("radius") or 0)
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 1 else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            kind,
            _emu(item["x"]),
            _emu(item["y"]),
            _emu(item["w"]),
            _emu(item["h"]),
        )
        shape.shadow.inherit = False
        if radius > 1:
            shortest = max(min(float(item["w"]), float(item["h"])), 1.0)
            shape.adjustments[0] = min(radius / shortest, 0.5)
        if item.get("fill"):
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(item["fill"])
        else:
            shape.fill.background()
        border = item.get("border")
        if border:
            shape.line.color.rgb = _rgb(border["color"])
            shape.line.width = Pt(float(border["width"]) * PT_PER_PX)
        else:
            shape.line.fill.background()
        shape.text_frame.text = ""


def _picture_crop(fit: str, box_w: float, box_h: float, image_path: Path) -> tuple[float, float]:
    """Crop fractions emulating object-fit: cover for a box aspect mismatch."""
    if fit != "cover":
        return 0.0, 0.0
    try:
        from PIL import Image

        with Image.open(image_path) as image:
            source_w, source_h = image.size
    except Exception:  # pragma: no cover - unreadable image
        return 0.0, 0.0
    if not source_w or not source_h or box_w <= 0 or box_h <= 0:
        return 0.0, 0.0
    box_ratio = box_w / box_h
    source_ratio = source_w / source_h
    if abs(box_ratio - source_ratio) < 0.01:
        return 0.0, 0.0
    if source_ratio > box_ratio:
        visible = box_ratio / source_ratio
        return (1 - visible) / 2, 0.0
    visible = source_ratio / box_ratio
    return 0.0, (1 - visible) / 2


def _add_images(slide, images: list[dict[str, Any]]) -> None:
    for item in images:
        path = item.get("path")
        if not path:
            continue
        picture = slide.shapes.add_picture(
            str(path),
            _emu(item["x"]),
            _emu(item["y"]),
            width=_emu(item["w"]),
            height=_emu(item["h"]),
        )
        crop_x, crop_y = _picture_crop(
            str(item.get("fit") or "fill"), float(item["w"]), float(item["h"]), Path(path)
        )
        if crop_x:
            picture.crop_left = crop_x
            picture.crop_right = crop_x
        if crop_y:
            picture.crop_top = crop_y
            picture.crop_bottom = crop_y


_DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"

# data-chart "type" -> python-pptx chart type. Names follow what a designer
# says ("column", "stacked bar") rather than the Office enum.
_CHART_TYPES: dict[str, str] = {
    "column": "COLUMN_CLUSTERED",
    "bar": "BAR_CLUSTERED",
    "stacked-column": "COLUMN_STACKED",
    "stacked-bar": "BAR_STACKED",
    "line": "LINE_MARKERS",
    "area": "AREA",
    "stacked-area": "AREA_STACKED",
    "pie": "PIE",
    "doughnut": "DOUGHNUT",
    "donut": "DOUGHNUT",
    "radar": "RADAR",
}
_ROUND_CHARTS = {"PIE", "DOUGHNUT"}
_LINE_CHARTS = {"LINE_MARKERS", "RADAR"}
_NUMBER_RE = re.compile(r"[-+]?\d[\d\s\u202f\u00a0.,]*")


def chart_number(raw: Any) -> float | None:
    """Read a number out of ``"1 240,5"``, ``"$3.2M"``, ``"48%"`` or ``12``."""
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    text = str(raw or "").strip()
    match = _NUMBER_RE.search(text)
    if not match:
        return None
    digits = match.group(0).replace(" ", "").replace("\u202f", "").replace("\u00a0", "")
    # "1.240,5" and "1,240.5" both mean 1240.5: the last separator is decimal
    # when it is followed by one or two digits, otherwise it groups thousands.
    if "," in digits and "." in digits:
        decimal = "," if digits.rfind(",") > digits.rfind(".") else "."
        digits = digits.replace("." if decimal == "," else ",", "").replace(decimal, ".")
    elif "," in digits:
        head, _, tail = digits.rpartition(",")
        digits = f"{head}.{tail}" if len(tail) in (1, 2) and head else digits.replace(",", "")
    elif digits.count(".") > 1:
        digits = digits.replace(".", "")
    try:
        value = float(digits)
    except ValueError:
        return None
    lowered = text.lower()
    for suffix, factor in (("bn", 1e9), ("md", 1e9), ("b", 1e9), ("m", 1e6), ("k", 1e3)):
        if re.search(rf"\d\s*{suffix}\b", lowered):
            return value * factor
    return value


def normalize_chart_spec(spec: dict[str, Any]) -> dict[str, Any] | None:
    """Turn the loose ``data-chart`` JSON into categories and numeric series.

    Accepted shapes: ``{"categories": [...], "series": [{"name", "values"}]}``,
    ``{"series": [{"label", "value"}, ...]}`` (one series of points), or
    ``{"points": [...]}``. Returns None when no number survives.
    """
    kind = str(spec.get("type") or "column").lower().replace("_", "-")
    if kind not in _CHART_TYPES:
        kind = "column"
    categories = [str(c) for c in (spec.get("categories") or [])]
    series_out: list[tuple[str, list[float | None]]] = []
    raw_series = spec.get("series") or spec.get("points") or []
    if raw_series and all(isinstance(s, dict) and "values" in s for s in raw_series):
        for entry in raw_series:
            values = [chart_number(v) for v in (entry.get("values") or [])]
            series_out.append((str(entry.get("name") or f"Series {len(series_out) + 1}"), values))
        if not categories and series_out:
            categories = [str(i + 1) for i in range(len(series_out[0][1]))]
    else:
        labels: list[str] = []
        values: list[float | None] = []
        for entry in raw_series:
            if isinstance(entry, dict):
                labels.append(str(entry.get("label") or entry.get("name") or entry.get("category") or ""))
                values.append(chart_number(entry.get("value")))
            elif isinstance(entry, (list, tuple)) and len(entry) >= 2:
                labels.append(str(entry[0]))
                values.append(chart_number(entry[1]))
            else:
                labels.append(str(len(labels) + 1))
                values.append(chart_number(entry))
        if not categories:
            categories = labels
        series_out.append((str(spec.get("name") or spec.get("title") or ""), values))
    series_out = [
        (name, values) for name, values in series_out if any(v is not None for v in values)
    ]
    if not series_out or not categories:
        return None
    width = len(categories)
    series_out = [(name, (values + [None] * width)[:width]) for name, values in series_out]
    number_format = spec.get("number_format") or spec.get("format")
    if not number_format:
        sample = " ".join(
            str(v) for s in (spec.get("series") or []) if isinstance(s, dict)
            for v in ([s.get("value")] + list(s.get("values") or []))
        )
        number_format = '0"%"' if "%" in sample else "#,##0.##"
    return {
        "type": kind,
        "categories": categories,
        "series": series_out,
        "number_format": str(number_format),
        "labels": spec.get("labels", True) is not False,
        "legend": spec.get("legend"),
        "colors": [str(c).lstrip("#").upper() for c in (spec.get("colors") or []) if c],
        "value_axis": spec.get("value_axis", True) is not False,
    }


def _no_fill(element) -> None:
    """Give a chart part no fill, so it sits on the slide decor as the page did."""
    from lxml import etree

    for existing in element.findall(f"{{{_CHART_NS}}}spPr"):
        element.remove(existing)
    properties = etree.SubElement(element, f"{{{_CHART_NS}}}spPr")
    etree.SubElement(properties, f"{{{_DRAWINGML_NS}}}noFill")
    line = etree.SubElement(properties, f"{{{_DRAWINGML_NS}}}ln")
    etree.SubElement(line, f"{{{_DRAWINGML_NS}}}noFill")


def _add_charts(slide, charts: list[dict[str, Any]], keep_fonts: bool = False) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt

    for item in charts:
        normalized = normalize_chart_spec(item.get("spec") or {})
        if normalized is None:
            continue
        data = CategoryChartData()
        data.categories = normalized["categories"]
        for name, values in normalized["series"]:
            data.add_series(name, values, number_format=normalized["number_format"])
        type_name = _CHART_TYPES[normalized["type"]]
        frame = slide.shapes.add_chart(
            getattr(XL_CHART_TYPE, type_name),
            _emu(item["x"]),
            _emu(item["y"]),
            _emu(item["w"]),
            _emu(item["h"]),
            data,
        )
        chart = frame.chart
        palette = normalized["colors"] or list(item.get("palette") or []) or ["4472C4", "ED7D31", "A5A5A5"]
        ink = str(item.get("ink") or "333333")
        muted = str(item.get("muted") or ink)
        size_pt = max(8.0, float(item.get("size") or 20) * PT_PER_PX)
        font_name = office_font(item.get("families") or [], item.get("font") or "Calibri", keep=keep_fonts)
        chart.font.name = font_name
        chart.font.size = Pt(size_pt)
        chart.font.color.rgb = _rgb(ink)
        chart.has_title = False
        _no_fill(chart._chartSpace)
        _no_fill(chart._chartSpace.chart.plotArea)

        multi = len(normalized["series"]) > 1
        round_chart = type_name in _ROUND_CHARTS
        wants_legend = normalized["legend"]
        chart.has_legend = bool(wants_legend) if wants_legend is not None else (multi or round_chart)
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(max(8.0, size_pt * 0.85))
            chart.legend.font.color.rgb = _rgb(muted)

        plot = chart.plots[0]
        if normalized["labels"]:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format = normalized["number_format"]
            labels.number_format_is_linked = False
            labels.font.size = Pt(max(8.0, size_pt * 0.9))
            labels.font.color.rgb = _rgb(ink)
            if round_chart:
                labels.show_percentage = normalized["number_format"].endswith('"%"')
                labels.show_value = not labels.show_percentage
        if not round_chart and type_name not in _LINE_CHARTS and hasattr(plot, "gap_width"):
            plot.gap_width = 60
        if round_chart:
            for index, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(palette[index % len(palette)])
        else:
            for index, series in enumerate(plot.series):
                colour = _rgb(palette[index % len(palette)])
                if type_name in _LINE_CHARTS:
                    series.format.line.color.rgb = colour
                    series.format.line.width = Pt(2.5)
                    series.smooth = False
                else:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = colour
                    series.format.line.fill.background()
            axis_size = Pt(max(8.0, size_pt * 0.85))
            category_axis = chart.category_axis
            category_axis.tick_labels.font.size = axis_size
            category_axis.tick_labels.font.color.rgb = _rgb(muted)
            category_axis.format.line.fill.background()
            category_axis.has_major_gridlines = False
            value_axis = chart.value_axis
            value_axis.visible = normalized["value_axis"]
            value_axis.tick_labels.font.size = axis_size
            value_axis.tick_labels.font.color.rgb = _rgb(muted)
            value_axis.tick_labels.number_format = normalized["number_format"]
            value_axis.tick_labels.number_format_is_linked = False
            value_axis.format.line.fill.background()
            value_axis.has_major_gridlines = True
            # Gridlines are the slide ink at a quarter strength: readable on the
            # theme's own paper and on an inverted slide alike.
            gridlines = value_axis.major_gridlines.format.line
            gridlines.color.rgb = _rgb(ink)
            gridlines.width = Pt(0.75)
            _set_alpha(gridlines, 25)


def _set_alpha(line_format, percent: int) -> None:
    """Add transparency to a line colour (python-pptx exposes none)."""
    from lxml import etree

    fill = line_format._ln.find(f"{{{_DRAWINGML_NS}}}solidFill") if line_format._ln is not None else None
    if fill is None:
        return
    colour = fill.find(f"{{{_DRAWINGML_NS}}}srgbClr")
    if colour is None:
        return
    for existing in colour.findall(f"{{{_DRAWINGML_NS}}}alpha"):
        colour.remove(existing)
    etree.SubElement(colour, f"{{{_DRAWINGML_NS}}}alpha").set("val", str(percent * 1000))


def _apply_gradient_fill(run, gradient: dict[str, Any]) -> None:
    """Give a run PowerPoint's own gradient text fill, kept fully editable."""
    from lxml import etree

    stops = gradient.get("stops") or []
    if len(stops) < 2:
        return
    properties = run.font._rPr
    for existing in properties.findall(f"{{{_DRAWINGML_NS}}}solidFill"):
        properties.remove(existing)
    fill = etree.SubElement(properties, f"{{{_DRAWINGML_NS}}}gradFill")
    stop_list = etree.SubElement(fill, f"{{{_DRAWINGML_NS}}}gsLst")
    for stop in stops:
        node = etree.SubElement(stop_list, f"{{{_DRAWINGML_NS}}}gs")
        node.set("pos", str(int(round(float(stop.get("pos") or 0) * 1000))))
        etree.SubElement(node, f"{{{_DRAWINGML_NS}}}srgbClr").set(
            "val", str(stop.get("hex") or "000000")
        )
    # CSS measures from "to top" clockwise, DrawingML from "to right".
    angle = (float(gradient.get("angle", 180)) - 90) % 360
    line = etree.SubElement(fill, f"{{{_DRAWINGML_NS}}}lin")
    line.set("ang", str(int(round(angle * 60000))))
    line.set("scaled", "0")
    # The fill has to precede a:latin and friends in the schema sequence.
    properties.remove(fill)
    properties.insert(0, fill)


_ALIGNMENTS = {
    "left": "LEFT",
    "start": "LEFT",
    "right": "RIGHT",
    "end": "RIGHT",
    "center": "CENTER",
    "justify": "JUSTIFY",
}


def _paragraph_lines(runs: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    """Split runs on explicit line breaks into per-paragraph run lists."""
    lines: list[list[dict[str, Any]]] = [[]]
    for run in runs:
        if run.get("newline"):
            lines.append([])
            continue
        for index, chunk in enumerate(str(run.get("text", "")).split("\n")):
            if index:
                lines.append([])
            if chunk:
                lines[-1].append({**run, "text": chunk})
    return [line for line in lines if line] or [[]]


def _add_texts(
    slide,
    texts: list[dict[str, Any]],
    keep_fonts: bool = False,
    title_shape=None,
) -> None:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Pt

    for item in texts:
        if title_shape is not None and item.get("title"):
            # The placeholder carries the structure; the measurement carries the
            # design, so it is moved onto the box the page actually drew rather
            # than left where the master would put it.
            box = title_shape
            box.left, box.top = _emu(item["x"]), _emu(item["y"])
            box.width, box.height = _emu(item["w"]), _emu(item["h"])
            title_shape = None
        else:
            box = slide.shapes.add_textbox(
                _emu(item["x"]),
                _emu(item["y"]),
                _emu(item["w"]),
                _emu(item["h"]),
            )
        frame = box.text_frame
        # A placeholder inherits the master's autofit, which would resize the
        # measured type the moment the text is set.
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.word_wrap = not item.get("singleLine")
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        alignment = _ALIGNMENTS.get(str(item.get("align") or "left"), "LEFT")
        line_height = item.get("lineHeight")
        for index, line in enumerate(_paragraph_lines(item.get("runs", []))):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = getattr(PP_ALIGN, alignment)
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            if line_height:
                paragraph.line_spacing = Pt(float(line_height) * PT_PER_PX)
            if item.get("rtl") and paragraph._pPr is not None:
                paragraph._pPr.set("rtl", "1")
            for run_data in line:
                run = paragraph.add_run()
                run.text = run_data.get("text", "")
                font = run.font
                font.name = office_font(
                    run_data.get("families") or [],
                    run_data.get("font") or "Calibri",
                    keep=keep_fonts,
                )
                font.size = Pt(float(run_data.get("size") or 16) * PT_PER_PX)
                font.bold = bool(run_data.get("bold"))
                font.italic = bool(run_data.get("italic"))
                font.underline = bool(run_data.get("underline"))
                if run_data.get("color"):
                    font.color.rgb = _rgb(run_data["color"])
                if run_data.get("gradient"):
                    _apply_gradient_fill(run, run_data["gradient"])
                spacing = float(run_data.get("spacing") or 0)
                if abs(spacing) > 0.05:
                    run.font._rPr.set("spc", str(int(round(spacing * PT_PER_PX * 100))))


def build_deck(slides: list[SlideData], output: Path, keep_fonts: bool = False) -> Path:
    """Assemble measured slides into a native PPTX file."""
    from pptx import Presentation

    presentation = Presentation()
    presentation.slide_width = _emu(SLIDE_WIDTH_PX)
    presentation.slide_height = _emu(SLIDE_HEIGHT_PX)
    blank = presentation.slide_layouts[6]
    # "Title Only": the one native layout that contributes a title placeholder
    # and nothing else, so the deck gains an outline without gaining a body
    # box nobody asked for.
    titled_layout = presentation.slide_layouts[5]
    deck_title = ""
    for data in slides:
        headed = any(text.get("title") for text in data.texts)
        slide = presentation.slides.add_slide(titled_layout if headed else blank)
        title_shape = slide.shapes.title if headed else None
        _add_backdrop(slide, data, presentation.slide_width, presentation.slide_height)
        _add_shapes(slide, data.shapes)
        _add_images(slide, data.images)
        _add_charts(slide, data.charts, keep_fonts=keep_fonts)
        _add_texts(slide, data.texts, keep_fonts=keep_fonts, title_shape=title_shape)
        if not deck_title and title_shape is not None:
            deck_title = title_shape.text_frame.text.strip()
        if title_shape is not None:
            element = title_shape._element
            tree = element.getparent()
            if title_shape.text_frame.text.strip():
                # The placeholder is cloned with the layout, so it sits under
                # the decor picture added after it. Moving it to the end of the
                # tree puts it back on top, where the page drew it.
                tree.append(element)
            else:  # pragma: no cover - a title was measured but carried no text
                tree.remove(element)
    # The first slide's title names the deck in the file properties: what the
    # reader sees in the recent-files list and what a search indexes.
    presentation.core_properties.title = (deck_title or output.stem)[:255]
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    return output


def collect_slides(inputs: Sequence[str]) -> list[Path]:
    """Expand files and directories into an ordered list of slide files."""
    slides: list[Path] = []
    for raw in inputs:
        path = Path(raw).expanduser()
        if path.is_dir():
            slides.extend(sorted(path.glob("slide_*.html")) or sorted(path.glob("*.html")))
        elif path.is_file():
            slides.append(path)
        else:
            raise ConversionError(f"Not found: {raw}")
    if not slides:
        raise ConversionError("No HTML slide to convert")
    return slides


def check_critique(slides: Sequence[Path], force: bool = False) -> None:
    """Refuse to ship a deck its own quality pass already rejected.

    ``ppt_design deck`` writes ``critique.json`` beside the slides and exits
    non-zero below the threshold, but an exit code is easy to ignore. This
    gate is not: a failing critique stops the conversion until the slides are
    fixed (or the caller passes ``--force`` because the user accepts the deck
    as is).
    """
    for folder in {slide.parent for slide in slides}:
        critique = folder / "critique.json"
        if not critique.is_file():
            continue
        try:
            data = json.loads(critique.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if data.get("pass") is not False:
            continue
        score = data.get("score")
        threshold = data.get("threshold")
        notes = "; ".join(str(note) for note in (data.get("notes") or [])[:4])
        message = (
            f"{critique}: the deck critique failed (score {score}, "
            f"threshold {threshold}). {notes}"
        )
        if force:
            print(f"html2pptx: {message} (converted anyway: --force)", file=sys.stderr)
            continue
        raise ConversionError(
            message
            + " - fix the slides (rework the deck JSON, re-run ppt_design deck)"
            " until the critique passes, then convert. Use --force only if the"
            " user explicitly accepts the deck as is."
        )


def report(measured: Sequence[SlideData]) -> list[str]:
    """Print what a reader would notice on the finished deck, slide by slide."""
    warnings: list[str] = []
    # Template furniture repeats on every slide, so an issue is listed once with
    # the slides it affects rather than once per slide.
    seen: dict[str, list[str]] = {}
    for data in measured:
        name = data.source.name
        if data.overflow > 2:
            warnings.append(
                f"{name}: content runs {data.overflow:.0f}px past the "
                f"{SLIDE_WIDTH_PX}x{SLIDE_HEIGHT_PX} frame and gets clipped; shorten the copy"
            )
        for issue in unreadable_texts(data.texts, data.backdrop, data.shapes):
            seen.setdefault(issue, []).append(name)
    for issue, slides_hit in seen.items():
        where = ", ".join(slides_hit[:3]) + (" and others" if len(slides_hit) > 3 else "")
        warnings.append(f"{where}: unreadable text {issue}; darken it or lighten the block")
    pictures = [
        picture
        for data in measured
        for picture in (Path(image["path"]) for image in data.images if image.get("path"))
    ]
    for issue in repeated_pictures(pictures):
        warnings.append(f"deck: {issue}; use distinct visuals or drop the duplicates")
    for warning in warnings:
        print(f"html2pptx: {warning}", file=sys.stderr)
    return warnings


def convert(
    inputs: Sequence[str],
    output: Path,
    *,
    chromium: str | None = None,
    workers: int = 4,
    timeout: int = 60,
    keep_workdir: Path | None = None,
    keep_fonts: bool = False,
    force: bool = False,
) -> Path:
    """Convert HTML slides into an editable deck and return the output path."""
    slides = collect_slides(inputs)
    check_critique(slides, force=force)
    binary = find_chromium(chromium)
    workdir = keep_workdir or Path(tempfile.mkdtemp(prefix="navin-html2pptx-"))
    workdir.mkdir(parents=True, exist_ok=True)
    try:
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            measured = list(
                pool.map(lambda slide: measure_slide(slide, binary, workdir, timeout), slides)
            )
        report(measured)
        built = build_deck(measured, output, keep_fonts=keep_fonts)
        _apply_notes(slides, built)
        return built
    finally:
        if keep_workdir is None:
            shutil.rmtree(workdir, ignore_errors=True)


_NOTES_RE = re.compile(r'data-notes="([^"]*)"')


def _apply_notes(slides: Sequence[Path], deck: Path) -> int:
    """Copy the speaker notes stamped on each slide (``data-notes``) into the deck.

    ppt_design writes them on the slide root; a deck converted without them
    would need a second command the operator has to remember.
    """
    import html as html_module

    notes: list[str] = []
    for slide in slides:
        try:
            match = _NOTES_RE.search(slide.read_text(encoding="utf-8", errors="ignore"))
        except OSError:
            match = None
        notes.append(html_module.unescape(match.group(1)).strip() if match else "")
    if not any(notes):
        return 0
    from pptx import Presentation

    presentation = Presentation(str(deck))
    written = 0
    for page, text in zip(presentation.slides, notes, strict=False):
        if text:
            page.notes_slide.notes_text_frame.text = text
            written += 1
    if written:
        presentation.save(str(deck))
    return written


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="html2pptx",
        description="Convert 1920x1080 HTML slides into an editable PPTX deck.",
    )
    parser.add_argument("inputs", nargs="+", help="Slide files, or a folder of slide_*.html")
    parser.add_argument("-o", "--output", required=True, help="Destination .pptx")
    parser.add_argument("--chromium", help="Path to a Chromium binary")
    parser.add_argument("--workers", type=int, default=4, help="Slides rendered in parallel")
    parser.add_argument("--timeout", type=int, default=60, help="Per-Chromium-run timeout")
    parser.add_argument("--keep-workdir", help="Keep intermediate files in this folder")
    parser.add_argument(
        "--keep-fonts",
        action="store_true",
        help="Keep CSS font names instead of mapping them onto Office fonts",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Convert even when critique.json rejected the slides",
    )
    args = parser.parse_args(argv)
    try:
        output = convert(
            args.inputs,
            Path(args.output).expanduser(),
            chromium=args.chromium,
            workers=args.workers,
            timeout=args.timeout,
            keep_workdir=Path(args.keep_workdir).expanduser() if args.keep_workdir else None,
            keep_fonts=args.keep_fonts,
            force=args.force,
        )
    except ConversionError as exc:
        print(f"html2pptx: {exc}", file=sys.stderr)
        return 1
    print(f"Wrote {output}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
