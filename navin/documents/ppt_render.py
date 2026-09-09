# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""Turn a semantic slide into finished HTML. The LLM never writes coordinates.

A lookbook layout still ships with sample copy so the picker has something to
show. A real deck must not. This module builds the slide body from the
semantic JSON in ``templates/ppt/_engine/semantic.schema.json``: title, items,
image, series. Sample phrases and ``image.png`` never enter the output.

If the image is missing, a theme panel fills that column rather than leaving
an empty frame or a leftover placeholder. Text and the visual sit in a flex
row. They do not stack on top of each other.
"""

from __future__ import annotations

import html
import json
import re
from pathlib import Path
from typing import Any, Mapping

_FAKE_IMAGES = {"", "image.png", "placeholder.png", "photo.png", "screenshot.png"}

# Copy that lives in the lookbook layouts. If any of it survives into a filled
# slide, the deck is still the sample, not the story.
SAMPLE_PHRASES = (
    "the title is the message",
    "deck kicker",
    "one sentence that states the outcome",
    "text holds the claim",
    "body copy. one idea. no filler.",
    "named fact.",
    "image leads, text follows",
    "caption-length argument.",
    "three to four equal claims",
    "01 card",
    "how the work actually moves",
    "inputs locked.",
    "working slice.",
    "arr is compounding",
    "the loop that compounds",
    "now, next, later",
    "shipping this quarter.",
    "where conversion drops",
    "visitors 100%",
    "the people who will deliver it",
    "the ask, in one sentence",
    "contact@example.com",
    "what we will decide today",
    "why this meeting exists.",
    "before and after the change",
    "named pain, with a number.",
    "the path is dated",
    "scope locked.",
    "enterprise revenue has tripled since 2024",
    "comparison the buyer can scan",
    "a sentence a real customer would say",
    "name, title, company",
    "full-bleed statement",
    "short proof under the claim.",
)


def _esc(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def _text(value: Any) -> str:
    return str(value or "").strip()


def _item_row(entry: Any) -> dict[str, str] | None:
    if isinstance(entry, str):
        label = entry.strip()
        if not label:
            return None
        return {"label": label, "detail": "", "value": "", "meta": "", "image": ""}
    if not isinstance(entry, dict):
        return None
    label = _text(
        entry.get("label")
        or entry.get("title")
        or entry.get("name")
        or entry.get("heading")
        or entry.get("phase")
    )
    detail = _text(
        entry.get("detail")
        or entry.get("body")
        or entry.get("text")
        or entry.get("description")
    )
    value = _text(entry.get("value") or entry.get("num") or entry.get("number"))
    meta = _text(
        entry.get("meta")
        or entry.get("delta")
        or entry.get("time")
        or entry.get("cite")
        or entry.get("role")
        or entry.get("step")
        or entry.get("phase")
    )
    image = _text(entry.get("image") or entry.get("src"))
    if not (label or detail or value):
        return None
    return {
        "label": label,
        "detail": detail,
        "value": value,
        "meta": meta,
        "image": image if image.lower() not in _FAKE_IMAGES else "",
    }


def items_of(slide: Mapping[str, Any]) -> list[dict[str, str]]:
    raw = slide.get("items") or []
    out: list[dict[str, str]] = []
    if not isinstance(raw, list):
        return out
    for entry in raw:
        row = _item_row(entry)
        if row:
            out.append(row)
    return out


def _join_nested(raw: Any) -> str:
    if not isinstance(raw, list):
        return _text(raw)
    lines: list[str] = []
    for entry in raw:
        if isinstance(entry, str) and entry.strip():
            lines.append(entry.strip())
            continue
        if isinstance(entry, dict):
            line = _text(
                entry.get("label")
                or entry.get("title")
                or entry.get("text")
                or entry.get("description")
                or entry.get("detail")
            )
            if line:
                lines.append(line)
    return "\n".join(lines)


def _columns_to_items(columns: Any) -> list[dict[str, str]]:
    if not isinstance(columns, list):
        return []
    out: list[dict[str, str]] = []
    for column in columns:
        if isinstance(column, str) and column.strip():
            out.append({"label": column.strip(), "detail": "", "value": "", "meta": "", "image": ""})
            continue
        if not isinstance(column, dict):
            continue
        label = _text(
            column.get("title") or column.get("label") or column.get("name") or column.get("heading")
        )
        detail = _join_nested(
            column.get("items") or column.get("points") or column.get("bullets")
        ) or _text(
            column.get("detail") or column.get("description") or column.get("body") or column.get("text")
        )
        if label or detail:
            out.append({"label": label, "detail": detail, "value": "", "meta": "", "image": ""})
    return out


def _alias_item_list(slide: Mapping[str, Any]) -> list[Any]:
    for key in ("features", "steps", "phases", "points", "cards", "bullets"):
        raw = slide.get(key)
        if isinstance(raw, list) and raw:
            return raw
    return []


def normalize_slide(slide: Mapping[str, Any]) -> dict[str, Any]:
    """Accept the field names agents actually write. The template still wins.

    A deck.json that uses ``features``, ``columns``, ``steps`` or
    ``description`` must fill the chosen theme, not render as a title and an
    empty panel. This does not invent a look. It maps content onto the
    semantic fields the theme layouts already know.
    """
    payload = dict(slide)
    if not _text(payload.get("body")):
        payload["body"] = _text(
            payload.get("description") or payload.get("copy") or payload.get("text")
        )
    if not _text(payload.get("kicker")):
        tagline = _text(payload.get("tagline") or payload.get("eyebrow"))
        if tagline:
            payload["kicker"] = tagline
    if not items_of(payload):
        columns = _columns_to_items(payload.get("columns"))
        if columns:
            payload["items"] = columns
        else:
            aliased = _alias_item_list(payload)
            if aliased:
                payload["items"] = aliased
    if not payload.get("footer") and _text(payload.get("contact")):
        payload["footer"] = {"left": "", "right": _text(payload.get("contact"))}
    return payload


def image_src(slide: Mapping[str, Any]) -> str:
    visual = slide.get("visual") if isinstance(slide.get("visual"), dict) else {}
    src = _text(slide.get("image") or visual.get("src") or visual.get("image"))
    if src.lower() in _FAKE_IMAGES:
        return ""
    return src


def _visual(slide: Mapping[str, Any]) -> dict[str, Any]:
    visual = slide.get("visual")
    return visual if isinstance(visual, dict) else {}


def _variant(slide: Mapping[str, Any], default: str) -> str:
    return _text(_visual(slide).get("variant")) or default


def _fit(slide: Mapping[str, Any], default: str = "cover") -> str:
    fit = _text(_visual(slide).get("fit"))
    return fit if fit in {"cover", "contain", "cover-face"} else default


def _kicker(slide: Mapping[str, Any]) -> str:
    kicker = _text(slide.get("kicker"))
    return f'<div class="nv-kicker">{_esc(kicker)}</div>' if kicker else ""


def _title(slide: Mapping[str, Any]) -> str:
    return f'<h1 class="nv-title">{_esc(_text(slide.get("title")))}</h1>'


def _body_p(slide: Mapping[str, Any], cls: str = "nv-body") -> str:
    body = _text(slide.get("body") or slide.get("description"))
    return f'<p class="{cls}">{_esc(body)}</p>' if body else ""


def _sub(slide: Mapping[str, Any]) -> str:
    sub = _text(slide.get("sub") or slide.get("subtitle"))
    if not sub:
        return ""
    return f'<p class="nv-sub">{_esc(sub)}</p>'


def _insight(slide: Mapping[str, Any]) -> str:
    insight = _text(slide.get("insight"))
    return f'<p class="nv-insight">{_esc(insight)}</p>' if insight else ""


def _footer(slide: Mapping[str, Any]) -> str:
    footer = slide.get("footer")
    left = right = ""
    if isinstance(footer, dict):
        left, right = _text(footer.get("left")), _text(footer.get("right"))
    elif isinstance(footer, list) and footer:
        left = _text(footer[0])
        right = _text(footer[1]) if len(footer) > 1 else ""
    if not left and not right:
        return ""
    return (
        f'<div class="nv-footer"><span>{_esc(left)}</span>'
        f"<span>{_esc(right)}</span></div>"
    )


def _img(slide: Mapping[str, Any], variant: str, *, grow: bool = True) -> str:
    src = image_src(slide)
    if not src:
        return ""
    fit = _fit(slide)
    if fit == "cover-face":
        fit = "cover"
        position = "object-position:center 18%;"
    else:
        kind = _text(_visual(slide).get("kind") or _visual(slide).get("type"))
        # Screenshots stay whole. Photos fill the frame. Never stretch.
        if kind in {"screenshot", "product", "ui"} and not _visual(slide).get("fit"):
            fit = "contain"
        position = "object-position:center center;"
    klass = "nv-img grow" if grow else "nv-img"
    alt = _esc(_text(_visual(slide).get("alt") or slide.get("title")))
    picture = (
        f'<div class="{klass}" data-variant="{_esc(variant)}" data-fit="{_esc(fit)}">'
        f'<img alt="{alt}" src="{_esc(src)}" style="{position}"></div>'
    )
    frame = _shot_kind(slide)
    if not frame:
        return picture
    return (
        f'<div class="nv-shot grow" data-kind="{_esc(frame)}">'
        f'<div class="nv-shot-bar" aria-hidden="true"><i></i><i></i><i></i></div>'
        f"{picture}</div>"
    )


def _shot_kind(slide: Mapping[str, Any]) -> str:
    """Product UI gets a CleanShot-style frame. Theme stills stay unframed."""
    visual = _visual(slide)
    kind = _text(visual.get("kind") or visual.get("type"))
    source = _text(visual.get("source"))
    if kind == "phone":
        return "phone"
    if kind in {"photo", "image"} or source == "theme":
        return ""
    if kind in {"screenshot", "product", "ui"}:
        return "browser"
    if _text(slide.get("layout")) == "product-hero" and image_src(slide):
        return "browser"
    return ""


def _list(slide: Mapping[str, Any], variant: str) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    parts = [
        f'<div class="nv-list nv-anim-stagger" data-variant="{_esc(variant)}" '
        f'data-n="{len(rows)}">'
    ]
    for row in rows:
        heading = row["label"] or row["value"]
        detail = row["detail"] or row["meta"]
        parts.append('<div class="nv-item">')
        if heading:
            parts.append(f"<h3>{_esc(heading)}</h3>")
        if detail:
            parts.append(f"<p>{_esc(detail)}</p>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _process(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    variant = _variant(slide, "chevron")
    parts = [f'<div class="nv-process nv-anim-stagger" data-variant="{_esc(variant)}">']
    for index, row in enumerate(rows, 1):
        number = row["meta"] or f"{index:02d}"
        parts.append('<div class="nv-step">')
        parts.append(f"<strong>{_esc(number)}</strong>")
        if row["label"]:
            parts.append(f"<h3>{_esc(row['label'])}</h3>")
        if row["detail"]:
            parts.append(f"<p>{_esc(row['detail'])}</p>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _cycle(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    hub = _text(slide.get("hub")) or _text(_visual(slide).get("hub")) or "Loop"
    n = max(3, min(8, len(rows)))
    parts = [
        f'<div class="nv-cycle" data-variant="{_esc(_variant(slide, "flywheel"))}" '
        f'data-n="{n}">',
        f'<div class="nv-hub">{_esc(hub)}</div>',
        '<div class="nv-ring"></div>',
    ]
    for row in rows[:n]:
        parts.append(f'<div class="nv-node"><h3>{_esc(row["label"] or row["value"])}</h3>')
        if row["detail"]:
            parts.append(f"<p>{_esc(row['detail'])}</p>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _timeline(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    variant = _variant(slide, "horizontal")
    parts = [f'<div class="nv-timeline nv-anim-stagger" data-variant="{_esc(variant)}">']
    for row in rows:
        parts.append("<div class=\"nv-tl\">")
        if row["meta"] or row["value"]:
            parts.append(f"<time>{_esc(row['meta'] or row['value'])}</time>")
        if row["label"]:
            parts.append(f"<h3>{_esc(row['label'])}</h3>")
        if row["detail"]:
            parts.append(f"<p>{_esc(row['detail'])}</p>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _roadmap(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    default = "now-next-later" if len(rows) == 3 else "quarterly"
    variant = _variant(slide, default)
    parts = [f'<div class="nv-roadmap" data-variant="{_esc(variant)}">']
    for row in rows:
        parts.append('<div class="nv-road">')
        parts.append(f"<h3>{_esc(row['label'] or row['meta'])}</h3>")
        if row["detail"] or row["value"]:
            parts.append(f"<p>{_esc(row['detail'] or row['value'])}</p>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _funnel(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    parts = [f'<div class="nv-funnel" data-variant="{_esc(_variant(slide, "conversion"))}">']
    for row in rows:
        label = " ".join(part for part in (row["label"], row["value"]) if part)
        parts.append(f'<div class="nv-stage">{_esc(label)}</div>')
    parts.append("</div>")
    return "".join(parts)


def _kpi(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows and _text(slide.get("value")):
        rows = [
            {
                "label": _text(slide.get("metric") or slide.get("body")),
                "detail": "",
                "value": _text(slide.get("value")),
                "meta": _text(slide.get("insight")),
                "image": "",
            }
        ]
    if not rows:
        return ""
    mapping = {1: "single", 2: "pair", 3: "triple", 4: "quad", 5: "five"}
    variant = _variant(slide, mapping.get(len(rows), "cards"))
    if variant == "tower":
        heights: list[int] = []
        for row in rows:
            try:
                heights.append(abs(float(re.sub(r"[^\d.]+", "", row["value"]) or "0")))
            except ValueError:
                heights.append(0)
        peak = max(heights) or 1
        parts = ['<div class="nv-kpi" data-variant="tower">']
        for row, height in zip(rows, heights, strict=False):
            pct = max(22, int(round(100 * height / peak)))
            parts.append(
                f'<div class="nv-stat" style="height:{pct}%">'
                f'<div class="nv-num">{_esc(row["value"])}</div>'
                f'<p>{_esc(row["label"])}</p></div>'
            )
        parts.append("</div>")
        return "".join(parts)
    parts = [f'<div class="nv-kpi" data-variant="{_esc(variant)}">']
    for row in rows:
        parts.append('<div class="nv-stat">')
        if row["value"]:
            parts.append(f'<div class="nv-num">{_esc(row["value"])}</div>')
        if row["label"]:
            parts.append(f"<p>{_esc(row['label'])}</p>")
        if row["meta"] or row["detail"]:
            parts.append(f'<p class="delta">{_esc(row["meta"] or row["detail"])}</p>')
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _series(slide: Mapping[str, Any]) -> list[dict[str, str]]:
    raw = slide.get("series") or []
    out: list[dict[str, str]] = []
    if not isinstance(raw, list):
        return out
    for entry in raw:
        if isinstance(entry, dict) and "points" in entry:
            for point in entry.get("points") or []:
                if isinstance(point, dict):
                    out.append(
                        {
                            "label": _text(point.get("label")),
                            "value": _text(point.get("value")),
                        }
                    )
            continue
        if isinstance(entry, dict):
            out.append(
                {
                    "label": _text(entry.get("label") or entry.get("name")),
                    "value": _text(entry.get("value")),
                }
            )
    if not out:
        for row in items_of(slide):
            if row["value"] or row["label"]:
                out.append({"label": row["label"] or row["meta"], "value": row["value"]})
    return out


_CHART_KINDS = {
    "column", "bar", "stacked-column", "stacked-bar", "line", "area", "stacked-area",
    "pie", "doughnut", "donut", "radar",
}


def _chart_number(raw: Any) -> float:
    text = str(raw or "")
    match = re.search(r"[-+]?\d[\d\s\u202f\u00a0.,]*", text)
    if not match:
        return 0.0
    digits = re.sub(r"[\s\u202f\u00a0]", "", match.group(0))
    if "," in digits and "." in digits:
        decimal = "," if digits.rfind(",") > digits.rfind(".") else "."
        digits = digits.replace("." if decimal == "," else ",", "").replace(decimal, ".")
    elif "," in digits:
        head, _, tail = digits.rpartition(",")
        digits = f"{head}.{tail}" if len(tail) in (1, 2) and head else digits.replace(",", "")
    try:
        return float(digits)
    except ValueError:
        return 0.0


def chart_spec(slide: Mapping[str, Any]) -> dict[str, Any] | None:
    """The ``data-chart`` payload html2pptx turns into a native, editable chart.

    A slide may carry a full ``chart`` object (``type``, ``categories``,
    ``series`` with ``name`` and ``values``, ``number_format``) for several
    series; otherwise its ``series`` / items become one series of points and
    the visual variant picks the chart type.
    """
    explicit = slide.get("chart") if isinstance(slide.get("chart"), dict) else {}
    kind = _text(explicit.get("type") or _variant(slide, "column")).lower().replace("_", "-")
    if kind not in _CHART_KINDS:
        kind = "column"
    spec: dict[str, Any] = {"type": kind}
    for key in ("number_format", "labels", "legend", "colors", "value_axis", "name"):
        if key in explicit:
            spec[key] = explicit[key]
    raw_series = explicit.get("series")
    if isinstance(raw_series, list) and raw_series and all(
        isinstance(entry, dict) and "values" in entry for entry in raw_series
    ):
        spec["categories"] = [_text(c) for c in (explicit.get("categories") or [])]
        spec["series"] = [
            {
                "name": _text(entry.get("name")),
                "values": [_chart_number(v) if not isinstance(v, (int, float)) else v for v in (entry.get("values") or [])],
            }
            for entry in raw_series
        ]
        if not spec["categories"]:
            spec["categories"] = [str(i + 1) for i in range(len(spec["series"][0]["values"]))]
        return spec
    # "chart.series" written as points (label / value) is the same one-series
    # chart as the slide-level series; do not lose it to the schema's shape.
    points = _series({"series": raw_series}) if isinstance(raw_series, list) else []
    if not points:
        points = _series(slide)
    if not points:
        return None
    spec["series"] = [{"label": p["label"], "value": p["value"]} for p in points]
    return spec


def _chart(slide: Mapping[str, Any]) -> str:
    spec = chart_spec(slide)
    if spec is None:
        return ""
    if "categories" in spec:
        labels = list(spec["categories"])
        numbers = [abs(float(v or 0)) for v in spec["series"][0]["values"]][: len(labels)]
    else:
        labels = [p["label"] or p["value"] for p in spec["series"]]
        numbers = [abs(_chart_number(p["value"])) for p in spec["series"]]
    if not numbers:
        return ""
    # The HTML drawing is the design proof in the browser, the QA pass and the
    # PDF; the converter drops it and places a native chart built from
    # data-chart. The drawing follows the chart type so both tell one story.
    payload = _esc(json.dumps(spec, ensure_ascii=False, separators=(",", ":")))
    kind = spec["type"]
    if kind in {"pie", "doughnut", "donut"}:
        return _chart_round(payload, labels, numbers)
    if kind in {"line", "area", "stacked-area", "radar"}:
        return _chart_lines(payload, spec, labels, numbers)
    peak = max(numbers) or 1
    variant = _variant(slide, "column")
    if variant not in {"column", "bar"}:
        variant = "bar" if kind in {"bar", "stacked-bar"} else "column"
    parts = [
        f'<div class="nv-chart nv-anim-stagger" data-variant="{_esc(variant)}" '
        f'data-chart="{payload}">'
    ]
    for label, number in zip(labels, numbers, strict=False):
        height = max(8, int(round(100 * number / peak)))
        parts.append(
            f'<div class="nv-bar" style="height:{height}%">'
            f"<span>{_esc(label)}</span></div>"
        )
    parts.append("</div>")
    return "".join(parts)


_CHART_TOKENS = ("var(--nv-accent)", "var(--nv-accent-2)", "var(--nv-muted)", "var(--nv-border)")


def _chart_round(payload: str, labels: list[str], numbers: list[float]) -> str:
    total = sum(numbers) or 1
    stops: list[str] = []
    legend: list[str] = []
    cursor = 0.0
    for index, (label, number) in enumerate(zip(labels, numbers, strict=False)):
        share = 100 * number / total
        token = _CHART_TOKENS[index % len(_CHART_TOKENS)]
        stops.append(f"{token} {cursor:.1f}% {cursor + share:.1f}%")
        cursor += share
        legend.append(
            f'<li><i style="background:{token}"></i>{_esc(label)}'
            f"<b>{share:.0f}%</b></li>"
        )
    return (
        f'<div class="nv-chart nv-chart-round nv-anim-stagger" data-chart="{payload}">'
        f'<div class="nv-donut" style="background:conic-gradient({", ".join(stops)})"></div>'
        f'<ul class="nv-chart-legend">{"".join(legend)}</ul></div>'
    )


def _chart_lines(
    payload: str, spec: Mapping[str, Any], labels: list[str], numbers: list[float]
) -> str:
    series = spec["series"] if "categories" in spec else [{"values": numbers}]
    all_values = [
        float(v or 0) for entry in series for v in (entry.get("values") or []) if v is not None
    ]
    peak = max(all_values) or 1
    width, height = 1000, 400
    count = max(len(labels), 2)
    step = width / (count - 1)
    shapes: list[str] = []
    fill = spec["type"] in {"area", "stacked-area"}
    for index, entry in enumerate(series):
        token = _CHART_TOKENS[index % len(_CHART_TOKENS)]
        points = [
            (i * step, height - (height - 20) * float(v or 0) / peak)
            for i, v in enumerate(entry.get("values") or [])
        ]
        if not points:
            continue
        path = " ".join(f"{x:.0f},{y:.0f}" for x, y in points)
        if fill:
            shapes.append(
                f'<polygon points="0,{height} {path} {points[-1][0]:.0f},{height}" '
                f'style="fill:{token};opacity:.25"/>'
            )
        shapes.append(
            f'<polyline points="{path}" style="fill:none;stroke:{token};stroke-width:6;'
            'stroke-linejoin:round;stroke-linecap:round"/>'
        )
        shapes.extend(
            f'<circle cx="{x:.0f}" cy="{y:.0f}" r="9" style="fill:{token}"/>' for x, y in points
        )
    axis = "".join(f"<span>{_esc(label)}</span>" for label in labels)
    legend = ""
    names = [_text(entry.get("name")) for entry in series]
    if len(series) > 1 and any(names):
        legend = '<ul class="nv-chart-legend">' + "".join(
            f'<li><i style="background:{_CHART_TOKENS[index % len(_CHART_TOKENS)]}"></i>'
            f"{_esc(name or f'Series {index + 1}')}</li>"
            for index, name in enumerate(names)
        ) + "</ul>"
    return (
        f'<div class="nv-chart nv-chart-lines nv-anim-stagger" data-chart="{payload}">'
        f'{legend}<svg viewBox="0 0 {width} {height}" preserveAspectRatio="none" role="img">'
        f'{"".join(shapes)}</svg><div class="nv-chart-axis">{axis}</div></div>'
    )


def _table(slide: Mapping[str, Any]) -> str:
    table = slide.get("table") if isinstance(slide.get("table"), dict) else {}
    headers = [str(cell) for cell in (table.get("headers") or [])]
    rows = table.get("rows") or []
    if not headers and items_of(slide):
        headers = ["Item", "Detail", "Value"]
        rows = [[row["label"], row["detail"], row["value"]] for row in items_of(slide)]
    if not headers:
        return ""
    variant = _variant(slide, "comparison")
    parts = [f'<table class="nv-table" data-variant="{_esc(variant)}"><thead><tr>']
    parts.extend(f"<th>{_esc(cell)}</th>" for cell in headers)
    parts.append("</tr></thead><tbody>")
    for row in rows:
        if not isinstance(row, list):
            continue
        parts.append("<tr>")
        parts.extend(f"<td>{_esc(cell)}</td>" for cell in row)
        parts.append("</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


def _comparison_copy(row: Mapping[str, str]) -> str:
    detail = row.get("detail") or row.get("value") or ""
    if "\n" in detail:
        return "".join(
            f"<p>{_esc(line)}</p>" for line in detail.split("\n") if line.strip()
        )
    return f"<p>{_esc(detail)}</p>" if detail else ""


def _comparison(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if len(rows) < 2:
        return ""
    if len(rows) > 2:
        variant = _variant(slide, "options")
        parts = [f'<div class="nv-compare" data-variant="{_esc(variant)}">']
        for row in rows:
            parts.append(
                f'<div class="nv-card"><h3>{_esc(row["label"])}</h3>'
                f"{_comparison_copy(row)}</div>"
            )
        parts.append("</div>")
        return "".join(parts)
    variant = _variant(slide, "before-after")
    left, right = rows[0], rows[1]
    return (
        f'<div class="nv-compare" data-variant="{_esc(variant)}">'
        f'<div class="nv-card"><h3>{_esc(left["label"])}</h3>'
        f"{_comparison_copy(left)}</div>"
        '<div class="nv-vs">VS</div>'
        f'<div class="nv-card"><h3>{_esc(right["label"])}</h3>'
        f"{_comparison_copy(right)}</div>"
        "</div>"
    )


def _matrix(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)[:4]
    if not rows:
        return ""
    axes = slide.get("axes") if isinstance(slide.get("axes"), dict) else {}
    axis = ""
    if axes:
        axis = (
            f'<div class="nv-axes"><span>{_esc(axes.get("y") or "")}</span>'
            f'<span>{_esc(axes.get("x") or "")}</span></div>'
        )
    parts = [axis, f'<div class="nv-matrix" data-variant="{_esc(_variant(slide, "swot"))}">']
    for row in rows:
        parts.append(
            f'<div class="nv-cell"><h3>{_esc(row["label"])}</h3>'
            f'<p>{_esc(row["detail"] or row["value"])}</p></div>'
        )
    parts.append("</div>")
    return "".join(parts)


def _hierarchy(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    top, rest = rows[0], rows[1:]
    parts = [
        f'<div class="nv-hierarchy" data-variant="{_esc(_variant(slide, "org"))}">',
        '<div class="nv-level">',
        f'<div class="nv-person"><strong>{_esc(top["label"])}</strong>'
        f'<span>{_esc(top["detail"] or top["meta"])}</span></div></div>',
    ]
    if rest:
        parts.append('<div class="nv-level">')
        for row in rest:
            parts.append(
                f'<div class="nv-person"><strong>{_esc(row["label"])}</strong>'
                f'<span>{_esc(row["detail"] or row["meta"])}</span></div>'
            )
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _architecture(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    parts = [f'<div class="nv-arch" data-variant="{_esc(_variant(slide, "system"))}">']
    for index, row in enumerate(rows):
        if index:
            parts.append('<div class="nv-arrow" aria-hidden="true">→</div>')
        parts.append(
            f'<div class="nv-node-box"><strong>{_esc(row["label"])}</strong>'
        )
        if row["detail"]:
            parts.append(f"<span>{_esc(row['detail'])}</span>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _quote(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    quote = _text(slide.get("quote") or slide.get("body") or (rows[0]["detail"] if rows else ""))
    cite = _text(slide.get("cite") or (rows[0]["label"] if rows else "") or (rows[0]["meta"] if rows else ""))
    if not quote:
        return ""
    portrait = _img(slide, "circle", grow=False)
    return (
        f'<div class="nv-quote" data-variant="{_esc(_variant(slide, "customer"))}">'
        f"{portrait}<blockquote>{_esc(quote)}</blockquote>"
        f"{f'<cite>{_esc(cite)}</cite>' if cite else ''}</div>"
    )


def _team(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    variant = _variant(slide, "single" if len(rows) == 1 else "grid")
    parts = [f'<div class="nv-people" data-variant="{_esc(variant)}">']
    for row in rows:
        photo = (
            f'<img class="nv-avatar" alt="{_esc(row["label"])}" src="{_esc(row["image"])}">'
            if row["image"]
            else ""
        )
        parts.append(
            f'<div class="nv-card">{photo}<h3>{_esc(row["label"])}</h3>'
            f'<p>{_esc(row["detail"] or row["meta"])}</p></div>'
        )
    parts.append("</div>")
    return "".join(parts)


def _gallery(slide: Mapping[str, Any]) -> str:
    srcs = [image_src(slide)] if image_src(slide) else []
    for row in items_of(slide):
        if row["image"]:
            srcs.append(row["image"])
    extras = slide.get("images")
    if isinstance(extras, list):
        srcs.extend(_text(item) for item in extras if _text(item).lower() not in _FAKE_IMAGES)
    srcs = [src for src in srcs if src]
    if not srcs:
        return ""
    shown = srcs[:6]
    parts = [f'<div class="nv-gallery" data-n="{len(shown)}">']
    for src in shown:
        parts.append(
            f'<div class="nv-img" data-fit="cover"><img alt="" src="{_esc(src)}"></div>'
        )
    parts.append("</div>")
    return "".join(parts)


def _pricing(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    if not rows:
        return ""
    parts = ['<div class="nv-price">']
    featured = 1 if len(rows) >= 3 else 0
    for index, row in enumerate(rows):
        klass = "nv-card is-featured" if index == featured else "nv-card"
        parts.append(f'<div class="{klass}"><h3>{_esc(row["label"])}</h3>')
        if row["value"]:
            parts.append(f'<div class="nv-num">{_esc(row["value"])}</div>')
        if row["detail"]:
            parts.append(f"<p>{_esc(row['detail'])}</p>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _map(slide: Mapping[str, Any]) -> str:
    rows = items_of(slide)
    parts = [f'<div class="nv-map" data-variant="{_esc(_variant(slide, "world"))}">']
    for index, row in enumerate(rows[:8]):
        left = 18 + (index * 11) % 64
        top = 22 + (index * 17) % 56
        parts.append(
            f'<div class="nv-pin" style="left:{left}%;top:{top}%" title="{_esc(row["label"])}"></div>'
        )
    if rows:
        parts.append('<div class="nv-card" style="position:absolute;left:40px;bottom:40px">')
        for row in rows[:4]:
            parts.append(f"<p>{_esc(row['label'])} {_esc(row['value'])}</p>")
        parts.append("</div>")
    parts.append("</div>")
    return "".join(parts)


def _cover_block(slide: Mapping[str, Any], layout_class: str) -> str:
    sub = _sub(slide) or (
        f'<p class="nv-sub">{_esc(_text(slide.get("body")))}</p>' if _text(slide.get("body")) else ""
    )
    picture = _img(slide, "full-bleed", grow=False)
    has_photo = " nv-has-photo" if picture else ""
    scrim = '<div class="nv-scrim" aria-hidden="true"></div>' if picture else ""
    return (
        f'<div class="nv-slide {layout_class}{has_photo}">'
        f"{picture}{scrim}"
        f"{_kicker(slide)}{_title(slide)}{sub}{_footer(slide)}</div>"
    )


def _panel_figure(slide: Mapping[str, Any]) -> str:
    """A measured figure for the side panel. Never the first word of the title."""
    rows = items_of(slide)
    value = _text(slide.get("value")) or (rows[0]["value"] if rows else "")
    if value:
        return value
    blob = " ".join(
        part
        for part in (
            _text(slide.get("body")),
            *(row["label"] + " " + row["detail"] + " " + row["value"] for row in rows),
        )
        if part
    )
    found = re.search(r"\d[\d\s.,]*\s*(?:%|h|j|x|k|m|md|€|\$)?", blob, re.I)
    if found:
        return found.group(0).strip()
    return ""


def _aside(slide: Mapping[str, Any], side: str) -> str:
    """The visual column: a real picture, or a theme panel when a number exists.

    A list of claims already fills the page. Do not invent a giant word from
    the title. That is how 'Qu'est-ce que Navin' became an empty blue slab.
    """
    picture = _img(slide, side)
    if picture:
        return picture
    rows = items_of(slide)
    value = _panel_figure(slide)
    if rows and len(rows) >= 2 and not value:
        return ""
    label = (
        _text(slide.get("insight"))
        or _text(slide.get("kicker"))
        or (rows[0]["label"] if rows else "")
        or _text(slide.get("body"))
    )
    if label and label == value:
        label = _text(slide.get("kicker")) or (rows[0]["detail"] if rows else "")
    if not value and not label:
        return ""
    if not value and not rows:
        return ""
    number = f'<div class="nv-num">{_esc(value)}</div>' if value else ""
    caption = f"<p>{_esc(label)}</p>" if label else ""
    return f'<div class="nv-panel grow" data-variant="{_esc(side)}">{number}{caption}</div>'


def _text_and_image(slide: Mapping[str, Any], image_side: str) -> str:
    """Text and a visual, side by side. A missing photo becomes a theme panel."""
    picture = _aside(slide, image_side)
    copy = [_body_p(slide), _list(slide, "vertical-block")]
    copy_html = "".join(part for part in copy if part)
    if picture:
        text_col = f'<div class="nv-col grow">{copy_html}</div>'
        row = (
            f"{picture}{text_col}" if image_side == "left" else f"{text_col}{picture}"
        )
        return (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f'<div class="nv-row">{row}</div></div>'
        )
    return (
        f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{copy_html}</div>'
    )


def _product_hero(slide: Mapping[str, Any]) -> str:
    """A real UI capture sits in a browser frame. A theme still does not."""
    payload = dict(slide)
    visual = dict(_visual(payload))
    src = image_src(payload)
    kind = _text(visual.get("kind") or visual.get("type"))
    if src and (kind in {"screenshot", "product", "ui", "phone"} or visual.get("source") != "theme"):
        visual.setdefault("kind", kind or "screenshot")
        visual.setdefault("fit", "contain")
        payload["visual"] = visual
        return _text_and_image(payload, "right")
    visual["kind"] = "photo"
    visual["fit"] = "cover"
    payload["visual"] = visual
    return _text_and_image(payload, "right")


def _split_points(body: str) -> list[dict[str, str]]:
    """Turn a paragraph into card-sized claims when the agent forgot ``items``."""
    text = _text(body)
    if not text:
        return []
    parts = [part.strip(" -•\t") for part in re.split(r"[\n•]+", text) if part.strip()]
    if len(parts) < 2:
        parts = [part.strip() for part in re.split(r"(?<=[.!?])\s+", text) if part.strip()]
    out: list[dict[str, str]] = []
    for part in parts[:6]:
        if len(part) < 4:
            continue
        out.append({"label": part, "detail": "", "value": "", "meta": "", "image": ""})
    return out


def enrich_slide(slide: Mapping[str, Any]) -> dict[str, Any]:
    """Stop a content slide collapsing to a title and one lonely sentence.

    A cover or a statement is allowed to hold one idea. A ``text`` slide that
    does the same looks unfinished: split a long body into cards, or promote
    a thin one to a 50/50 panel so the page still has a composition.
    """
    payload = normalize_slide(slide)
    layout = _text(payload.get("layout")) or "text"
    items = items_of(payload)
    body = _text(payload.get("body"))
    has_visual = bool(
        image_src(payload)
        or payload.get("series")
        or payload.get("table")
        or (isinstance(payload.get("chart"), dict) and payload["chart"].get("series"))
    )

    if layout == "text" and not items and not has_visual:
        points = _split_points(body)
        if len(points) >= 2:
            payload["items"] = points
            payload["layout"] = "cards"
            return payload
        payload["layout"] = "text-image"
        return payload

    if layout in {"text-image", "image-text"} and not items and not has_visual:
        points = _split_points(body)
        if len(points) >= 2:
            payload["items"] = points
        return payload

    if layout in {"cards", "feature-grid"} and len(items) < 2:
        points = items + _split_points(body)
        if len(points) >= 2:
            payload["items"] = points
        return payload

    return payload


def render_body(slide: Mapping[str, Any]) -> str:
    """HTML for the ``<body>`` of one semantic slide. No sample leftovers."""
    slide = enrich_slide(normalize_slide(slide))
    layout = _text(slide.get("layout")) or "text"
    builders = {
        "cover": lambda: _cover_block(slide, "nv-layout-cover"),
        "hero": lambda: _cover_block(slide, "nv-layout-cover"),
        "statement": lambda: _cover_block(slide, "nv-layout-statement"),
        "section-break": lambda: _cover_block(slide, "nv-layout-section-break"),
        "closing": lambda: _cover_block(slide, "nv-layout-cover"),
        "text": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f"{_body_p(slide)}{_list(slide, 'vertical-block')}</div>"
        ),
        "text-image": lambda: _text_and_image(slide, "right"),
        "image-text": lambda: _text_and_image(slide, "left"),
        "full-image": lambda: (
            f'<div class="nv-slide nv-layout-full-image">'
            f'{_img(slide, "full-bleed", grow=False) or ""}'
            f'<div class="nv-content">{_kicker(slide)}{_title(slide)}'
            f"{_sub(slide) or _body_p(slide, 'nv-sub')}</div></div>"
        ),
        "full-bleed-hero": lambda: (
            f'<div class="nv-slide nv-layout-full-bleed-hero">'
            f'{_img(slide, "full-bleed", grow=False)}'
            f'<div class="nv-content">{_kicker(slide)}{_title(slide)}'
            f"{_sub(slide) or _body_p(slide, 'nv-sub')}</div></div>"
        ),
        "product-hero": lambda: _product_hero(slide),
        "gallery": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_gallery(slide)}</div>'
        ),
        "cards": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f"{_list(slide, _variant(slide, 'card'))}</div>"
        ),
        "feature-grid": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f"{_list(slide, 'grid')}</div>"
        ),
        "agenda": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f"{_list(slide, 'numbered')}</div>"
        ),
        "big-numbers": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_kpi(slide)}</div>'
        ),
        "kpi": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_kpi(slide)}</div>'
        ),
        "dashboard": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f"{_insight(slide)}{_kpi(slide)}</div>"
        ),
        "chart": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f"{_insight(slide)}{_chart(slide)}</div>"
        ),
        "data-story": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}'
            f"{_insight(slide)}{_chart(slide)}{_body_p(slide)}</div>"
        ),
        "table": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_table(slide)}</div>'
        ),
        "comparison": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_comparison(slide)}</div>'
        ),
        "pricing": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_pricing(slide)}</div>'
        ),
        "timeline": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_timeline(slide)}</div>'
        ),
        "roadmap": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_roadmap(slide)}</div>'
        ),
        "process": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_process(slide)}</div>'
        ),
        "cycle": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_cycle(slide)}</div>'
        ),
        "funnel": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_funnel(slide)}</div>'
        ),
        "hierarchy": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_hierarchy(slide)}</div>'
        ),
        "architecture": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_architecture(slide)}</div>'
        ),
        "matrix": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_matrix(slide)}</div>'
        ),
        "map": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_map(slide)}</div>'
        ),
        "team": lambda: (
            f'<div class="nv-slide">{_kicker(slide)}{_title(slide)}{_team(slide)}</div>'
        ),
        "quote": lambda: f'<div class="nv-slide">{_quote(slide)}</div>',
        "case-study": lambda: _text_and_image(slide, "right"),
    }
    build = builders.get(layout)
    if build is None:
        build = builders["text"]
    return _stamp_slide(build(), slide)


_OUTLINE_LAYOUTS = {"agenda"}
_OUTLINE_SKIP = {
    "cover",
    "hero",
    "closing",
    "agenda",
    "quote",
    "full-bleed-hero",
}
_OUTLINE_TITLE_RE = re.compile(
    r"^(sommaire|agenda|contents|table of contents|table des matieres|"
    r"au programme|outline)$",
    re.I,
)


def outline_copy(language: str, *hints: str) -> dict[str, str]:
    """Labels for the agenda slide. Content titles stay as the agent wrote them."""
    lang = (language or "").lower()
    blob = " ".join(hints).lower()
    french = lang.startswith("fr") or bool(
        re.search(r"\b(les|des|une|pour|avec|dans|cette|equipe|risques)\b", blob)
    )
    if french:
        return {"kicker": "Sommaire", "title": "Au programme"}
    if lang.startswith("ar"):
        return {"kicker": "الفهرس", "title": "محاور العرض"}
    return {"kicker": "Agenda", "title": "What we will cover"}


def stagecraft_copy(language: str, *hints: str) -> dict[str, str]:
    """Auto-inserted gallery and curtain copy follows the deck language."""
    lang = (language or "").lower()
    blob = " ".join(hints).lower()
    french = lang.startswith("fr") or bool(
        re.search(r"\b(les|des|une|pour|avec|dans|cette|equipe|risques)\b", blob)
    )
    if french:
        return {
            "gallery_kicker": "Preuve",
            "gallery_title": "Le travail, pas un collage",
            "curtain_title": "Ce qui change maintenant",
        }
    if lang.startswith("ar"):
        return {
            "gallery_kicker": "دليل",
            "gallery_title": "العمل الظاهر",
            "curtain_title": "ما الذي يتغير الان",
        }
    return {
        "gallery_kicker": "Evidence",
        "gallery_title": "The work, not a collage",
        "curtain_title": "What changes now",
    }


def _is_outline_slide(slide: Mapping[str, Any]) -> bool:
    layout = _text(slide.get("layout")).lower()
    if layout in _OUTLINE_LAYOUTS:
        return True
    title = _text(slide.get("title"))
    kicker = _text(slide.get("kicker"))
    return bool(_OUTLINE_TITLE_RE.match(title) or _OUTLINE_TITLE_RE.match(kicker))


def _outline_entries(slides: list[Mapping[str, Any]]) -> list[dict[str, str]]:
    chapters = [
        slide
        for slide in slides
        if _text(slide.get("layout")) == "section-break" and _text(slide.get("title"))
    ]
    source = chapters if len(chapters) >= 2 else [
        slide
        for slide in slides
        if _text(slide.get("layout")) not in _OUTLINE_SKIP and _text(slide.get("title"))
        and not _is_outline_slide(slide)
    ]
    seen: set[str] = set()
    out: list[dict[str, str]] = []
    for slide in source:
        title = _text(slide.get("title"))
        key = title.lower()
        if key in seen:
            continue
        seen.add(key)
        detail = _text(slide.get("kicker") or slide.get("sub") or slide.get("body"))
        if len(detail) > 90:
            detail = detail[:87].rsplit(" ", 1)[0] + "."
        out.append({"label": title, "detail": detail, "value": "", "meta": "", "image": ""})
        if len(out) >= 10:
            break
    return out


def ensure_outline(
    slides: list[Mapping[str, Any]],
    *,
    language: str = "",
    outline_title: str = "",
) -> list[dict[str, Any]]:
    """Put a sommaire after the cover. A deck without one has no map.

    Cover, statement and closing may stand alone. From three content slides
    up, the room needs to see the path. If the agent already wrote an agenda,
    it is kept and filled when it has no items.
    """
    deck = [dict(slide) for slide in slides]
    entries = _outline_entries(deck)
    if len(entries) < 3:
        return deck
    copy = outline_copy(
        language,
        outline_title,
        *( _text(slide.get("title")) for slide in deck[:4] ),
    )
    if outline_title:
        copy["title"] = outline_title
    existing = next((index for index, slide in enumerate(deck) if _is_outline_slide(slide)), -1)
    if existing >= 0:
        current = deck[existing]
        if not items_of(current):
            current["items"] = entries
            current.setdefault("layout", "agenda")
            current.setdefault("kicker", copy["kicker"])
            current.setdefault("title", copy["title"])
        return deck
    insert_at = 0
    if deck and _text(deck[0].get("layout")) in {"cover", "hero", "full-bleed-hero"}:
        insert_at = 1
    deck.insert(
        insert_at,
        {
            "layout": "agenda",
            "kicker": copy["kicker"],
            "title": copy["title"],
            "items": entries,
        },
    )
    return deck


_TONE_ALIASES = {
    "light": "light",
    "dark": "dark",
    "hero": "hero-light",
    "hero-light": "hero-light",
    "herolight": "hero-light",
    "hero-dark": "hero-dark",
    "herodark": "hero-dark",
}
_TONE_LIGHT = {"light", "hero-light"}
_TONE_DARK = {"dark", "hero-dark"}
_STAGE_SKIP = {
    "cover",
    "hero",
    "full-bleed-hero",
    "agenda",
    "closing",
}
_DATA_HERO = {"big-numbers", "kpi", "dashboard", "data-story"}
STYLE_THEMES = {
    "magazine": "editorial_luxe",
    "swiss": "black_and_white_clean",
    "corporate": "textbook",
}


def normalize_tone(value: Any) -> str:
    raw = _text(value).lower().replace("_", "-").replace(" ", "-")
    return _TONE_ALIASES.get(raw, "")


def _tone_family(tone: str) -> str:
    if tone in _TONE_DARK:
        return "dark"
    if tone in _TONE_LIGHT:
        return "light"
    return ""


def _stamp_slide(body: str, slide: Mapping[str, Any]) -> str:
    """Mark filled slides with tone and speaker notes the converter can keep."""
    tone = normalize_tone(slide.get("tone"))
    notes = _text(slide.get("notes"))

    def repl(match: re.Match[str]) -> str:
        classes = match.group(1)
        extra = ""
        if tone and f"nv-tone-{tone}" not in classes:
            extra += f" nv-tone-{tone}"
        attrs = [f'class="nv-slide{classes}{extra}"', 'data-filled="1"']
        if tone:
            attrs.append(f'data-tone="{_esc(tone)}"')
        if notes:
            attrs.append(f'data-notes="{_esc(notes)}"')
        return f"<div {' '.join(attrs)}>"

    stamped = re.sub(r'<div class="nv-slide([^"]*)">', repl, body, count=1)
    if "data-filled=" not in stamped:
        stamped = re.sub(
            r'class="nv-slide([^"]*)"',
            r'class="nv-slide\1" data-filled="1"',
            stamped,
            count=1,
        )
    return stamped


def ensure_rhythm(slides: list[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Alternate light / dark / hero so the deck breathes.

    Three pages of the same family tire the room. Cover and curtains take a
    hero tone. Content pages flip every two slides. Eight pages and up must
    show both a dark hero and a light hero.
    """
    deck = [dict(slide) for slide in slides]
    for index, slide in enumerate(deck):
        if normalize_tone(slide.get("tone")):
            slide["tone"] = normalize_tone(slide.get("tone"))
            continue
        layout = _text(slide.get("layout"))
        if layout in {"cover", "hero", "full-bleed-hero"}:
            slide["tone"] = "hero-dark" if index == 0 else "hero-light"
        elif layout in {"statement", "section-break", "quote", "closing"}:
            slide["tone"] = "hero-dark" if index % 2 == 0 else "hero-light"
        else:
            slide["tone"] = "dark" if (index // 2) % 2 else "light"

    for index in range(2, len(deck)):
        family = _tone_family(deck[index]["tone"])
        if (
            family
            and family == _tone_family(deck[index - 1]["tone"])
            and family == _tone_family(deck[index - 2]["tone"])
        ):
            deck[index]["tone"] = "hero-dark" if family == "light" else "hero-light"

    if len(deck) >= 8:
        tones = {normalize_tone(slide.get("tone")) for slide in deck}
        if "hero-dark" not in tones:
            deck[0]["tone"] = "hero-dark"
        if "hero-light" not in tones:
            for slide in deck:
                if _text(slide.get("layout")) in {"section-break", "statement", "closing"}:
                    slide["tone"] = "hero-light"
                    break
            else:
                deck[min(3, len(deck) - 1)]["tone"] = "hero-light"
    return deck


def _content_count(slides: list[Mapping[str, Any]]) -> int:
    return sum(1 for slide in slides if _text(slide.get("layout")) not in _STAGE_SKIP)


def _collect_images(slides: list[Mapping[str, Any]]) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for slide in slides:
        candidates = [image_src(slide)]
        extras = slide.get("images")
        if isinstance(extras, list):
            candidates.extend(_text(item) for item in extras)
        for row in items_of(slide):
            if row.get("image"):
                candidates.append(row["image"])
        for src in candidates:
            key = src.lower()
            if not src or key in _FAKE_IMAGES or key in seen:
                continue
            seen.add(key)
            found.append(src)
    return found


def ensure_stagecraft(
    slides: list[Mapping[str, Any]],
    *,
    language: str = "",
) -> list[dict[str, Any]]:
    """From eight content slides up, force a curtain, a data hero, a photo grid.

    A deck that is only cards looks like a document. One chapter break, one
    giant number, and a gallery when photos exist are the minimum stagecraft.
    """
    deck = [dict(slide) for slide in slides]
    if _content_count(deck) < 8:
        return deck
    layouts = {_text(slide.get("layout")) for slide in deck}
    copy = stagecraft_copy(
        language,
        *(_text(slide.get("title")) for slide in deck[:6]),
    )

    if "section-break" not in layouts:
        insert_at = 1
        if deck and _text(deck[0].get("layout")) in {"cover", "hero", "full-bleed-hero"}:
            insert_at = 2 if len(deck) > 1 and _is_outline_slide(deck[1]) else 1
        insert_at = min(insert_at + max(1, (len(deck) - insert_at) // 3), len(deck))
        nxt = deck[insert_at] if insert_at < len(deck) else {}
        title = _text(nxt.get("title")) or copy["curtain_title"]
        if len(title) > 42:
            title = title[:39].rsplit(" ", 1)[0]
        deck.insert(
            insert_at,
            {
                "layout": "section-break",
                "kicker": "02",
                "title": title,
                "tone": "hero-dark",
            },
        )
        layouts.add("section-break")

    if not layouts & _DATA_HERO:
        promoted = False
        for slide in deck:
            if _text(slide.get("layout")) in _STAGE_SKIP | {"section-break", "gallery"}:
                continue
            rows = items_of(slide)
            if any(row.get("value") for row in rows):
                slide["layout"] = "big-numbers"
                promoted = True
                break
        if not promoted:
            for slide in deck:
                if _text(slide.get("layout")) in {"cards", "text", "feature-grid"} and items_of(slide):
                    slide["layout"] = "big-numbers"
                    break

    if "gallery" not in layouts:
        images = _collect_images(deck)
        if len(images) >= 2:
            insert_at = 1
            if deck and _text(deck[0].get("layout")) in {"cover", "hero", "full-bleed-hero"}:
                insert_at = 2 if len(deck) > 1 and _is_outline_slide(deck[1]) else 1
            for index, slide in enumerate(deck):
                if image_src(slide):
                    insert_at = min(index + 1, len(deck))
                    break
            deck.insert(
                insert_at,
                {
                    "layout": "gallery",
                    "kicker": copy["gallery_kicker"],
                    "title": copy["gallery_title"],
                    "images": images[:6],
                    "tone": "light",
                },
            )
    return deck


def stagecraft_gaps(slides: list[Mapping[str, Any]]) -> list[str]:
    """Deck-level holes the room notices on an 8+ page talk."""
    if _content_count(slides) < 8:
        return []
    layouts = {_text(slide.get("layout")) for slide in slides}
    gaps: list[str] = []
    if "section-break" not in layouts:
        gaps.append("missing-section-break")
    if not layouts & _DATA_HERO:
        gaps.append("missing-data-hero")
    if "gallery" not in layouts and _collect_images(slides):
        gaps.append("missing-gallery")
    return gaps


_TOPIC_TITLE = re.compile(
    r"^(overview|introduction|agenda|sommaire|conclusion|next steps|"
    r"prochaines etapes|equipe|team|marche|market|produit|product|"
    r"solution|probleme|problem|resultats|results|contexte|context|"
    r"about|a propos|merci|thank you|the number|"
    r"qu.est-ce que\b.*|what is\b.*|"
    r"sommaire de\b.*|les \d+ piliers\b.*|"
    r"architecture\b.*|capacit[eé]s\b.*|cycle d.\b.*|"
    r"feuille de route\b.*|impact et\b.*|pr[eé]sentation\b.*)\s*$",
    re.I,
)

_POSTER_LAYOUTS = {
    "cover",
    "hero",
    "statement",
    "section-break",
    "closing",
    "quote",
    "full-image",
    "full-bleed-hero",
    "agenda",
}


def _slide_is_hollow(slide: Mapping[str, Any]) -> bool:
    payload = normalize_slide(slide)
    layout = _text(payload.get("layout"))
    if layout in _POSTER_LAYOUTS:
        return False
    if items_of(payload) or image_src(payload) or payload.get("series") or payload.get("table"):
        return False
    if isinstance(payload.get("chart"), dict) and payload["chart"].get("series"):
        return False
    extras = payload.get("images")
    if isinstance(extras, list) and any(
        _text(item) and _text(item).lower() not in _FAKE_IMAGES for item in extras
    ):
        return False
    if _text(payload.get("body") or payload.get("quote")):
        return False
    return True


def _is_assertion(title: str) -> bool:
    text = _text(title)
    if not text or _TOPIC_TITLE.match(text):
        return False
    if re.search(r"\d", text):
        return True
    return len(text.split()) >= 4


def critique_deck(slides: list[Mapping[str, Any]]) -> dict[str, Any]:
    """Five scores the room uses. Under 85 is not a deliverable."""
    deck = [normalize_slide(slide) for slide in slides]
    count = len(deck)
    layouts = [_text(slide.get("layout")) for slide in deck]
    titles = [_text(slide.get("title")) for slide in deck]
    content = _content_count(deck)
    notes: list[str] = []

    story = 0
    if "cover" in layouts:
        story += 6
    else:
        notes.append("No cover. Open with a hook, not a topic list.")
    if content < 3 or "agenda" in layouts:
        story += 4
    else:
        notes.append("No sommaire after the cover.")
    if titles:
        story += min(10, round(10 * sum(1 for title in titles if _is_assertion(title)) / len(titles)))
    if titles and sum(1 for title in titles if _is_assertion(title)) < max(1, (len(titles) + 1) // 2):
        notes.append("Most titles are topic labels. The title is the message.")

    typography = 0
    if layouts and layouts[0] == "cover":
        typography += 8
    if {"cover", "statement", "section-break"} & set(layouts):
        typography += 6
    short = sum(1 for title in titles if 0 < len(title.split()) <= 12)
    if titles and short == len(titles):
        typography += 6
    elif short:
        typography += 3
    else:
        notes.append("Titles run long. Poster pages need 8 to 12 words.")

    rhythm = 0
    families = [
        _tone_family(normalize_tone(slide.get("tone")) or "light") or "light"
        for slide in deck
    ]
    if count < 3 or len(set(families)) >= 2:
        rhythm += 10
    else:
        notes.append("The deck stays in one light or dark family.")
    flat = any(
        families[index] == families[index + 1] == families[index + 2]
        for index in range(max(0, len(families) - 2))
    )
    if not flat:
        rhythm += 10
    else:
        notes.append("Three slides in a row share the same family.")

    evidence = 0
    if set(layouts) & _DATA_HERO:
        evidence += 8
    else:
        notes.append("No data hero. Give the room one giant number.")
    measured = any(
        _text(slide.get("layout")) in {"chart", "dashboard", "data-story"}
        or slide.get("series")
        or (isinstance(slide.get("chart"), dict) and slide["chart"].get("series"))
        or any(row.get("value") for row in items_of(slide))
        for slide in deck
    )
    if measured or content < 4:
        evidence += 6
    else:
        notes.append("No chart and no measured items.")
    if _collect_images(deck) or "gallery" in layouts:
        evidence += 6
    elif content < 8:
        evidence += 4
    else:
        notes.append("No photograph. A talk without evidence looks invented.")

    stage = 0
    if content < 8 or "section-break" in layouts:
        stage += 8
    else:
        notes.append("No chapter curtain.")
    noted = sum(1 for slide in deck if _text(slide.get("notes")))
    if count and noted >= (count + 1) // 2:
        stage += 6
    elif noted:
        stage += 3
        notes.append("Speaker notes cover fewer than half the slides.")
    else:
        notes.append("No speaker notes.")
    photos = _collect_images(deck)
    if content < 8 or "gallery" in layouts or len(photos) < 2:
        stage += 6
    else:
        notes.append("Photos exist but no gallery.")

    dimensions = {
        "story": story,
        "type": typography,
        "rhythm": rhythm,
        "evidence": evidence,
        "stage": stage,
    }
    score = sum(dimensions.values())
    hollow = [
        index
        for index, slide in enumerate(deck, 1)
        if _slide_is_hollow(slide)
    ]
    if hollow:
        notes.append(
            "Hollow slides: "
            + ", ".join(str(index) for index in hollow)
            + ". Fill the chosen template: items, a photo, or a chart."
        )
    topic_heavy = bool(
        titles
        and sum(1 for title in titles if _is_assertion(title))
        < max(1, (len(titles) + 1) // 2)
    )
    notes_thin = bool(content >= 8 and noted < (count + 1) // 2)
    passed = (
        score >= 85
        and not hollow
        and not topic_heavy
        and not notes_thin
    )
    return {
        "score": score,
        "threshold": 85,
        "pass": passed,
        "dimensions": dimensions,
        "notes": notes,
    }


_PREVIEW_HTML = """<!DOCTYPE html>
<html lang="und">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>%%TITLE%%</title>
<style>
  :root { color-scheme: dark; }
  html, body { margin: 0; height: 100%; background: #111; color: #f4f4f4; overflow: hidden; font-family: Inter, Segoe UI, sans-serif; }
  body[data-mode="audience"] .desk, body[data-mode="overview"] .desk, body[data-mode="overview"] .stage { display: none; }
  body[data-mode="presenter"] .stage, body[data-mode="overview"] .stage { display: none; }
  body[data-mode="presenter"] .desk { display: grid; }
  body[data-mode="overview"] .grid { display: grid; }
  body[data-chrome="off"] .bar { display: none; }
  .stage { width: 100vw; height: 100vh; display: flex; align-items: center; justify-content: center; }
  .desk { display: none; height: 100vh; grid-template-columns: minmax(0, 1.6fr) minmax(320px, 0.9fr); gap: 16px; padding: 16px 16px 56px; }
  .pane { background: #1a1a1a; border-radius: 12px; overflow: hidden; display: flex; align-items: center; justify-content: center; position: relative; }
  .side { display: flex; flex-direction: column; gap: 12px; min-width: 0; }
  .next-wrap { flex: 0 0 38%; }
  .notes { flex: 1; background: #1a1a1a; border-radius: 12px; padding: 16px 18px; overflow: auto; font-size: 18px; line-height: 1.45; white-space: pre-wrap; }
  .notes h2, .next-wrap h2 { margin: 0 0 8px; font-size: 11px; letter-spacing: 0.14em; text-transform: uppercase; color: #999; }
  iframe.slide { width: 1920px; height: 1080px; border: 0; background: #000; transform-origin: center center; }
  .bar { position: fixed; left: 0; right: 0; bottom: 0; display: flex; align-items: center; gap: 12px; padding: 10px 20px; background: rgba(0,0,0,.62); font-size: 13px; z-index: 2; }
  .bar button { background: transparent; color: inherit; border: 1px solid #666; padding: 4px 10px; cursor: pointer; }
  .timer { margin-left: auto; font-variant-numeric: tabular-nums; }
  .grid { display: none; position: fixed; inset: 0 0 48px; padding: 20px; grid-template-columns: repeat(auto-fill, minmax(220px, 1fr)); gap: 12px; overflow: auto; }
  .grid button { background: #1a1a1a; color: inherit; border: 1px solid #333; border-radius: 10px; padding: 18px 14px; text-align: left; cursor: pointer; }
  .grid button[aria-current="true"] { border-color: #f4f4f4; }
</style>
</head>
<body data-mode="audience" data-chrome="on">
<div class="stage"><iframe class="slide" id="frame" title="%%TITLE%%" src="%%FIRST%%"></iframe></div>
<div class="desk">
  <div class="pane"><iframe class="slide" id="now" title="Current" src="%%FIRST%%"></iframe></div>
  <div class="side">
    <div class="pane next-wrap">
      <h2>Next</h2>
      <iframe class="slide" id="peek" title="Next" src="%%FIRST%%"></iframe>
    </div>
    <div class="notes"><h2>Notes</h2><div id="script"></div></div>
  </div>
</div>
<div class="grid" id="overview"></div>
<div class="bar">
  <button type="button" id="prev" aria-label="Previous">Prev</button>
  <span id="pos">1 / %%COUNT%%</span>
  <button type="button" id="next" aria-label="Next">Next</button>
  <button type="button" id="mode" aria-label="Presenter">Presenter (S)</button>
  <span>%%TITLE%%</span>
  <span class="timer" id="timer">00:00</span>
</div>
<script>
const slides = %%FILES%%;
const notes = %%NOTES%%;
let index = 0;
let started = Date.now();
const frame = document.getElementById("frame");
const now = document.getElementById("now");
const peek = document.getElementById("peek");
const pos = document.getElementById("pos");
const script = document.getElementById("script");
const overview = document.getElementById("overview");
function scaleFrame(node, width, height) {
  if (!node || !node.parentElement) return;
  const box = node.parentElement.getBoundingClientRect();
  const scale = Math.min(box.width / 1920, box.height / 1080);
  node.style.transform = "scale(" + scale + ")";
}
function fit() {
  const chrome = document.body.getAttribute("data-chrome") === "off" ? 0 : 48;
  if (document.body.getAttribute("data-mode") === "audience") {
    const scale = Math.min(window.innerWidth / 1920, (window.innerHeight - chrome) / 1080);
    frame.style.transform = "scale(" + scale + ")";
  }
  scaleFrame(now, 1920, 1080);
  scaleFrame(peek, 1920, 1080);
}
function show(next) {
  index = (next + slides.length) % slides.length;
  const here = slides[index];
  const after = slides[(index + 1) % slides.length];
  frame.src = here;
  now.src = here;
  peek.src = after;
  pos.textContent = (index + 1) + " / " + slides.length;
  script.textContent = notes[index] || "No speaker notes on this slide.";
  Array.from(overview.children).forEach((button, offset) => {
    button.setAttribute("aria-current", offset === index ? "true" : "false");
  });
  fit();
}
function setMode(mode) {
  document.body.setAttribute("data-mode", mode);
  document.getElementById("mode").textContent = mode === "presenter" ? "Audience (S)" : "Presenter (S)";
  fit();
}
slides.forEach((name, offset) => {
  const button = document.createElement("button");
  button.type = "button";
  button.textContent = (offset + 1) + ". " + name;
  button.onclick = () => { show(offset); setMode("audience"); };
  overview.appendChild(button);
});
document.getElementById("prev").onclick = () => show(index - 1);
document.getElementById("next").onclick = () => show(index + 1);
document.getElementById("mode").onclick = () => {
  setMode(document.body.getAttribute("data-mode") === "presenter" ? "audience" : "presenter");
};
window.addEventListener("keydown", (event) => {
  if (event.key === "ArrowRight" || event.key === " " || event.key === "PageDown") show(index + 1);
  if (event.key === "ArrowLeft" || event.key === "PageUp") show(index - 1);
  if (event.key === "Home") show(0);
  if (event.key === "End") show(slides.length - 1);
  if (event.key === "s" || event.key === "S") {
    setMode(document.body.getAttribute("data-mode") === "presenter" ? "audience" : "presenter");
  }
  if (event.key === "o" || event.key === "O") {
    setMode(document.body.getAttribute("data-mode") === "overview" ? "audience" : "overview");
  }
  if (event.key === "f" || event.key === "F") {
    const on = document.body.getAttribute("data-chrome") !== "off";
    document.body.setAttribute("data-chrome", on ? "off" : "on");
    fit();
  }
});
setInterval(() => {
  const elapsed = Math.floor((Date.now() - started) / 1000);
  const minutes = String(Math.floor(elapsed / 60)).padStart(2, "0");
  const seconds = String(elapsed % 60).padStart(2, "0");
  document.getElementById("timer").textContent = minutes + ":" + seconds;
}, 1000);
window.addEventListener("resize", fit);
show(0);
</script>
</body>
</html>
"""


def write_swipe_preview(
    out_dir: Path, *, title: str = "", notes: list[str] | None = None
) -> Path:
    """Audience swipe, presenter desk (S), overview (O), timer."""
    slides = sorted(out_dir.glob("slide_*.html"))
    if not slides:
        raise FileNotFoundError(f"{out_dir}: no slide_*.html to preview")
    names = [path.name for path in slides]
    heading = _esc(title or "Deck preview")
    spoken = list(notes if notes is not None else notes_from_html_dir(out_dir))
    while len(spoken) < len(names):
        spoken.append("")
    html = (
        _PREVIEW_HTML.replace("%%TITLE%%", heading)
        .replace("%%FIRST%%", _esc(names[0]))
        .replace("%%COUNT%%", str(len(names)))
        .replace("%%FILES%%", json.dumps(names, ensure_ascii=True))
        .replace("%%NOTES%%", json.dumps(spoken[: len(names)], ensure_ascii=True))
    )
    path = out_dir / "index.html"
    path.write_text(html, encoding="utf-8")
    return path


def notes_from_html_dir(out_dir: Path) -> list[str]:
    """Speaker notes stamped on each slide_XX.html, in file order."""
    found: list[str] = []
    for path in sorted(out_dir.glob("slide_*.html")):
        html_text = path.read_text(encoding="utf-8")
        match = re.search(r'data-notes="([^"]*)"', html_text)
        if not match:
            found.append("")
            continue
        found.append(html.unescape(match.group(1)))
    return found


def apply_speaker_notes(pptx_path: Path, notes: list[str]) -> int:
    """Write speaker notes into an editable PPTX. Returns how many slides got notes."""
    from pptx import Presentation

    presentation = Presentation(str(pptx_path))
    written = 0
    for slide, text in zip(presentation.slides, notes):
        if not str(text or "").strip():
            continue
        slide.notes_slide.notes_text_frame.text = str(text).strip()
        written += 1
    if written:
        presentation.save(str(pptx_path))
    return written


def leftover_sample_phrases(html_text: str) -> list[str]:
    lowered = html_text.lower()
    return [phrase for phrase in SAMPLE_PHRASES if phrase in lowered]


def has_placeholder_image(html_text: str) -> bool:
    return any(f'src="{name}"' in html_text.lower() for name in _FAKE_IMAGES if name)


def load_slide_json(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("slide JSON must be an object")
    return normalize_slide(data)


def load_deck_json(path: Path) -> tuple[str, list[dict[str, Any]], dict[str, str]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    meta = {"language": "", "outline_title": ""}
    if isinstance(data, list):
        slides = data
        theme = ""
    elif isinstance(data, dict):
        theme = _text(data.get("theme"))
        slides = data.get("slides") or []
        meta["language"] = _text(data.get("language"))
        meta["outline_title"] = _text(data.get("outline_title") or data.get("agenda_title"))
    else:
        raise ValueError("deck JSON must be an object or an array")
    if not isinstance(slides, list) or not slides:
        raise ValueError("deck JSON needs a non-empty slides array")
    return theme, [normalize_slide(item) for item in slides if isinstance(item, dict)], meta
