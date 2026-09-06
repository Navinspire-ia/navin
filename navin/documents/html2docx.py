"""Convert an HTML document master into a native, editable DOCX file.

The template library ships A4 documents as HTML (``document.html``, one
``.page`` section per page). Word is a flow format, not a canvas, so this
converter does not reproduce the geometry the way the deck converter does: it
reads the laid-out page in Chromium and rebuilds it as Word constructs -
headings, paragraphs, bullet and numbered lists, tables, pictures - carrying
over the typography, colors, shading and borders it measured.

What comes out is a document somebody can actually work in: styles in the
navigation pane, selectable text, tables with real cells, and no screenshot
anywhere.

Usage::

    python3 -m navin.documents.html2docx document.html -o report.docx
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Sequence

if __package__:
    from . import _dom, _word
    from ._assets import resolve_asset
    from ._chromium import ConversionError, find_chromium, measure
    from ._fonts import office_font
else:  # copied as a standalone folder into a user workspace
    import _dom
    import _word
    from _assets import resolve_asset
    from _chromium import ConversionError, find_chromium, measure
    from _fonts import office_font

# A4 at 96 dpi, the viewport the masters are designed against.
PAGE_WIDTH_PX = 794
PAGE_HEIGHT_PX = 1123
PT_PER_PX = 0.75

_MEASURE_BODY = r"""
  const out = { pages: [], scale: 1 };
  let uid = 0;

  const roots = pageRoots([".page", ".sheet", "body > section", "body > article", "body > main"]);

  const px = (v) => round(parseFloat(v) || 0);

  // Authoring directives. A master asks for Word behaviour it cannot express in
  // CSS - a refreshable table of contents, a landscape section, a heading that
  // must not be left alone at the foot of a page - by tagging the element.
  // ENTER applies once, to the next block emitted after the tagged container;
  // INHERIT flows down to every block inside it; SELF stays on its own element.
  const ENTER_KEYS = ["toc", "section", "columns", "break", "numbering", "cover"];
  const INHERIT_KEYS = ["keep", "role"];
  const SELF_KEYS = ["caption", "bookmark", "ref"];
  // Read where the container is met rather than carried onto a block:
  // data-doc-row="stack" tells a side-by-side row to come through as its rows of
  // content instead of a table, which is how a header band keeps a real heading.
  const STRUCTURE_KEYS = ["row"];
  let pending = null;

  const directivesOf = (el) => {
    if (!el || typeof el.getAttribute !== "function") return {};
    const found = {};
    for (const key of [...ENTER_KEYS, ...INHERIT_KEYS, ...SELF_KEYS, ...STRUCTURE_KEYS]) {
      const value = el.getAttribute("data-doc-" + key);
      if (value !== null && String(value).trim()) found[key] = String(value).trim();
    }
    return found;
  };

  const pick = (source, keys) => {
    const out = {};
    for (const key of keys) if (source[key]) out[key] = source[key];
    return out;
  };

  const emit = (blocks, block, self, inherited) => {
    const doc = { ...(inherited || {}) };
    if (pending) {
      Object.assign(doc, pending);
      pending = null;
    }
    Object.assign(doc, self || {});
    if (Object.keys(doc).length) block.doc = doc;
    blocks.push(block);
  };

  const tocLevels = (value) => {
    if (!value || value === "true" || value === "1" || value === "yes") return "1-3";
    return value;
  };

  const sideBorders = (cs) => {
    const sides = {};
    for (const side of ["Top", "Right", "Bottom", "Left"]) {
      const width = parseFloat(cs[`border${side}Width`]) || 0;
      const paint = color(cs[`border${side}Color`]);
      const style = cs[`border${side}Style`];
      if (width > 0.4 && paint && style !== "none" && style !== "hidden") {
        sides[side.toLowerCase()] = { width: round(width), color: paint.hex, style };
      }
    }
    return Object.keys(sides).length ? sides : null;
  };

  const boxOf = (el, cs) => {
    const fill = color(cs.backgroundColor);
    const borders = sideBorders(cs);
    if (!fill && !borders) return null;
    return { id: `b${uid++}`, fill: fill && fill.alpha > 0.15 ? fill.hex : null, borders };
  };

  const paragraphStyle = (el, cs, rect, frame) => ({
    align: cs.textAlign,
    rtl: cs.direction === "rtl",
    indentLeft: round(Math.max(0, rect.left - frame.left)),
    indentRight: round(Math.max(0, frame.right - rect.right)),
    lineHeight: Number.isFinite(parseFloat(cs.lineHeight)) ? px(cs.lineHeight) : null,
    top: round(rect.top),
    bottom: round(rect.bottom),
  });

  const listItems = (list, frame) => {
    const items = [];
    for (const li of Array.from(list.children)) {
      if (li.tagName !== "LI") continue;
      const cs = getComputedStyle(li);
      const rect = li.getBoundingClientRect();
      if (!visible(cs, rect)) continue;
      const nested = li.querySelector("ul, ol");
      const runs = runsOf(nested ? stripped(li, nested) : li);
      if (runs.length) {
        items.push({ runs, level: 0, style: paragraphStyle(li, cs, rect, frame) });
      }
      if (nested) {
        for (const child of listItems(nested, frame)) {
          items.push({ ...child, level: child.level + 1 });
        }
      }
    }
    return items;
  };

  // A list item that carries a sub-list must contribute its own text only.
  const stripped = (li, nested) => {
    const clone = li.cloneNode(true);
    for (const sub of Array.from(clone.querySelectorAll("ul, ol"))) sub.remove();
    return clone;
  };

  const cellData = (cell, frame) => {
    const cs = getComputedStyle(cell);
    const rect = cell.getBoundingClientRect();
    const fill = color(cs.backgroundColor);
    return {
      runs: runsOf(cell),
      header: cell.tagName === "TH",
      colspan: parseInt(cell.getAttribute("colspan") || "1", 10) || 1,
      rowspan: parseInt(cell.getAttribute("rowspan") || "1", 10) || 1,
      width: round(rect.width),
      align: cs.textAlign,
      valign: cs.verticalAlign,
      fill: fill && fill.alpha > 0.15 ? fill.hex : null,
      borders: sideBorders(cs),
      padding: {
        top: px(cs.paddingTop), right: px(cs.paddingRight),
        bottom: px(cs.paddingBottom), left: px(cs.paddingLeft),
      },
    };
  };

  const tableData = (table, frame) => {
    const rows = [];
    for (const tr of Array.from(table.querySelectorAll("tr"))) {
      const cs = getComputedStyle(tr);
      const rect = tr.getBoundingClientRect();
      if (!visible(cs, rect)) continue;
      const cells = Array.from(tr.children)
        .filter((cell) => cell.tagName === "TD" || cell.tagName === "TH")
        .map((cell) => cellData(cell, frame));
      if (cells.length) {
        rows.push({ cells, head: !!tr.closest("thead"), height: round(rect.height) });
      }
    }
    if (!rows.length) return null;
    const rect = table.getBoundingClientRect();
    return {
      kind: "table",
      rows,
      width: round(rect.width),
      indentLeft: round(Math.max(0, rect.left - frame.left)),
      top: round(rect.top),
      bottom: round(rect.bottom),
    };
  };

  const gridTable = (el, kids, columns, frame) => {
    const rows = [];
    for (let i = 0; i < kids.length; i += columns) {
      const cells = kids.slice(i, i + columns).map((kid) => {
        const cs = getComputedStyle(kid);
        const rect = kid.getBoundingClientRect();
        const fill = color(cs.backgroundColor);
        return {
          runs: runsOf(kid),
          header: false,
          colspan: 1,
          rowspan: 1,
          width: round(rect.width),
          align: cs.textAlign,
          valign: "top",
          fill: fill && fill.alpha > 0.15 ? fill.hex : null,
          borders: sideBorders(cs),
          padding: { top: 2, right: 4, bottom: 2, left: 0 },
        };
      });
      while (cells.length < columns) {
        cells.push({ runs: [], header: false, colspan: 1, rowspan: 1, width: cells[0].width,
          align: "left", valign: "top", fill: null, borders: null,
          padding: { top: 2, right: 4, bottom: 2, left: 0 } });
      }
      rows.push({ cells, head: false, height: 0 });
    }
    const rect = el.getBoundingClientRect();
    return {
      kind: "table", rows, borderless: true,
      width: round(rect.width),
      indentLeft: round(Math.max(0, rect.left - frame.left)),
      top: round(rect.top), bottom: round(rect.bottom),
    };
  };

  const collect = (el, frame, blocks, box, flow) => {
    const cs = getComputedStyle(el);
    const rect = el.getBoundingClientRect();
    if (SKIP_TAGS.has(el.tagName)) return;
    if (!visible(cs, rect)) return;

    const attrs = directivesOf(el);
    const enter = pick(attrs, ENTER_KEYS);
    const nextFlow = { ...(flow || {}), ...pick(attrs, INHERIT_KEYS) };
    const self = { ...pick(attrs, SELF_KEYS), ...enter };

    if (attrs.toc) {
      emit(blocks, {
        kind: "toc",
        levels: tocLevels(attrs.toc),
        runs: hasText(el) ? runsOf(el) : [],
        top: round(rect.top),
        bottom: round(rect.bottom),
      }, self, nextFlow);
      return;
    }

    if (el.tagName === "IMG") {
      const src = el.currentSrc || el.src || "";
      if (embeddable(src)) {
        emit(blocks, {
          kind: "image", src,
          width: round(rect.width), height: round(rect.height),
          indentLeft: round(Math.max(0, rect.left - frame.left)),
          top: round(rect.top), bottom: round(rect.bottom), box,
        }, self, nextFlow);
      }
      return;
    }
    if (el.tagName === "HR") {
      emit(blocks, {
        kind: "divider", color: (color(cs.borderTopColor) || { hex: "CCCCCC" }).hex,
        top: round(rect.top), bottom: round(rect.bottom),
        indentLeft: round(Math.max(0, rect.left - frame.left)),
      }, self, nextFlow);
      return;
    }
    if (el.tagName === "TABLE") {
      const table = tableData(el, frame);
      if (table) emit(blocks, table, self, nextFlow);
      return;
    }
    if (el.tagName === "UL" || el.tagName === "OL") {
      const items = listItems(el, frame);
      if (items.length) {
        emit(blocks, {
          kind: "list", ordered: el.tagName === "OL", items, box,
          top: round(rect.top), bottom: round(rect.bottom),
        }, self, nextFlow);
      }
      return;
    }

    const own = boxOf(el, cs);
    const inherited = own ? own : box;
    const blockKids = Array.from(el.children).filter(
      (kid) => !SKIP_TAGS.has(kid.tagName) && isBlockish(kid) &&
        visible(getComputedStyle(kid), kid.getBoundingClientRect()) &&
        (hasText(kid) || kid.tagName === "IMG" || kid.tagName === "HR" ||
         kid.tagName === "TABLE" || kid.querySelector("img")),
    );

    if (blockKids.length) {
      if (Object.keys(enter).length) pending = { ...(pending || {}), ...enter };
      // A row whose own child is a heading is a title bar, not a grid of
      // cells: collapsing it into a table would bury the heading in a cell,
      // where Word gives it no outline level and the table of contents cannot
      // see it. A heading deeper down (a card grid) still travels as a cell,
      // because there the grid is the content.
      const titleBar = attrs.row === "stack" ||
        Array.from(el.children).some((kid) => /^H[1-6]$/.test(kid.tagName));
      const columns = titleBar ? 0 : columnsOf(el, cs);
      if (columns > 1) {
        emit(blocks, gridTable(el, blockKids, columns, frame), self, nextFlow);
        return;
      }
      // Walk the nodes in order: a container can hold text of its own between
      // two child blocks, and that text is content like any other.
      let inline = [];
      const flush = () => {
        if (!inline.length) return;
        const runs = runsOf(inline);
        const nodes = inline;
        inline = [];
        if (!runs.some((run) => (run.text || "").trim())) return;
        const range = document.createRange();
        range.setStartBefore(nodes[0]);
        range.setEndAfter(nodes[nodes.length - 1]);
        const inlineRect = range.getBoundingClientRect();
        emit(blocks, {
          kind: "paragraph", level: 0, runs, box: inherited,
          style: paragraphStyle(el, cs, inlineRect.width ? inlineRect : rect, frame),
          top: round(inlineRect.top || rect.top),
          bottom: round(inlineRect.bottom || rect.top),
        }, {}, nextFlow);
      };
      for (const node of Array.from(el.childNodes)) {
        if (node.nodeType === 3) {
          if ((node.textContent || "").trim()) inline.push(node);
          continue;
        }
        if (node.nodeType !== 1 || SKIP_TAGS.has(node.tagName)) continue;
        if (blockKids.includes(node)) {
          flush();
          collect(node, frame, blocks, inherited, nextFlow);
          continue;
        }
        inline.push(node);
      }
      flush();
      return;
    }
    if (!hasText(el)) {
      // A thin decorated strip is a rule, whatever tag it uses.
      if (rect.height <= 6 && rect.width > 40 && color(cs.backgroundColor)) {
        emit(blocks, {
          kind: "divider", color: color(cs.backgroundColor).hex,
          top: round(rect.top), bottom: round(rect.bottom),
          indentLeft: round(Math.max(0, rect.left - frame.left)),
        }, self, nextFlow);
      }
      return;
    }
    const runs = runsOf(el);
    if (!runs.length) return;
    const heading = /^H([1-6])$/.exec(el.tagName);
    emit(blocks, {
      kind: "paragraph",
      level: heading ? parseInt(heading[1], 10) : 0,
      runs,
      box: inherited,
      style: paragraphStyle(el, cs, rect, frame),
      top: round(rect.top),
      bottom: round(rect.bottom),
    }, self, nextFlow);
  };

  for (const root of roots) {
    const cs = getComputedStyle(root);
    const rect = root.getBoundingClientRect();
    if (!visible(cs, rect)) continue;
    const frame = contentBox(root, cs, rect);
    const frameBox = { left: frame.x, right: frame.x + frame.w };
    const blocks = [];
    const floating = [];
    const pageDoc = directivesOf(root);
    const pageFlow = pick(pageDoc, INHERIT_KEYS);
    for (const kid of Array.from(root.children)) {
      const kidCs = getComputedStyle(kid);
      const kidRect = kid.getBoundingClientRect();
      if (!visible(kidCs, kidRect)) continue;
      // Absolutely placed strips at the very bottom are running footers.
      const pinned = kidCs.position === "absolute" || kidCs.position === "fixed";
      if (pinned && kidRect.bottom > rect.bottom - rect.height * 0.12) {
        const runs = [];
        for (const part of Array.from(kid.children).length ? Array.from(kid.children) : [kid]) {
          const partRuns = runsOf(part);
          if (partRuns.length) {
            if (runs.length) runs.push({ text: "\t", tab: true });
            runs.push(...partRuns);
          }
        }
        if (runs.length) floating.push({ runs, align: getComputedStyle(kid).textAlign });
        continue;
      }
      collect(kid, frameBox, blocks, null, pageFlow);
    }
    out.pages.push({
      width: round(rect.width),
      height: round(rect.height),
      margins: {
        top: round(frame.y - rect.top),
        right: round(rect.right - (frame.x + frame.w)),
        bottom: round(Math.max(rect.bottom - (frame.y + frame.h), 12)),
        left: round(frame.x - rect.left),
      },
      background: (color(getComputedStyle(root).backgroundColor) || {}).hex || null,
      blocks,
      footer: floating[0] || null,
      doc: Object.keys(pageDoc).length ? pageDoc : null,
    });
  }
  publish(out);
"""


@dataclass
class PageData:
    """One measured ``.page`` section, as a flow of Word-ready blocks."""

    width: float
    height: float
    margins: dict[str, float]
    blocks: list[dict[str, Any]] = field(default_factory=list)
    footer: dict[str, Any] | None = None
    doc: dict[str, Any] | None = None


def page_top(page: PageData) -> float:
    """Vertical origin of a page, so block offsets are page-relative."""
    return min((float(block.get("top", 0)) for block in page.blocks), default=0.0)


def _signature(block: dict[str, Any]) -> str:
    """Text identity of a block, used to spot page furniture."""
    if block.get("kind") == "paragraph":
        text = "".join(str(run.get("text", "")) for run in block.get("runs") or [])
        return f"p:{block.get('level')}:{' '.join(text.split()).lower()}"
    if block.get("kind") == "list":
        items = [
            "".join(str(run.get("text", "")) for run in (item.get("runs") or []))
            for item in block.get("items") or []
        ]
        return "l:" + "|".join(" ".join(item.split()).lower() for item in items)
    if block.get("kind") == "table":
        cells = [
            "".join(str(run.get("text", "")) for run in (cell.get("runs") or []))
            for row in block.get("rows") or []
            for cell in row.get("cells") or []
        ]
        return "t:" + "|".join(" ".join(cell.split()).lower() for cell in cells)
    return f"{block.get('kind')}:{block.get('top')}"


def strip_page_furniture(pages: list[PageData]) -> list[dict[str, Any]]:
    """Lift banners repeated on every page out of the flow.

    An HTML master repeats its letterhead on each ``.page`` because every page
    stands alone. Word does not work that way: a single flow with the banner
    typed six times reads as a mistake. Repeated top-of-page blocks move into
    the running header, and the extra copies disappear.
    """
    if len(pages) < 2:
        return []
    threshold = max(2, len(pages) // 2)
    counts: dict[str, int] = {}
    for page in pages:
        cutoff = float(page.margins.get("top", 0)) + page.height * 0.22
        for block in page.blocks:
            if float(block.get("top", 0)) - page_top(page) <= cutoff:
                counts[_signature(block)] = counts.get(_signature(block), 0) + 1
    repeated = {key for key, count in counts.items() if count >= threshold}
    if not repeated:
        return []
    header: list[dict[str, Any]] = []
    seen: set[str] = set()
    for page in pages:
        kept = []
        for block in page.blocks:
            key = _signature(block)
            if key in repeated:
                if key not in seen:
                    seen.add(key)
                    header.append(block)
                continue
            kept.append(block)
        page.blocks = kept
    return header


def measure_document(source: Path, chromium: str, workdir: Path, timeout: int) -> list[PageData]:
    """Measure an HTML document and return its pages."""
    payload = measure(
        source,
        _dom.script(_MEASURE_BODY, PAGE_WIDTH_PX, 0),
        chromium=chromium,
        workdir=workdir,
        timeout=timeout,
        viewport=(PAGE_WIDTH_PX, PAGE_HEIGHT_PX * 3),
        prefix="doc",
    )
    pages = [
        PageData(
            width=float(page.get("width") or PAGE_WIDTH_PX),
            height=float(page.get("height") or PAGE_HEIGHT_PX),
            margins=page.get("margins") or {},
            blocks=page.get("blocks") or [],
            footer=page.get("footer"),
            doc=page.get("doc") if isinstance(page.get("doc"), dict) else None,
        )
        for page in payload.get("pages", [])
    ]
    if not pages:
        raise ConversionError(f"{source.name}: no page section found")
    for page in pages:
        for block in page.blocks:
            if block.get("kind") == "image":
                block["path"] = resolve_asset(block.get("src", ""), source, workdir, id(block))
    return pages


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_ALIGNMENTS = {
    "left": "LEFT", "start": "LEFT", "right": "RIGHT", "end": "RIGHT",
    "center": "CENTER", "justify": "JUSTIFY",
}
_BORDER_SIZES = {"top": "top", "right": "right", "bottom": "bottom", "left": "left"}


def _element(tag: str, **attributes: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    node = OxmlElement(f"w:{tag}")
    for key, value in attributes.items():
        node.set(qn(f"w:{key}"), value)
    return node


def _shade(properties, color: str) -> None:
    properties.append(_element("shd", val="clear", color="auto", fill=color))


def _paragraph_borders(paragraph, borders: dict[str, Any], sides: Sequence[str]) -> None:
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    holder = properties.find(qn("w:pBdr"))
    if holder is None:
        holder = _element("pBdr")
        properties.append(holder)
    for side in sides:
        spec = borders.get(side)
        if not spec:
            continue
        # Word measures borders in eighths of a point, and refuses under 2.
        size = max(2, min(96, int(round(float(spec["width"]) * PT_PER_PX * 8))))
        holder.append(
            _element(
                _BORDER_SIZES[side],
                val="single",
                sz=str(size),
                space="4",
                color=str(spec["color"]),
            )
        )


def _apply_runs(paragraph, runs: Sequence[dict[str, Any]], keep_fonts: bool, scale: float) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    for data in runs:
        if data.get("newline"):
            paragraph.add_run().add_break()
            continue
        if data.get("tab"):
            paragraph.add_run().add_tab()
            continue
        text = str(data.get("text", ""))
        if not text and not data.get("field"):
            continue
        run = paragraph.add_run(text)
        if data.get("field"):
            run._element.set("data-field", str(data["field"]))
        font = run.font
        name = office_font(data.get("families") or [], data.get("font") or "Calibri", keep_fonts)
        font.name = name
        # python-docx only sets the ascii/hAnsi slots; complex scripts fall
        # back to the theme font without these.
        properties = run._element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = _element("rFonts")
            properties.insert(0, fonts)
        for slot in ("cs", "eastAsia"):
            fonts.set(qn(f"w:{slot}"), name)
        font.size = Pt(max(1.0, float(data.get("size") or 16) * PT_PER_PX * scale))
        font.bold = bool(data.get("bold"))
        font.italic = bool(data.get("italic"))
        font.underline = bool(data.get("underline"))
        if data.get("strike"):
            font.strike = True
        if data.get("color"):
            font.color.rgb = RGBColor.from_string(str(data["color"]))
        if data.get("highlight"):
            _shade(properties, str(data["highlight"]))
        spacing = float(data.get("spacing") or 0)
        if abs(spacing) > 0.05:
            properties.append(_element("spacing", val=str(int(round(spacing * 20 * PT_PER_PX)))))


def _style_paragraph(paragraph, style: dict[str, Any], scale: float, gap: float) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    fmt = paragraph.paragraph_format
    fmt.alignment = getattr(
        WD_ALIGN_PARAGRAPH, _ALIGNMENTS.get(str(style.get("align") or "left"), "LEFT")
    )
    fmt.left_indent = Pt(float(style.get("indentLeft") or 0) * PT_PER_PX * scale)
    fmt.right_indent = Pt(float(style.get("indentRight") or 0) * PT_PER_PX * scale)
    fmt.space_before = Pt(max(0.0, gap) * PT_PER_PX * scale)
    fmt.space_after = Pt(0)
    line_height = style.get("lineHeight")
    if line_height:
        fmt.line_spacing = Pt(float(line_height) * PT_PER_PX * scale)
    if style.get("rtl"):
        paragraph._p.get_or_add_pPr().append(_element("bidi", val="1"))


def _box_sides(block: dict[str, Any], previous_box: str | None, next_box: str | None) -> list[str]:
    """Which sides of a shared container border this paragraph carries."""
    box = block.get("box") or {}
    identifier = box.get("id")
    sides = ["left", "right"]
    if identifier != previous_box:
        sides.append("top")
    if identifier != next_box:
        sides.append("bottom")
    return sides


def _add_table(document, block: dict[str, Any], scale: float, keep_fonts: bool, gap: float):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt

    rows = block.get("rows") or []
    columns = max(sum(cell.get("colspan", 1) for cell in row["cells"]) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    properties.append(_element("tblLayout", type="fixed"))
    if block.get("indentLeft"):
        properties.append(
            _element("tblInd", w=str(int(float(block["indentLeft"]) * 15 * scale)), type="dxa")
        )
    # The default Table Grid borders would draw lines the design does not have.
    borders = _element("tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_element(side, val="none", sz="0", space="0", color="auto"))
    properties.append(borders)

    widths = [0.0] * columns
    for row in rows:
        index = 0
        for cell in row["cells"]:
            span = cell.get("colspan", 1)
            if span == 1 and index < columns:
                widths[index] = max(widths[index], float(cell.get("width") or 0))
            index += span
    total = sum(widths) or 1.0
    available = float(block.get("width") or PAGE_WIDTH_PX)
    for index, width in enumerate(widths):
        points = Pt((width or total / columns) / total * available * PT_PER_PX * scale)
        for row in table.rows:
            row.cells[index].width = points

    for row_index, row in enumerate(rows):
        column = 0
        for cell_data in row["cells"]:
            if column >= columns:
                break
            cell = table.cell(row_index, column)
            span = min(cell_data.get("colspan", 1), columns - column)
            if span > 1:
                cell = cell.merge(table.cell(row_index, column + span - 1))
            row_span = cell_data.get("rowspan", 1)
            if row_span > 1 and row_index + row_span - 1 < len(rows):
                cell = cell.merge(table.cell(row_index + row_span - 1, column))
            _fill_cell(cell, cell_data, block, scale, keep_fonts)
            column += span
    if rows and rows[0].get("head"):
        table.rows[0]._tr.get_or_add_trPr().append(_element("tblHeader", val="true"))
    for row in table.rows:
        _word.keep_row_together(row)
    if gap:
        _spacer(document, gap, scale)
    return table


def _fill_cell(cell, data: dict[str, Any], block: dict[str, Any], scale: float, keep: bool) -> None:
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    properties = cell._tc.get_or_add_tcPr()
    if data.get("fill"):
        _shade(properties, str(data["fill"]))
    borders = data.get("borders")
    if borders and not block.get("borderless"):
        holder = _element("tcBorders")
        for side, spec in borders.items():
            size = max(2, min(96, int(round(float(spec["width"]) * PT_PER_PX * 8))))
            holder.append(
                _element(side, val="single", sz=str(size), space="0", color=str(spec["color"]))
            )
        properties.append(holder)
    padding = data.get("padding") or {}
    if padding:
        margins = _element("tcMar")
        for side in ("top", "left", "bottom", "right"):
            value = int(round(float(padding.get(side) or 0) * 15 * scale))
            margins.append(_element(side, w=str(value), type="dxa"))
        properties.append(margins)
    valign = str(data.get("valign") or "top")
    cell.vertical_alignment = (
        WD_ALIGN_VERTICAL.CENTER if valign in {"middle", "center"}
        else WD_ALIGN_VERTICAL.BOTTOM if valign == "bottom"
        else WD_ALIGN_VERTICAL.TOP
    )
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = getattr(
        WD_ALIGN_PARAGRAPH, _ALIGNMENTS.get(str(data.get("align") or "left"), "LEFT")
    )
    _apply_runs(paragraph, data.get("runs") or [], keep, scale)


def _spacer(document, gap: float, scale: float) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(max(1.0, gap * PT_PER_PX * scale))
    paragraph.add_run("")


def _heading_style(document, level: int) -> str:
    name = f"Heading {min(level, 9)}"
    try:
        document.styles[name]
        return name
    except KeyError:  # pragma: no cover - default template always has them
        return "Normal"


def _truthy(value: str | None) -> bool:
    return (value or "").strip().lower() in {"1", "true", "yes", "on"}


def _block_text(block: dict[str, Any]) -> str:
    return "".join(str(run.get("text", "")) for run in block.get("runs") or [])


def _is_caption(block: dict[str, Any] | None) -> bool:
    if not block:
        return False
    doc = _word.parse_directives(block.get("doc"))
    return bool(doc.get("caption") or (doc.get("role") or "").lower() == "caption")


def _layout_of(
    doc: dict[str, str], fallback: tuple[bool, int] = (False, 1)
) -> tuple[bool, int]:
    """Read landscape/columns from a directive payload."""
    landscape, columns = fallback
    section = (doc.get("section") or "").lower()
    if section in {"landscape", "land", "horizontal"}:
        landscape = True
    elif section in {"portrait", "port", "vertical"}:
        landscape = False
    raw = doc.get("columns")
    if raw and str(raw).isdigit():
        columns = max(1, min(3, int(raw)))
    return landscape, columns


def _page_layout(page: PageData) -> tuple[bool, int]:
    return _layout_of(_word.parse_directives(page.doc))


def _numbering_spec(pages: Sequence[PageData]) -> str | None:
    for page in pages:
        sources: list[Any] = [page.doc]
        sources.extend(block.get("doc") for block in page.blocks)
        for source in sources:
            raw = _word.parse_directives(source).get("numbering")
            if not raw:
                continue
            if raw.lower() in {"true", "1", "yes", "on"}:
                return "1-4"
            return raw
    return None


def _apply_page_size(
    section, page: PageData, scale: float, landscape: bool, columns: int
) -> None:
    from docx.shared import Emu, Pt

    section.page_width = Emu(int(210 * 36000))
    section.page_height = Emu(int(297 * 36000))
    margins = page.margins
    section.top_margin = Pt(float(margins.get("top", 54)) * PT_PER_PX * scale)
    section.bottom_margin = Pt(float(margins.get("bottom", 54)) * PT_PER_PX * scale)
    section.left_margin = Pt(float(margins.get("left", 54)) * PT_PER_PX * scale)
    section.right_margin = Pt(float(margins.get("right", 54)) * PT_PER_PX * scale)
    _word.set_orientation(section, landscape)
    _word.set_columns(section, columns)


@dataclass
class _WriteState:
    bookmark_id: int = 0
    layout: tuple[bool, int] = (False, 1)
    used_fields: bool = False


def _ensure_layout(
    document, page: PageData, scale: float, wanted: tuple[bool, int], state: _WriteState
) -> None:
    if wanted == state.layout:
        return
    from docx.enum.section import WD_SECTION

    section = document.add_section(WD_SECTION.NEW_PAGE)
    _apply_page_size(section, page, scale, wanted[0], wanted[1])
    state.layout = wanted


def _apply_pagination(
    paragraph,
    block: dict[str, Any],
    *,
    keep_next: bool = False,
    keep_lines: bool = False,
    page_break: bool = False,
) -> None:
    doc = _word.parse_directives(block.get("doc"))
    extra_next, extra_lines = _word.keep_flags(doc.get("keep") or "")
    brk = (doc.get("break") or "").lower()
    _word.set_pagination(
        paragraph,
        keep_next=keep_next or extra_next,
        keep_lines=keep_lines or extra_lines,
        page_break_before=page_break or brk in {"page", "before"},
        widow_control=True,
    )


def _attach_bookmark(paragraph, block: dict[str, Any], state: _WriteState) -> None:
    name = _word.parse_directives(block.get("doc")).get("bookmark")
    if not name:
        return
    state.bookmark_id += 1
    _word.add_bookmark(paragraph, name, state.bookmark_id)


def _attach_ref(paragraph, block: dict[str, Any], state: _WriteState) -> None:
    name = _word.parse_directives(block.get("doc")).get("ref")
    if not name:
        return
    _word.add_cross_reference(paragraph, name, placeholder=name)
    state.used_fields = True


_TITLE_RE = re.compile(r"<title[^>]*>(.*?)</title>", re.S | re.I)
_LANG_RE = re.compile(r"<html\b[^>]*\blang=[\"']([A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})*)[\"']", re.I)


def document_meta(source: Path) -> dict[str, str]:
    """Title and language declared by the HTML, for the DOCX properties.

    Word shows the title in the file properties and the recent-files list, and
    uses the language for spell checking and hyphenation; a document without
    them looks machine-made the moment it is opened. ``lang="und"`` (the
    template placeholder for "not chosen yet") counts as no language.
    """
    try:
        html = source.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return {}
    meta: dict[str, str] = {}
    match = _TITLE_RE.search(html)
    if match:
        title = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", match.group(1))).strip()
        if title:
            meta["title"] = title
    match = _LANG_RE.search(html)
    if match and match.group(1).lower() not in {"und", "zxx", "mul"}:
        meta["lang"] = match.group(1)
    return meta


def _apply_document_meta(document, meta: dict[str, str]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    title = meta.get("title")
    if title:
        document.core_properties.title = title[:255]
    lang = meta.get("lang")
    if not lang:
        return
    document.core_properties.language = lang
    # The default run properties are what Word consults for every run without
    # its own w:lang, so the proofing language follows the document's.
    styles = document.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        return
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.insert(0, rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    node = rpr.find(qn("w:lang"))
    if node is None:
        node = OxmlElement("w:lang")
        rpr.append(node)
    node.set(qn("w:val"), lang)
    if lang.lower().startswith(("ar", "he", "fa", "ur")):
        node.set(qn("w:bidi"), lang)


def build_document(
    pages: Sequence[PageData],
    output: Path,
    keep_fonts: bool = False,
    meta: dict[str, str] | None = None,
) -> Path:
    """Assemble measured pages into a native DOCX file."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    if meta:
        _apply_document_meta(document, meta)

    pages = list(pages)
    banner = strip_page_furniture(pages)
    numbering = _numbering_spec(pages)
    if numbering:
        levels = list(_word.iter_levels(numbering))
        _word.bind_heading_numbering(document, levels[-1] if levels else 4)
    state = _WriteState()

    for index, page in enumerate(pages):
        scale = PAGE_WIDTH_PX / page.width if page.width else 1.0
        wanted = _page_layout(page)
        if index == 0:
            section = document.sections[0]
            _apply_page_size(section, page, scale, wanted[0], wanted[1])
            state.layout = wanted
            running = next((p.footer for p in pages if p.footer), None)
            if running:
                _write_footer(section, running, scale, keep_fonts)
            if banner:
                _write_banner(section, banner, scale, keep_fonts)
            if _truthy(_word.parse_directives(page.doc).get("cover")):
                section.different_first_page = True
        elif wanted != state.layout:
            _ensure_layout(document, page, scale, wanted, state)
        else:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _write_blocks(document, page, scale, keep_fonts, state)

    if state.used_fields:
        _word.request_field_update(document)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def _write_footer(section, footer: dict[str, Any], scale: float, keep_fonts: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Pt

    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    width = section.page_width - section.left_margin - section.right_margin
    paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    paragraph.paragraph_format.space_before = Pt(0)
    _apply_runs(paragraph, _numbered_runs(footer.get("runs") or []), keep_fonts, scale)
    _fill_page_fields(paragraph)


_PAGE_NUMBER_RE = re.compile(
    r"(?P<label>\b(?:page|pág|pag|seite|pagina|صفحة)\b\s*)?"
    r"(?P<current>\d{1,3})\s*(?P<sep>/|of|sur|von|de|di|van)\s*(?P<total>\d{1,3})",
    re.IGNORECASE,
)


def _numbered_runs(runs: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    """Mark "Page 3 / 8" in a footer so Word can number the pages itself."""
    out: list[dict[str, Any]] = []
    done = False
    for run in runs:
        text = str(run.get("text", ""))
        match = None if done else _PAGE_NUMBER_RE.search(text)
        if not match:
            out.append(run)
            continue
        done = True
        before = text[: match.start()] + (match.group("label") or "")
        separator = f" {match.group('sep')} "
        if before:
            out.append({**run, "text": before})
        out.append({**run, "text": "", "field": "PAGE"})
        out.append({**run, "text": separator})
        out.append({**run, "text": "", "field": "NUMPAGES"})
        after = text[match.end():]
        if after:
            out.append({**run, "text": after})
    return out


def _fill_page_fields(paragraph) -> None:
    """Turn the marked runs into real PAGE / NUMPAGES fields."""
    for run in paragraph.runs:
        instruction = run._element.get("data-field")
        if not instruction:
            continue
        del run._element.attrib["data-field"]
        begin = _element("fldChar", fldCharType="begin")
        code = _element("instrText", space="preserve")
        code.text = f" {instruction} "
        end = _element("fldChar", fldCharType="end")
        for node in (begin, code, end):
            run._element.append(node)


def _write_banner(section, blocks: Sequence[dict[str, Any]], scale: float, keep: bool) -> None:
    """Put the letterhead repeated on every master page into the header."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Pt

    header = section.header
    header.is_linked_to_previous = False
    width = section.page_width - section.left_margin - section.right_margin
    written = 0
    for block in blocks:
        runs = _banner_runs(block)
        if not runs:
            continue
        paragraph = header.paragraphs[0] if written == 0 else header.add_paragraph()
        written += 1
        _style_paragraph(paragraph, block.get("style") or {}, scale, 0)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        if any(run.get("tab") for run in runs):
            paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        _apply_runs(paragraph, runs, keep, scale)


def _banner_runs(block: dict[str, Any]) -> list[dict[str, Any]]:
    """Flatten a repeated block into the single line a header holds.

    Letterheads are often laid out as a row of cells; a header has no room for
    a table, so the cells become tab-separated runs on one line.
    """
    if block.get("kind") == "paragraph":
        return list(block.get("runs") or [])
    if block.get("kind") == "table":
        runs: list[dict[str, Any]] = []
        for row in block.get("rows") or []:
            for cell in row.get("cells") or []:
                cell_runs = cell.get("runs") or []
                if not cell_runs:
                    continue
                if runs:
                    runs.append({"text": "\t", "tab": True})
                runs.extend(cell_runs)
        return runs
    if block.get("kind") == "list":
        runs = []
        for item in block.get("items") or []:
            if runs:
                runs.append({"text": " · "})
            runs.extend(item.get("runs") or [])
        return runs
    return []


def _write_blocks(
    document, page: PageData, scale: float, keep_fonts: bool, state: _WriteState | None = None
) -> None:
    from docx.shared import Emu, Pt

    state = state or _WriteState()
    blocks = page.blocks
    # Blocks carry viewport coordinates: on page four they start around
    # y=3400, so spacing is measured from the top of their own page.
    cursor = page_top(page)
    for position, block in enumerate(blocks):
        gap = min(max(0.0, float(block.get("top", cursor)) - cursor), page.height * 0.5)
        kind = block.get("kind")
        nxt = blocks[position + 1] if position + 1 < len(blocks) else None
        doc = _word.parse_directives(block.get("doc"))
        if doc.get("section") or doc.get("columns"):
            _ensure_layout(document, page, scale, _layout_of(doc, state.layout), state)

        if kind == "toc":
            title = _block_text(block).strip()
            if title:
                heading = document.add_paragraph(style="Heading 1")
                _style_paragraph(heading, block.get("style") or {}, scale, gap)
                _apply_runs(heading, block.get("runs") or [], keep_fonts, scale)
                _apply_pagination(heading, block, keep_next=True, keep_lines=True)
                gap = 0
            paragraph = document.add_paragraph()
            _word.add_table_of_contents(paragraph, str(block.get("levels") or "1-3"))
            state.used_fields = True
            cursor = float(block.get("bottom", cursor))
            continue
        if kind == "table":
            table = _add_table(document, block, scale, keep_fonts, gap=0)
            if table is not None and _is_caption(nxt):
                last = table.rows[-1].cells[0].paragraphs[-1]
                _apply_pagination(last, block, keep_next=True)
            cursor = float(block.get("bottom", cursor))
            continue
        if kind == "image":
            path = block.get("path")
            if path:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(gap * PT_PER_PX * scale)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.left_indent = Pt(
                    float(block.get("indentLeft") or 0) * PT_PER_PX * scale
                )
                paragraph.add_run().add_picture(
                    str(path), width=Emu(int(float(block["width"]) * 9525 * scale))
                )
                _apply_pagination(
                    paragraph, block, keep_next=_is_caption(nxt), keep_lines=True
                )
                _attach_bookmark(paragraph, block, state)
            cursor = float(block.get("bottom", cursor))
            continue
        if kind == "divider":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(gap * PT_PER_PX * scale)
            paragraph.paragraph_format.space_after = Pt(0)
            _paragraph_borders(
                paragraph, {"bottom": {"width": 1.2, "color": block.get("color", "CCCCCC")}},
                ["bottom"],
            )
            cursor = float(block.get("bottom", cursor))
            continue
        if kind == "list":
            cursor = _write_list(document, block, scale, keep_fonts, gap, cursor)
            continue

        style = block.get("style") or {}
        level = int(block.get("level") or 0)
        caption = _is_caption(block)
        paragraph = document.add_paragraph(
            style=_heading_style(document, level) if level and not caption else None
        )
        _style_paragraph(paragraph, style, scale, gap)
        if caption:
            label = doc.get("caption") or "Figure"
            if _truthy(label):
                label = "Figure"
            paragraph.add_run(f"{label} ")
            _word.add_sequence(paragraph, re.sub(r"\s+", "", label) or "Figure")
            paragraph.add_run(" - ")
            state.used_fields = True
        _apply_runs(paragraph, block.get("runs") or [], keep_fonts, scale)
        _attach_bookmark(paragraph, block, state)
        _attach_ref(paragraph, block, state)
        box = block.get("box")
        if box:
            previous = (blocks[position - 1].get("box") or {}).get("id") if position else None
            following = (nxt.get("box") or {}).get("id") if nxt else None
            sides = _box_sides(block, previous, following)
            if box.get("fill"):
                _shade(paragraph._p.get_or_add_pPr(), str(box["fill"]))
            if box.get("borders"):
                _paragraph_borders(paragraph, box["borders"], sides)
        keep_next = level >= 1 or (caption and nxt is not None and nxt.get("kind") == "image")
        _apply_pagination(paragraph, block, keep_next=keep_next, keep_lines=level >= 1 or caption)
        cursor = float(block.get("bottom", cursor))


def _write_list(
    document, block: dict[str, Any], scale: float, keep_fonts: bool, gap: float, cursor: float
) -> float:
    from docx.shared import Pt

    ordered = bool(block.get("ordered"))
    for index, item in enumerate(block.get("items") or []):
        level = min(int(item.get("level") or 0), 2)
        base = "List Number" if ordered else "List Bullet"
        name = base if level == 0 else f"{base} {level + 1}"
        try:
            document.styles[name]
        except KeyError:  # pragma: no cover - default template has 1..3
            name = base
        paragraph = document.add_paragraph(style=name)
        style = item.get("style") or {}
        _style_paragraph(paragraph, style, scale, gap if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(0)
        _apply_runs(paragraph, item.get("runs") or [], keep_fonts, scale)
    return float(block.get("bottom", cursor))


def check_qa(page: Path, force: bool = False) -> None:
    """Refuse to ship a document its own quality pass already rejected.

    ``word_qa`` writes ``qa.json`` beside the document and exits non-zero
    below the threshold, but an exit code is easy to ignore. This gate is
    not: a failing report stops the conversion until the pages are fixed
    (or the caller passes ``--force`` because the user accepts the document
    as is).
    """
    qa = page.parent / "qa.json"
    if not qa.is_file():
        return
    try:
        data = json.loads(qa.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return
    if data.get("pass") is not False:
        return
    score = data.get("score")
    threshold = data.get("threshold")
    rework = ", ".join(str(n) for n in (data.get("rework") or [])[:6])
    message = (
        f"{qa}: the document QA failed (score {score}, threshold {threshold}"
        + (f", pages to rework: {rework}" if rework else "")
        + ")."
    )
    if force:
        print(f"html2docx: {message} (converted anyway: --force)", file=sys.stderr)
        return
    raise ConversionError(
        message
        + " Fix the pages, re-run word_qa until it passes, then convert."
        " Use --force only if the user explicitly accepts the document as is."
    )


def convert(
    source: str,
    output: Path,
    *,
    chromium: str | None = None,
    timeout: int = 90,
    keep_workdir: Path | None = None,
    keep_fonts: bool = False,
    force: bool = False,
) -> Path:
    """Convert an HTML document into an editable DOCX and return its path."""
    page = Path(source).expanduser()
    if page.is_dir():
        candidates = sorted(page.glob("document.html")) or sorted(page.glob("*.html"))
        if not candidates:
            raise ConversionError(f"No HTML document in {page}")
        page = candidates[0]
    if not page.is_file():
        raise ConversionError(f"Not found: {source}")
    check_qa(page, force=force)
    binary = find_chromium(chromium)
    workdir = keep_workdir or Path(tempfile.mkdtemp(prefix="navin-html2docx-"))
    workdir.mkdir(parents=True, exist_ok=True)
    try:
        pages = measure_document(page, binary, workdir, timeout)
        return build_document(pages, output, keep_fonts=keep_fonts, meta=document_meta(page))
    finally:
        if keep_workdir is None:
            shutil.rmtree(workdir, ignore_errors=True)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="html2docx",
        description="Convert an A4 HTML document into an editable DOCX file.",
    )
    parser.add_argument("input", help="document.html, or the folder holding it")
    parser.add_argument("-o", "--output", required=True, help="Destination .docx")
    parser.add_argument("--chromium", help="Path to a Chromium binary")
    parser.add_argument("--timeout", type=int, default=90, help="Chromium timeout")
    parser.add_argument("--keep-workdir", help="Keep intermediate files in this folder")
    parser.add_argument(
        "--keep-fonts",
        action="store_true",
        help="Keep CSS font names instead of mapping them onto Office fonts",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Convert even when qa.json rejected the document",
    )
    args = parser.parse_args(argv)
    try:
        output = convert(
            args.input,
            Path(args.output).expanduser(),
            chromium=args.chromium,
            timeout=args.timeout,
            keep_workdir=Path(args.keep_workdir).expanduser() if args.keep_workdir else None,
            keep_fonts=args.keep_fonts,
            force=args.force,
        )
    except ConversionError as exc:
        print(f"html2docx: {exc}", file=sys.stderr)
        return 1
    print(f"Wrote {output}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
