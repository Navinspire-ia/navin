"""Meeting translation, document export and calendar helpers."""

from __future__ import annotations

import asyncio
import base64
import io
import os
import re
import zipfile
from collections.abc import Mapping
from datetime import UTC, datetime
from typing import Any

from navin.agent.skill_routing import build_action_skill_context
from navin.webui.meeting_api import MeetingError, _ask, strip_report_fences

_TRANSLATE_SYSTEM = (
    "Translate meeting content accurately. Preserve Markdown structure, speaker labels, "
    "timecodes, names, numbers, decisions and action ownership. Output only the translation. "
    "Use plain hyphens. Do not summarize or add facts."
)
_CLEANUP_SYSTEM = (
    "Clean a meeting transcript with high accuracy. Fix punctuation, casing and obvious "
    "recognition errors only. Preserve every fact, speaker label, timecode and uncertainty. "
    "Do not summarize, translate or claim acoustic speaker identification. Output only text."
)
_MAX_TEXT_TRANSFORM_CHARS = 12000


def _check_transform_length(text: str) -> None:
    if len(text) > _MAX_TEXT_TRANSFORM_CHARS:
        raise MeetingError(
            f"translation and cleanup accept at most {_MAX_TEXT_TRANSFORM_CHARS} characters per call; "
            "split the text into smaller sections. No transformed text was generated.",
            status=413,
        )


async def translate_payload(
    *, text: str, target_language: str, source_language: str = "", kind: str = "transcript"
) -> dict[str, Any]:
    if not text.strip():
        raise MeetingError("text is required")
    if not target_language.strip():
        raise MeetingError("target language is required")
    _check_transform_length(text)
    prompt = (
        f"Content kind: {kind}\nSource language: {source_language or 'auto'}\n"
        f"Target language: {target_language}\n\n{text}"
    )
    skills = await asyncio.to_thread(build_action_skill_context, "meeting", "translate")
    translated, model, route = await _ask(
        skills.augment_system(_TRANSLATE_SYSTEM), prompt, max_tokens=5000, timeout_s=150.0
    )
    value = strip_report_fences(translated)
    if not value:
        raise MeetingError("the model returned an empty translation", status=502)
    return {"text": value, "model": model, "route": route, "kind": kind, "skill_context": skills.metadata}


async def cleanup_payload(*, transcript: str, language: str = "") -> dict[str, Any]:
    if not transcript.strip():
        raise MeetingError("transcript is required")
    _check_transform_length(transcript)
    prompt = f"Language: {language or 'auto'}\n\n{transcript}"
    skills = await asyncio.to_thread(build_action_skill_context, "meeting", "cleanup")
    cleaned, model, route = await _ask(
        skills.augment_system(_CLEANUP_SYSTEM), prompt, max_tokens=5000, timeout_s=150.0
    )
    value = strip_report_fences(cleaned)
    if not value:
        raise MeetingError("the model returned an empty transcript", status=502)
    return {
        "text": value,
        "model": model,
        "route": route,
        "skill_context": skills.metadata,
        "accuracy": "llm_text_cleanup",
        "acoustic_diarization": False,
    }


def _normalize_docx(data: bytes) -> bytes:
    source = io.BytesIO(data)
    target = io.BytesIO()
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as outgoing:
        for name in sorted(incoming.namelist()):
            info = zipfile.ZipInfo(name)
            info.date_time = (1980, 1, 1, 0, 0, 0)
            info.external_attr = 0o600 << 16
            outgoing.writestr(info, incoming.read(name))
    return target.getvalue()


def docx_payload(
    *, title: str, report: str, transcript: str = "", notes: str = ""
) -> dict[str, Any]:
    """Build deterministic DOCX bytes using the bundled document dependency."""
    try:
        from docx import Document
    except ImportError as exc:
        raise MeetingError("python-docx is not installed", status=503) from exc
    document = Document()
    core = document.core_properties
    core.title = title.strip() or "Meeting"
    core.author = "Navin"
    core.created = datetime(2000, 1, 1, tzinfo=UTC)
    core.modified = datetime(2000, 1, 1, tzinfo=UTC)
    document.add_heading(title.strip() or "Meeting", level=0)
    for heading, value in (("Report", report), ("Notes", notes), ("Transcript", transcript)):
        if not value.strip():
            continue
        document.add_heading(heading, level=1)
        for line in value.replace("\r\n", "\n").splitlines():
            clean = re.sub(r"^#{1,6}\s+", "", line).strip()
            if clean:
                document.add_paragraph(clean)
    buffer = io.BytesIO()
    document.save(buffer)
    data = _normalize_docx(buffer.getvalue())
    filename = re.sub(r"[^A-Za-z0-9._-]+", "-", title.strip()).strip("-") or "meeting"
    return {
        "filename": f"{filename}.docx",
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "data_base64": base64.b64encode(data).decode("ascii"),
        "bytes": len(data),
    }


def _ics_escape(value: Any) -> str:
    return (
        str(value or "")
        .replace("\\", "\\\\")
        .replace("\r", "")
        .replace("\n", "\\n")
        .replace(",", "\\,")
        .replace(";", "\\;")
    )


def calendar_payload(
    events: list[Mapping[str, Any]],
    *,
    provider: str = "ics",
    credentials: Mapping[str, Any] | None = None,
    allow_network: bool = False,
) -> dict[str, Any]:
    """Return an ICS fallback or a provider-ready, network-disabled result."""
    selected = provider.strip().lower() or "ics"
    if selected in {"google", "microsoft"}:
        configured = bool(credentials) or bool(
            os.environ.get(
                "GOOGLE_CALENDAR_CREDENTIALS"
                if selected == "google"
                else "MICROSOFT_CALENDAR_CREDENTIALS"
            )
        )
        if configured and allow_network and os.environ.get("CI", "").lower() not in {"1", "true"}:
            raise MeetingError(
                "calendar network sync requires an installed integration adapter", status=501
            )
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Navin//Meetings//EN"]
    for index, event in enumerate(events):
        uid = _ics_escape(event.get("id") or f"meeting-{index}@navin.local")
        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:{uid}",
                f"DTSTART:{_ics_escape(event.get('start') or '')}",
                f"DTEND:{_ics_escape(event.get('end') or '')}",
                f"SUMMARY:{_ics_escape(event.get('title') or 'Meeting')}",
                f"DESCRIPTION:{_ics_escape(event.get('description') or '')}",
                f"LOCATION:{_ics_escape(event.get('url') or event.get('location') or '')}",
                "END:VEVENT",
            ]
        )
    lines.append("END:VCALENDAR")
    return {
        "provider": selected,
        "configured": bool(credentials),
        "synced": False,
        "fallback": "ics",
        "ics": "\r\n".join(lines) + "\r\n",
    }
