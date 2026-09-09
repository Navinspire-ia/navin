# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""Ads Manager exports (CSV / TSV / XLSX) and MCP rows -> normalized ``AdRow``.

Handles what real exports look like: UTF-16 tab-separated Google downloads,
French locale numbers ("1 234,56"), report preambles ("Campaign report
(1 sept. 2026 - 7 sept. 2026)"), "Total:" trailer rows, currency suffixes in
headers ("Amount spent (EUR)") and Google Ads API field paths from MCP rows
(``metrics.cost_micros``).
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from navin.ads.models import AdRow, Platform

# Canonical column -> header aliases (normalized: lowercase, no parenthesis
# groups, punctuation collapsed). English, French and API field names.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "account": (
        "account", "account name", "customer", "customer name", "customer id", "account id",
        "advertiser", "advertiser name", "advertiser id", "ad account", "ad account name",
        "ad account id", "compte", "nom du compte", "annonceur", "customer.descriptive_name",
        "customer.id", "account_name", "account_id",
    ),
    "campaign": (
        "campaign", "campaign name", "campagne", "nom de la campagne", "nom de campagne",
        "campaign.name", "campaign_name", "campaign id name",
    ),
    "campaign_id": ("campaign id", "id de campagne", "campaign.id", "campaign_id"),
    "ad_group": (
        "ad group", "ad group name", "adgroup", "adgroup name", "ad set", "ad set name",
        "adset name", "adset", "groupe d annonces", "nom du groupe d annonces",
        "ensemble de publicites", "nom de l ensemble de publicites", "ad_group.name",
        "ad_group_name", "adgroup_name", "campaign group name", "campaign group",
        "nom du groupe de campagnes",
    ),
    "ad": (
        "ad", "ad name", "ad id", "creative", "creative name", "headline", "annonce",
        "nom de la publicite", "nom de l annonce", "ad_name", "ad_id", "ad_group_ad.ad.name",
        "ad_group_ad.ad.id", "creative id", "creative name",
    ),
    "keyword": (
        "keyword", "search keyword", "keyword text", "mot cle", "mot-cle", "mots cles",
        "ad_group_criterion.keyword.text", "keyword_text",
    ),
    "match_type": (
        "match type", "keyword match type", "type de correspondance",
        "ad_group_criterion.keyword.match_type", "match_type",
    ),
    "search_term": (
        "search term", "search query", "query", "terme de recherche", "requete de recherche",
        "search_term_view.search_term", "search_term", "search term text",
    ),
    "date": (
        "date", "day", "jour", "reporting starts", "reporting_starts", "date start",
        "start date", "debut", "segments.date", "week", "semaine", "time period",
        "stat_time_day",
    ),
    "status": (
        "status", "campaign status", "ad group status", "ad set delivery", "delivery status",
        "delivery", "etat", "statut", "campaign.status", "ad_group.status",
        "ad_group_criterion.status", "campaign state", "ad group state",
    ),
    "impressions": (
        "impressions", "impr", "imprs", "impression", "metrics.impressions", "affichages",
        "impressions totales", "total impressions",
    ),
    "clicks": (
        "clicks", "clics", "link clicks", "clics sur un lien", "metrics.clicks", "clicks all",
        "clics tous", "total clicks",
    ),
    "cost": (
        "cost", "cout", "spend", "amount spent", "montant depense", "total spent", "total spend",
        "depenses", "metrics.cost", "metrics.cost_micros", "cost_micros", "spent", "budget spent",
        "amount spent total",
    ),
    "conversions": (
        "conversions", "conv", "all conv", "all conversions", "results", "resultats",
        "purchases", "achats", "leads", "prospects", "metrics.conversions", "conversions all",
        "total conversions", "website conversions", "conversions site web", "complete payment",
        "external website conversions", "conversion",
    ),
    "conversion_value": (
        "conv value", "conversion value", "conversions value", "all conv value", "total conv value",
        "valeur de conv", "valeur de conversion", "valeur des conversions", "revenue", "revenu",
        "purchases conversion value", "purchase conversion value", "valeur de conversion des achats",
        "total purchase value", "metrics.conversions_value", "conversions_value",
        "website purchases conversion value", "purchase value", "conversion_value",
        "total conversion value", "results value",
    ),
    "quality_score": (
        "quality score", "quality score qs", "niveau de qualite", "qs",
        "ad_group_criterion.quality_info.quality_score", "quality_score", "qualite",
    ),
    "impression_share": (
        "search impr share", "search impression share", "impression share", "impr share",
        "taux d impressions sur le reseau de recherche", "taux d impressions",
        "metrics.search_impression_share", "search_impression_share",
    ),
    "frequency": ("frequency", "frequence", "repetition"),
    "reach": ("reach", "couverture", "portee", "unique reach"),
    "video_views": ("video views", "vues de video", "vues", "views", "video plays", "3-second video plays"),
    "daily_budget": (
        "budget", "daily budget", "budget quotidien", "campaign budget", "ad set budget",
        "campaign_budget.amount_micros", "budget amount", "budget total", "lifetime budget",
    ),
    "cpc": ("avg cpc", "cpc", "cost per click", "cpc moyen", "cout par clic", "metrics.average_cpc"),
    "ctr": ("ctr", "click-through rate", "taux de clics", "metrics.ctr", "ctr all"),
    "cpa": ("cost conv", "cost per conversion", "cost per result", "cout par conversion", "cout par resultat", "cpa"),
    "roas": ("roas", "purchase roas", "conv value cost", "return on ad spend"),
}

# When an export carries two columns for one metric, these aliases win.
_PREFERRED_ALIASES: dict[str, frozenset[str]] = {
    "clicks": frozenset({"link clicks", "clics sur un lien"}),
    "cost": frozenset({"amount spent", "montant depense", "cost", "spend", "total spent"}),
    "date": frozenset({"day", "date", "jour", "reporting starts"}),
}

_METRIC_COLUMNS = frozenset({"impressions", "clicks", "cost", "conversions", "conversion_value"})
_ENTITY_COLUMNS = frozenset({"account", "campaign", "ad_group", "ad", "keyword", "search_term"})
_EXTRA_COLUMNS = frozenset({
    "quality_score", "impression_share", "frequency", "reach", "video_views", "daily_budget",
})
_DERIVED_COLUMNS = frozenset({"cpc", "ctr", "cpa", "roas", "campaign_id"})

# Header vocabulary that points at one platform. Filenames add hints.
_PLATFORM_SIGNATURES: dict[Platform, tuple[str, ...]] = {
    Platform.META: (
        "ad set name", "ad set", "amount spent", "reporting starts", "results", "link clicks",
        "purchases conversion value", "purchase roas", "ensemble de publicites", "montant depense",
        "cost per result", "ad set delivery",
    ),
    Platform.LINKEDIN: (
        "campaign group name", "campaign group", "total spent", "leads", "member",
        "nom du groupe de campagnes",
    ),
    Platform.TIKTOK: (
        "ad group name", "total purchase value", "complete payment", "cpc destination",
        "video views", "6-second video views", "2-second video views", "cost per result",
        "advertiser name", "advertiser id",
    ),
    Platform.MICROSOFT: (
        "spend", "quality score", "account name", "revenue", "campaign name", "ad group",
        "delivered match type", "campaign status", "ad distribution", "top vs other",
    ),
    Platform.GOOGLE: (
        "impr", "conv value", "avg cpc", "search impr share", "campaign", "ad group",
        "search term", "search keyword", "cost", "conversions", "all conv", "quality score",
        "campaign type", "currency code", "metrics.cost_micros", "campaign.name",
    ),
    Platform.REDDIT: ("subreddit", "community", "ad group", "spend", "cpv", "reddit"),
}

_FILENAME_HINTS: dict[Platform, tuple[str, ...]] = {
    Platform.GOOGLE: ("google", "gads", "adwords"),
    Platform.MICROSOFT: ("microsoft", "bing", "msads"),
    Platform.META: ("meta", "facebook", "instagram", "fb-"),
    Platform.LINKEDIN: ("linkedin",),
    Platform.TIKTOK: ("tiktok",),
    Platform.REDDIT: ("reddit",),
}

_TOTAL_ROW_RE = re.compile(r"^(total|totaux|totals?)\b", re.IGNORECASE)
_PAREN_RE = re.compile(r"\([^)]*\)")
_NON_WORD_RE = re.compile(r"[^a-z0-9._]+")
_MULTI_SPACE_RE = re.compile(r"\s+")
_NUMBER_JUNK_RE = re.compile(r"[^0-9,.\-]")


@dataclass(slots=True)
class IngestResult:
    rows: list[AdRow]
    platform: Platform
    columns: dict[str, str] = field(default_factory=dict)
    unknown_columns: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    source: str = ""

    @property
    def data_gaps(self) -> list[dict[str, str]]:
        gaps: list[dict[str, str]] = []
        missing = [name for name in ("cost", "clicks", "impressions") if name not in self.columns]
        if missing:
            gaps.append({
                "service": f"export:{self.source or 'rows'}",
                "reason": "Missing core columns: " + ", ".join(missing),
            })
        if "conversions" not in self.columns:
            gaps.append({
                "service": f"export:{self.source or 'rows'}",
                "reason": "No conversions column: CPA, CVR and ROAS cannot be judged",
            })
        elif "conversion_value" not in self.columns:
            gaps.append({
                "service": f"export:{self.source or 'rows'}",
                "reason": "No conversion value column: ROAS cannot be judged",
            })
        return gaps


def normalize_header(header: str) -> str:
    """Lowercase, strip accents / parenthesis groups / punctuation."""
    text = _strip_accents(str(header or "")).lower().strip()
    text = _PAREN_RE.sub(" ", text)
    text = text.replace("'", " ").replace("\u2019", " ").replace("/", " ")
    text = _NON_WORD_RE.sub(" ", text)
    text = _MULTI_SPACE_RE.sub(" ", text).strip()
    text = text.rstrip(".").strip()
    if text.endswith(" ."):
        text = text[:-2].strip()
    return text


def _strip_accents(text: str) -> str:
    import unicodedata

    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


_ALIAS_INDEX: dict[str, str] = {}
for _canonical, _aliases in COLUMN_ALIASES.items():
    for _alias in _aliases:
        _ALIAS_INDEX.setdefault(normalize_header(_alias), _canonical)


def canonical_column(header: str) -> str | None:
    key = normalize_header(header)
    if not key:
        return None
    if key in _ALIAS_INDEX:
        return _ALIAS_INDEX[key]
    # "Impr." -> "impr", "Conv. value" -> "conv value": drop trailing dots per token.
    compact = " ".join(token.rstrip(".") for token in key.split())
    if compact in _ALIAS_INDEX:
        return _ALIAS_INDEX[compact]
    # Currency / unit suffix without parenthesis: "Cost EUR", "Spend USD".
    tokens = compact.split()
    if len(tokens) >= 2 and tokens[-1] in {"eur", "usd", "gbp", "chf", "cad", "mad", "aed", "sar"}:
        return _ALIAS_INDEX.get(" ".join(tokens[:-1]))
    return None


def parse_number(raw: Any) -> float | None:
    """Parse export numbers: '1,234.5', '1 234,56', '12.5%', '<10%', '--', '€12'."""
    if raw is None:
        return None
    if isinstance(raw, bool):
        return float(raw)
    if isinstance(raw, int | float):
        return float(raw)
    text = str(raw).strip()
    if not text or text in {"--", "-", "n/a", "na", "none", "null"}:
        return None
    percent = text.endswith("%")
    text = text.replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    text = text.lstrip("<>~≈")
    text = _NUMBER_JUNK_RE.sub("", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        # Last separator wins as the decimal mark.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", text):
            text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
    try:
        value = float(text)
    except ValueError:
        return None
    return value / 100 if percent else value


def detect_platform(headers: list[str], *, filename: str = "") -> Platform:
    normalized = {normalize_header(header) for header in headers}
    joined = " ".join(sorted(normalized))
    scores: dict[Platform, float] = {}
    for platform, signatures in _PLATFORM_SIGNATURES.items():
        score = 0.0
        for signature in signatures:
            sig = normalize_header(signature)
            if sig in normalized:
                score += 2
            elif sig and sig in joined:
                score += 0.5
        scores[platform] = score
    name = Path(filename).name.lower() if filename else ""
    for platform, hints in _FILENAME_HINTS.items():
        # A filename naming the platform beats generic header vocabulary.
        if any(hint in name for hint in hints):
            scores[platform] = scores.get(platform, 0) + 12
    if any(header.startswith(("metrics.", "campaign.", "ad_group.", "segments.")) for header in normalized):
        scores[Platform.GOOGLE] = scores.get(Platform.GOOGLE, 0) + 6
    # Strong disambiguators between the search platforms.
    if "spend" in normalized and "impr" not in normalized and "cost" not in normalized:
        scores[Platform.GOOGLE] = scores.get(Platform.GOOGLE, 0) - 3
    if "impr" in normalized or "conv value" in normalized or "avg cpc" in normalized:
        scores[Platform.MICROSOFT] = scores.get(Platform.MICROSOFT, 0) - 3
    best = max(scores.items(), key=lambda item: (item[1], -list(_PLATFORM_SIGNATURES).index(item[0])))
    return best[0] if best[1] >= 2 else Platform.UNKNOWN


def _decode_bytes(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    # Google Ads UTF-16LE without BOM shows as NUL-interleaved ASCII.
    if len(data) >= 4 and data[1:2] == b"\x00" and data[3:4] == b"\x00":
        try:
            return data.decode("utf-16-le")
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _sniff_delimiter(sample: str) -> str:
    candidates = [",", ";", "\t", "|"]
    lines = [line for line in sample.splitlines() if line.strip()][:20]
    if not lines:
        return ","
    best, best_score = ",", -1.0
    for candidate in candidates:
        counts = [line.count(candidate) for line in lines]
        if not counts or max(counts) == 0:
            continue
        # Prefer the delimiter that is present on most lines with a stable count.
        present = sum(1 for count in counts if count > 0)
        consistency = 1.0 / (1 + (max(counts) - min(count for count in counts if count > 0)))
        score = present * 10 + consistency + max(counts) * 0.01
        if score > best_score:
            best, best_score = candidate, score
    return best


def read_table_text(text: str) -> list[list[str]]:
    text = text.lstrip("\ufeff")
    delimiter = _sniff_delimiter(text[:20000])
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[cell.strip() for cell in row] for row in reader]


def read_table_file(path: str | Path) -> list[list[str]]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".xls":
        raise ValueError(f"{source.name}: legacy .xls is not supported, export as .xlsx or .csv")
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(source)
    return read_table_text(_decode_bytes(source.read_bytes()))


def _read_xlsx(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ValueError("openpyxl is required to read XLSX exports") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows: list[list[str]] = []
        for values in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else _cell_text(value) for value in values]
            if any(cell for cell in cells):
                rows.append(cells)
        return rows
    finally:
        workbook.close()


def _cell_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        try:
            return value.date().isoformat() if hasattr(value, "date") else value.isoformat()
        except Exception:  # noqa: BLE001
            return str(value)
    return str(value).strip()


def find_header_row(rows: list[list[str]]) -> int:
    """Index of the first row that reads like a metrics header (skip preambles)."""
    best_index, best_score = -1, 0
    for index, row in enumerate(rows[:30]):
        canon = {canonical_column(cell) for cell in row if cell}
        canon.discard(None)
        metrics = len(canon & (_METRIC_COLUMNS | _DERIVED_COLUMNS))
        entities = len(canon & _ENTITY_COLUMNS)
        score = metrics + entities * 2
        if metrics >= 1 and entities >= 1 and score > best_score:
            best_index, best_score = index, score
            if score >= 6:
                break
    return best_index


def parse_table(
    rows: list[list[str]],
    *,
    platform: Platform | str | None = None,
    source: str = "",
) -> IngestResult:
    header_index = find_header_row(rows)
    if header_index < 0:
        raise ValueError(
            "No recognizable export header (need a campaign / ad group column plus "
            "impressions, clicks or cost)"
        )
    headers = rows[header_index]
    detected = _coerce_platform(platform) or detect_platform(headers, filename=source)
    columns: dict[str, str] = {}
    positions: dict[str, int] = {}
    unknown: list[str] = []
    for index, header in enumerate(headers):
        if not header:
            continue
        canonical = canonical_column(header)
        if canonical is None:
            unknown.append(header)
            continue
        if canonical in positions:
            # First match wins unless a preferred alias shows up later
            # ("Link clicks" beats "Clicks (all)" on Meta exports).
            if normalize_header(header) in _PREFERRED_ALIASES.get(canonical, ()) and (
                normalize_header(columns[canonical]) not in _PREFERRED_ALIASES.get(canonical, ())
            ):
                positions[canonical] = index
                columns[canonical] = header
            continue
        positions[canonical] = index
        columns[canonical] = header
    warnings: list[str] = []
    parsed: list[AdRow] = []
    for raw in rows[header_index + 1:]:
        if not any(cell for cell in raw):
            continue
        first = next((cell for cell in raw if cell), "")
        if _TOTAL_ROW_RE.match(first):
            continue
        record = {name: (raw[idx] if idx < len(raw) else "") for name, idx in positions.items()}
        row = _row_from_record(record, detected, source, columns)
        if row is None:
            continue
        parsed.append(row)
    if not parsed:
        warnings.append("Header recognized but no data rows were parsed")
    return IngestResult(
        rows=parsed, platform=detected, columns=columns, unknown_columns=unknown,
        warnings=warnings, source=source,
    )


def parse_records(
    records: list[dict[str, Any]],
    *,
    platform: Platform | str | None = None,
    source: str = "mcp",
) -> IngestResult:
    """Normalize JSON rows (MCP tool output, API dumps). Nested keys are flattened."""
    flat = [_flatten(record) for record in records if isinstance(record, dict)]
    if not flat:
        raise ValueError("rows must be a non-empty list of objects")
    headers: list[str] = []
    for record in flat:
        for key in record:
            if key not in headers:
                headers.append(key)
    table = [headers] + [[_cell_text(record.get(key, "")) if record.get(key) is not None else "" for key in headers] for record in flat]
    return parse_table(table, platform=platform, source=source)


def ingest_file(path: str | Path, *, platform: Platform | str | None = None) -> IngestResult:
    source = Path(path)
    return parse_table(read_table_file(source), platform=platform, source=source.name)


def ingest_text(text: str, *, platform: Platform | str | None = None, source: str = "inline") -> IngestResult:
    return parse_table(read_table_text(text), platform=platform, source=source)


def _coerce_platform(value: Platform | str | None) -> Platform | None:
    if value is None or value == "":
        return None
    if isinstance(value, Platform):
        return value
    text = str(value).strip().lower()
    aliases = {
        "google": Platform.GOOGLE, "google ads": Platform.GOOGLE, "gads": Platform.GOOGLE,
        "adwords": Platform.GOOGLE, "microsoft": Platform.MICROSOFT, "bing": Platform.MICROSOFT,
        "microsoft ads": Platform.MICROSOFT, "msads": Platform.MICROSOFT, "meta": Platform.META,
        "facebook": Platform.META, "instagram": Platform.META, "meta ads": Platform.META,
        "linkedin": Platform.LINKEDIN, "linkedin ads": Platform.LINKEDIN, "tiktok": Platform.TIKTOK,
        "tiktok ads": Platform.TIKTOK, "reddit": Platform.REDDIT, "reddit ads": Platform.REDDIT,
        "auto": None, "unknown": None,
    }
    if text in aliases:
        return aliases[text]
    raise ValueError(f"unknown platform: {value}")


def _flatten(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in record.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict):
            out.update(_flatten(value, name + "."))
        else:
            out[name] = value
    return out


def _row_from_record(
    record: dict[str, str],
    platform: Platform,
    source: str,
    columns: dict[str, str],
) -> AdRow | None:
    text = {
        name: record.get(name, "").strip()
        for name in ("account", "campaign", "ad_group", "ad", "keyword", "match_type", "search_term", "date", "status")
    }
    numbers: dict[str, float] = {}
    for name in _METRIC_COLUMNS | _EXTRA_COLUMNS:
        if name not in record:
            continue
        value = parse_number(record[name])
        if value is None:
            continue
        numbers[name] = value
    if "cost" in numbers and "cost" in columns and "micros" in normalize_header(columns["cost"]):
        numbers["cost"] = numbers["cost"] / 1_000_000
    if "daily_budget" in numbers and "micros" in normalize_header(columns.get("daily_budget", "")):
        numbers["daily_budget"] = numbers["daily_budget"] / 1_000_000
    if "impression_share" in numbers and numbers["impression_share"] > 1:
        numbers["impression_share"] = numbers["impression_share"] / 100
    has_entity = any(text[name] for name in ("account", "campaign", "ad_group", "ad", "keyword", "search_term"))
    has_metric = any(name in numbers for name in _METRIC_COLUMNS)
    if not has_entity and not has_metric:
        return None
    if not has_entity:
        # Metric-only lines are trailers ("Total: Account") in disguise.
        return None
    extras = {name: numbers[name] for name in _EXTRA_COLUMNS if name in numbers}
    return AdRow(
        platform=platform,
        account=text["account"],
        campaign=text["campaign"],
        ad_group=text["ad_group"],
        ad=text["ad"],
        keyword=text["keyword"],
        match_type=text["match_type"].lower(),
        search_term=text["search_term"],
        date=_normalize_date(text["date"]),
        status=text["status"].lower(),
        impressions=numbers.get("impressions", 0.0),
        clicks=numbers.get("clicks", 0.0),
        cost=numbers.get("cost", 0.0),
        conversions=numbers.get("conversions", 0.0),
        conversion_value=numbers.get("conversion_value", 0.0),
        extras=extras,
        source=source,
    )


_DATE_PATTERNS = (
    (re.compile(r"^(\d{4})-(\d{2})-(\d{2})"), lambda m: f"{m.group(1)}-{m.group(2)}-{m.group(3)}"),
    (re.compile(r"^(\d{4})/(\d{2})/(\d{2})"), lambda m: f"{m.group(1)}-{m.group(2)}-{m.group(3)}"),
    (re.compile(r"^(\d{2})/(\d{2})/(\d{4})"), lambda m: f"{m.group(3)}-{m.group(2)}-{m.group(1)}"),
    (re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})"), lambda m: f"{m.group(3)}-{m.group(2)}-{m.group(1)}"),
    (re.compile(r"^(\d{8})$"), lambda m: f"{m.group(1)[:4]}-{m.group(1)[4:6]}-{m.group(1)[6:]}"),
)


def _normalize_date(raw: str) -> str:
    text = raw.strip()
    if not text:
        return ""
    for pattern, build in _DATE_PATTERNS:
        match = pattern.match(text)
        if match:
            return build(match)
    return text
