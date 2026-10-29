# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only
"""Tests for the docx, xlsx, pptx and pdf built-in skills.

Each skill ships a SKILL.md and a scripts/ folder. The tests check the
frontmatter and the house rules, then run every script on small fixtures
built in tmp_path (no network, no LibreOffice required).
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SKILLS = ROOT / "navin" / "skills"
DOC_SKILLS = ("docx", "xlsx", "pptx", "pdf")
TEXT_SUFFIXES = {".md", ".py", ".txt", ".json", ".html", ".css", ".js"}
LONG_DASHES = re.compile("[\u2013\u2014]")
FRONTMATTER = re.compile(r"^---\s*\n(.*?)\n---\s*\n", re.DOTALL)
PY = sys.executable


def _load_validator():
    path = SKILLS / "skill-development" / "scripts" / "validate_skill.py"
    spec = importlib.util.spec_from_file_location("validate_skill", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.dont_write_bytecode = True
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = False
    return module


def _run(*args: str, expect: int | None = 0) -> subprocess.CompletedProcess:
    proc = subprocess.run(
        [PY, *args],
        capture_output=True,
        text=True,
        timeout=60,
        env={"PYTHONDONTWRITEBYTECODE": "1", "PATH": "/usr/bin:/bin"},
    )
    if expect is not None:
        assert proc.returncode == expect, f"{args}\nstdout: {proc.stdout}\nstderr: {proc.stderr}"
    return proc


def _script(skill: str, name: str) -> str:
    return str(SKILLS / skill / "scripts" / name)


def _json_out(proc: subprocess.CompletedProcess) -> dict:
    return json.loads(proc.stdout)


# ---------------------------------------------------------------- fixtures


@pytest.fixture()
def sample_docx(tmp_path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Contract", 1)
    doc.add_paragraph("Payment due in 30 days after invoice.")
    doc.add_paragraph("Second paragraph.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell A"
    path = tmp_path / "in.docx"
    doc.save(path)
    return path


@pytest.fixture()
def sample_xlsx(tmp_path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Qty", "Price", "Total"])
    ws.append(["A", 2, 3.5, "=B2*C2"])
    ws.append(["B", 0, 4, "=C3/B3"])
    ws["F1"] = "=Missing!A1"
    ws["G1"] = "#REF!"
    path = tmp_path / "in.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def sample_pdf(tmp_path: Path) -> Path:
    from reportlab.pdfgen import canvas

    path = tmp_path / "in.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(100, 700, "Hello PDF world")
    c.acroForm.textfield(name="Name", x=100, y=600, width=200, height=20)
    c.acroForm.checkbox(name="Agree", x=100, y=560, buttonStyle="check")
    c.showPage()
    c.drawString(100, 700, "Page two")
    c.save()
    return path


@pytest.fixture()
def sample_deck(tmp_path: Path) -> Path:
    spec = _json_out(_run(_script("pptx", "pptx_build.py"), "--spec-example"))
    spec["slides"] = [s for s in spec["slides"] if s["type"] != "image"]
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps(spec), encoding="utf-8")
    out = tmp_path / "deck.pptx"
    report = _json_out(_run(_script("pptx", "pptx_build.py"), "--spec", str(spec_path), "--out", str(out), "--json"))
    assert report["ok"] and report["slides"] == len(spec["slides"])
    return out


# ------------------------------------------------------------ house rules


@pytest.mark.parametrize("skill", DOC_SKILLS)
def test_frontmatter_and_lint(skill: str):
    folder = SKILLS / skill
    text = (folder / "SKILL.md").read_text(encoding="utf-8")
    match = FRONTMATTER.match(text)
    assert match, f"{skill}: missing frontmatter"
    import yaml

    meta = yaml.safe_load(match.group(1))
    assert meta["name"] == skill
    assert 150 <= len(meta["description"]) <= 1024
    assert meta["metadata"]["navin"]["category"] == "documents"
    assert "## References" in text and "## Workflow" in text
    report = _load_validator().lint(folder)
    assert report["passed"], report["critical"]
    assert report["stats"]["trigger_phrases"] >= 3


@pytest.mark.parametrize("skill", DOC_SKILLS)
def test_no_long_dashes_or_upstream_names(skill: str):
    for path in (SKILLS / skill).rglob("*"):
        if not path.is_file() or path.suffix not in TEXT_SUFFIXES or "__pycache__" in path.parts:
            continue
        text = path.read_text(encoding="utf-8")
        assert not LONG_DASHES.search(text), f"long dash in {path}"
        assert not re.search(r"claude|anthropic", text, re.IGNORECASE), f"upstream name in {path}"


@pytest.mark.parametrize("skill", DOC_SKILLS)
def test_scripts_are_executable_and_have_help(skill: str):
    scripts = sorted((SKILLS / skill / "scripts").glob("*.py"))
    assert scripts
    for script in scripts:
        assert script.stat().st_mode & 0o111, f"{script} is not executable"
        proc = _run(str(script), "--help")
        assert "usage" in proc.stdout.lower()


def test_shared_office_helpers_are_byte_identical():
    for name in ("office_pkg.py", "office_convert.py"):
        digests = {hashlib.sha256((SKILLS / s / "scripts" / name).read_bytes()).hexdigest() for s in ("docx", "xlsx", "pptx")}
        assert len(digests) == 1, f"{name} differs between docx, xlsx and pptx"


# ------------------------------------------------------------------- docx


def test_docx_text_dump_and_replace(sample_docx: Path, tmp_path: Path):
    rows = _json_out(_run(_script("docx", "docx_text.py"), "dump", str(sample_docx), "--format", "json"))
    texts = [p["text"] for p in rows]
    assert "Payment due in 30 days after invoice." in texts
    assert any(p.get("in_table") for p in rows)
    out = tmp_path / "replaced.docx"
    report = _json_out(
        _run(_script("docx", "docx_text.py"), "replace", str(sample_docx), "--find", "30 days", "--with", "45 days", "--out", str(out), "--json")
    )
    assert report["replacements"] == 1
    assert "45 days" in _run(_script("docx", "docx_text.py"), "dump", str(out), "--format", "markdown").stdout
    missing = _run(_script("docx", "docx_text.py"), "replace", str(sample_docx), "--find", "absent text", "--with", "x", "--out", str(tmp_path / "m.docx"), expect=1)
    assert missing.returncode == 1


def test_docx_tracked_changes_round_trip(sample_docx: Path, tmp_path: Path):
    tracked = tmp_path / "tracked.docx"
    _run(_script("docx", "docx_track.py"), "replace", str(sample_docx), "--find", "30 days", "--with", "45 days", "--author", "Navin", "--out", str(tracked))
    counts = _json_out(_run(_script("docx", "docx_track.py"), "count", str(tracked), "--json"))
    assert counts["ins"] == 1 and counts["del"] == 1
    xml = zipfile.ZipFile(tracked).read("word/document.xml").decode("utf-8")
    assert "<w:ins " in xml and "<w:del " in xml and "w:delText" in xml

    accepted = tmp_path / "accepted.docx"
    _run(_script("docx", "docx_track.py"), "accept", str(tracked), "--out", str(accepted))
    text = _run(_script("docx", "docx_text.py"), "dump", str(accepted), "--format", "markdown").stdout
    assert "45 days" in text and "30 days" not in text
    assert _json_out(_run(_script("docx", "docx_track.py"), "count", str(accepted), "--json"))["ins"] == 0

    rejected = tmp_path / "rejected.docx"
    _run(_script("docx", "docx_track.py"), "reject", str(tracked), "--out", str(rejected))
    text = _run(_script("docx", "docx_text.py"), "dump", str(rejected), "--format", "markdown").stdout
    assert "30 days" in text and "45 days" not in text


def test_docx_comments(sample_docx: Path, tmp_path: Path):
    # python-docx >= 1.2 ships Document.add_comment; the project pins it, so the
    # script must work here rather than be skipped.
    out = tmp_path / "commented.docx"
    report = _json_out(
        _run(_script("docx", "docx_comments.py"), "add", str(sample_docx), "--anchor", "Payment", "--text", "Check with finance", "--author", "Navin", "--out", str(out))
    )
    assert report["ok"]
    listing = _run(_script("docx", "docx_comments.py"), "list", str(out)).stdout
    assert "Check with finance" in listing and "Navin" in listing


def test_office_pkg_unpack_pack_validate(sample_docx: Path, tmp_path: Path):
    work = tmp_path / "work"
    report = _json_out(_run(_script("docx", "office_pkg.py"), "unpack", str(sample_docx), str(work)))
    assert report["ok"] and report["malformed_xml"] == []
    assert (work / "word" / "document.xml").is_file()
    packed = tmp_path / "packed.docx"
    _run(_script("docx", "office_pkg.py"), "pack", str(work), str(packed))
    with zipfile.ZipFile(packed) as zf:
        assert zf.namelist()[0] == "[Content_Types].xml"
    report = _json_out(_run(_script("docx", "office_pkg.py"), "validate", str(packed), "--json"))
    assert report["ok"] and report["errors"] == []
    from docx import Document

    assert "Payment due in 30 days after invoice." in [p.text for p in Document(str(packed)).paragraphs]

    # A dangling relationship target must be refused by pack and reported by validate.
    rels = work / "word" / "_rels" / "document.xml.rels"
    dangling = '<Relationship Id="rIdZ" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/missing.png"/>'
    rels.write_text(rels.read_text(encoding="utf-8").replace("</Relationships>", dangling + "</Relationships>"), encoding="utf-8")
    broken = tmp_path / "broken.docx"
    report = _json_out(_run(_script("docx", "office_pkg.py"), "pack", str(work), str(broken), expect=1))
    assert not report["ok"] and any("missing.png" in e for e in report["errors"])
    with zipfile.ZipFile(sample_docx) as src, zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            payload = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                payload = payload.replace(b"</Relationships>", dangling.encode() + b"</Relationships>")
            dst.writestr(item, payload)
    report = _json_out(_run(_script("docx", "office_pkg.py"), "validate", str(broken), "--json", expect=1))
    assert not report["ok"] and any("missing.png" in e for e in report["errors"])

    # Malformed XML inside a part: unpack keeps the raw bytes and reports it (exit 1).
    malformed = tmp_path / "malformed.docx"
    with zipfile.ZipFile(sample_docx) as src, zipfile.ZipFile(malformed, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            payload = src.read(item.filename)
            if item.filename == "word/document.xml":
                payload = payload[:-10]
            dst.writestr(item, payload)
    report = _json_out(_run(_script("docx", "office_pkg.py"), "unpack", str(malformed), str(tmp_path / "work2"), expect=1))
    assert not report["ok"] and any(e.startswith("word/document.xml") for e in report["malformed_xml"])
    assert (tmp_path / "work2" / "word" / "document.xml").is_file()


def test_office_convert_reports_missing_soffice(sample_docx: Path):
    proc = _run(_script("docx", "office_convert.py"), str(sample_docx), "--to", "pdf", "--json", expect=None)
    if proc.returncode == 0:
        assert Path(_json_out(proc)["output"]).suffix == ".pdf"
    else:
        assert proc.returncode == 3
        assert "soffice" in (proc.stdout + proc.stderr).lower()


# ------------------------------------------------------------------- xlsx


def test_xlsx_inspect_finds_formulas_and_errors(sample_xlsx: Path):
    report = _json_out(_run(_script("xlsx", "xlsx_inspect.py"), str(sample_xlsx)))
    sheet = report["sheets"][0]
    assert sheet["name"] == "Data" and sheet["formulas"] == 3 and sheet["header_row"] == 1
    assert any(c["cell"] == "G1" for c in sheet["error_cells"])
    assert "Data" in _run(_script("xlsx", "xlsx_inspect.py"), str(sample_xlsx), "--markdown").stdout


def test_xlsx_recalc_static_flags_missing_sheet(sample_xlsx: Path, tmp_path: Path):
    out = tmp_path / "rc.xlsx"
    proc = _run(_script("xlsx", "xlsx_recalc.py"), str(sample_xlsx), "--static", "--out", str(out), expect=1)
    report = _json_out(proc)
    assert report["recalculated"] is False and report["ok"] is False
    kinds = {(c["cell"], c["kind"]) for c in report["error_cells"]}
    assert ("F1", "missing_sheet") in kinds and ("G1", "literal_error") in kinds


def test_xlsx_style_applies_header_fill(sample_xlsx: Path, tmp_path: Path):
    out = tmp_path / "styled.xlsx"
    report = _json_out(_run(_script("xlsx", "xlsx_style.py"), str(sample_xlsx), "--out", str(out), "--number-format", "C=#,##0.00"))
    assert report["ok"] and report["cells_styled"] > 0
    from openpyxl import load_workbook

    ws = load_workbook(out)["Data"]
    assert ws["A1"].fill.fgColor.rgb.lower().endswith("17171a")
    assert ws["A1"].font.bold and ws.freeze_panes
    assert ws["C2"].number_format == "#,##0.00"
    assert ws["D2"].value == "=B2*C2"


# ------------------------------------------------------------------- pptx


def test_pptx_build_inventory_replace(sample_deck: Path, tmp_path: Path):
    inv = _json_out(_run(_script("pptx", "pptx_inventory.py"), str(sample_deck), "--slide", "2"))
    assert inv["slide_count"] == 11
    shapes = inv["slides"][0]["shapes"]
    title = next(s for s in shapes if s["text"] == "Agenda")
    edits = tmp_path / "edits.json"
    edits.write_text(json.dumps([
        {"slide": 2, "shape_id": title["id"], "text": "Ordre du jour"},
        {"slide": 2, "shape_name": title["name"], "find": "jour", "replace": "JOUR"},
    ]), encoding="utf-8")
    out = tmp_path / "edited.pptx"
    _run(_script("pptx", "pptx_replace.py"), str(sample_deck), "--edits", str(edits), "--delete-slide", "11", "--out", str(out))
    inv2 = _json_out(_run(_script("pptx", "pptx_inventory.py"), str(out)))
    assert inv2["slide_count"] == 10
    texts = [s["text"] for s in inv2["slides"][1]["shapes"]]
    assert "Ordre du JOUR" in texts and "Agenda" not in texts


def test_pptx_thumbnails_outline_without_soffice(sample_deck: Path):
    proc = _run(_script("pptx", "pptx_thumbnails.py"), str(sample_deck), "--outline")
    assert "Quarterly review" in proc.stdout and "Agenda" in proc.stdout


def test_pptx_build_fails_on_missing_image(tmp_path: Path):
    spec = _json_out(_run(_script("pptx", "pptx_build.py"), "--spec-example"))
    spec["slides"] = [s for s in spec["slides"] if s["type"] == "image"]
    spec["slides"][0]["image"] = str(tmp_path / "nope.png")
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps(spec), encoding="utf-8")
    proc = _run(_script("pptx", "pptx_build.py"), "--spec", str(spec_path), "--out", str(tmp_path / "x.pptx"), expect=None)
    assert proc.returncode != 0
    assert "nope.png" in (proc.stdout + proc.stderr)


def test_pptx_package_validates(sample_deck: Path):
    report = _json_out(_run(_script("pptx", "office_pkg.py"), "validate", str(sample_deck), "--json"))
    assert report["ok"] and report["errors"] == []


# -------------------------------------------------------------------- pdf


def test_pdf_info_and_text(sample_pdf: Path):
    info = _json_out(_run(_script("pdf", "pdf_pages.py"), "info", str(sample_pdf)))
    assert info["pages"] == 2 and info["has_form"] is True and info["encrypted"] is False
    text = _json_out(_run(_script("pdf", "pdf_text.py"), str(sample_pdf), "--format", "json", "--tables"))
    pages = text["pages"]
    assert "Hello PDF world" in pages[0]["text"] and "Page two" in pages[1]["text"]
    md = _run(_script("pdf", "pdf_text.py"), str(sample_pdf), "--pages", "2", "--format", "markdown").stdout
    assert "Page two" in md and "Hello PDF" not in md


def test_pdf_forms_fields_fill_flatten(sample_pdf: Path, tmp_path: Path):
    fields = _json_out(_run(_script("pdf", "pdf_forms.py"), "fields", str(sample_pdf), "--json"))
    by_name = {f["name"]: f for f in fields}
    assert by_name["Name"]["type"] == "text"
    assert by_name["Agree"]["type"] == "checkbox" and "/Yes" in by_name["Agree"]["options"]
    data = tmp_path / "values.json"
    data.write_text(json.dumps({"Name": "Jane Doe", "Agree": "/Yes"}), encoding="utf-8")
    out = tmp_path / "filled.pdf"
    report = _json_out(_run(_script("pdf", "pdf_forms.py"), "fill", str(sample_pdf), "--data", str(data), "--out", str(out)))
    assert set(report["filled"]) == {"Name", "Agree"} and report["unknown_fields"] == []
    ghost = tmp_path / "ghost.json"
    ghost.write_text(json.dumps({"Name": "Jane Doe", "Ghost": "x"}), encoding="utf-8")
    report = _json_out(_run(_script("pdf", "pdf_forms.py"), "fill", str(sample_pdf), "--data", str(ghost), "--out", str(tmp_path / "g.pdf"), expect=1))
    assert report["unknown_fields"] == ["Ghost"] and report["filled"] == ["Name"]
    values = {f["name"]: f["value"] for f in _json_out(_run(_script("pdf", "pdf_forms.py"), "fields", str(out), "--json"))}
    assert values["Name"] == "Jane Doe" and values["Agree"] == "/Yes"
    flat = tmp_path / "flat.pdf"
    report = _json_out(_run(_script("pdf", "pdf_forms.py"), "fill", str(sample_pdf), "--data", str(data), "--out", str(flat), "--flatten"))
    assert report["flattened"] is True
    import fitz

    with fitz.open(flat) as doc:
        assert "Jane Doe" in doc[0].get_text()
        assert list(doc[0].widgets()) == []


def test_pdf_pages_split_merge_rotate_extract_render(sample_pdf: Path, tmp_path: Path):
    parts = tmp_path / "parts"
    report = _json_out(_run(_script("pdf", "pdf_pages.py"), "split", str(sample_pdf), "--every", "1", "--out-dir", str(parts)))
    files = [o["file"] for o in report["outputs"]]
    assert len(files) == 2
    merged = tmp_path / "merged.pdf"
    assert _json_out(_run(_script("pdf", "pdf_pages.py"), "merge", *files, "--out", str(merged)))["pages"] == 2
    rotated = tmp_path / "rotated.pdf"
    _run(_script("pdf", "pdf_pages.py"), "rotate", str(merged), "--degrees", "90", "--pages", "1", "--out", str(rotated))
    from pypdf import PdfReader

    reader = PdfReader(str(rotated))
    assert reader.pages[0].rotation == 90 and reader.pages[1].rotation == 0
    extracted = tmp_path / "p2.pdf"
    _run(_script("pdf", "pdf_pages.py"), "extract", str(merged), "--pages", "2", "--out", str(extracted))
    assert len(PdfReader(str(extracted)).pages) == 1
    png_dir = tmp_path / "png"
    report = _json_out(_run(_script("pdf", "pdf_pages.py"), "render", str(merged), "--dpi", "40", "--out-dir", str(png_dir)))
    assert len(report["outputs"]) == 2 and all(Path(o["file"]).stat().st_size > 0 for o in report["outputs"])


def test_pdf_overlay_text_and_grid(sample_pdf: Path, tmp_path: Path):
    out = tmp_path / "overlay.pdf"
    report = _json_out(_run(_script("pdf", "pdf_overlay.py"), str(sample_pdf), "--text", "Stamped by Navin", "--x", "100", "--y", "500", "--page", "2", "--out", str(out)))
    assert report["items_written"] == 1
    import fitz

    with fitz.open(out) as doc:
        assert "Stamped by Navin" in doc[1].get_text() and "Stamped by Navin" not in doc[0].get_text()
    grid_dir = tmp_path / "grid"
    report = _json_out(_run(_script("pdf", "pdf_overlay.py"), str(sample_pdf), "--grid", "--step", "100", "--dpi", "40", "--out-dir", str(grid_dir)))
    assert Path(report["output"]).is_file() and report["step"] == 100


# ---------------------------------------------------------------- routing


ROUTING_PROMPTS = [
    ("docx", "Ajoute des commentaires et des modifications suivies dans ce document Word"),
    ("docx", "Extract the text of report.docx and add tracked changes"),
    ("docx", "Lis ce fichier .docx et remplace les paragraphes de la section 2"),
    ("xlsx", "Crée un fichier Excel avec des formules et recalcule les totaux"),
    ("xlsx", "Inspect this spreadsheet .xlsx: sheets, formulas, errors"),
    ("xlsx", "Mets en forme le tableau Excel et ajoute un style aux en-têtes"),
    ("pptx", "Remplace les textes de cette présentation PowerPoint et génère les vignettes"),
    ("pptx", "Build a pptx deck from this JSON outline and inventory the slides"),
    ("pptx", "Fais l'inventaire des slides du .pptx et remplace le logo"),
    ("pdf", "Remplis ce formulaire PDF et aplatis les champs"),
    ("pdf", "Extract tables from this PDF and merge the pages"),
    ("pdf", "Fusionne ces deux PDF et ajoute un filigrane"),
]


@pytest.fixture(scope="module")
def skill_router():
    """The three paths the agent uses to pick a playbook, on the real catalog."""
    from navin.agent.skills import SkillsLoader
    from navin.agent.skills_graph import build_skills_graph
    from navin.agent.tools.skill_catalog import SkillCatalogTool

    loader = SkillsLoader(ROOT)
    return loader, build_skills_graph(loader), SkillCatalogTool(workspace=ROOT)


@pytest.mark.parametrize(("skill", "prompt"), ROUTING_PROMPTS, ids=[f"{s}-{i}" for i, (s, _p) in enumerate(ROUTING_PROMPTS)])
def test_document_requests_route_to_their_skill(skill_router, skill: str, prompt: str):
    """A French or English request lands on the document skill through `skill
    action=find`, the graph ranking and the unprompted per-turn suggestion."""
    loader, graph, tool = skill_router
    found = [row["name"] for row in loader.search_skills(prompt, limit=5)]
    ranked = [row["name"] for row in graph.suggest(prompt, limit=5)]
    suggested = [row["name"] for row in tool._suggest_for_request(prompt)]
    assert skill in found[:3], f"find: {found}"
    assert skill in ranked[:3], f"graph: {ranked}"
    assert skill in suggested, f"auto-suggest: {suggested}"
