# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""preview_document: real render, approximation, and honest refusals."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from navin.documents import _pdf, preview_document
from navin.documents._office import OfficeUnavailableError


def _text_pdf(path: Path, pages: int = 3) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    page = canvas.Canvas(str(path), pagesize=A4)
    page.setTitle("Report")
    for number in range(1, pages + 1):
        page.drawString(72, 800, f"Page {number}")
        page.showPage()
    page.save()


def _deck(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    for title in ("One", "Two"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = "Body copy"
    presentation.save(str(path))


@unittest.skipUnless(_pdf.renderer_available(), "no PDF renderer (PyMuPDF, pypdfium2 or pdftoppm)")
class PdfRenderTests(unittest.TestCase):
    def test_a_pdf_renders_to_pages_and_a_labelled_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "report.pdf"
            _text_pdf(pdf, pages=3)
            result = preview_document.render(pdf, Path(tmp) / "previews", max_edge=600)
            self.assertEqual(result["mode"], "real")
            self.assertEqual(result["total_pages"], 3)
            self.assertEqual(len(result["pages"]), 3)
            self.assertTrue(all(Path(p).is_file() for p in result["pages"]))
            self.assertTrue(Path(result["sheet"]).is_file())
            from PIL import Image

            with Image.open(result["pages"][0]) as image:
                self.assertLessEqual(max(image.size), 600)

    def test_max_pages_caps_the_render_but_reports_the_total(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "report.pdf"
            _text_pdf(pdf, pages=4)
            result = preview_document.render(pdf, Path(tmp) / "p", max_pages=2, max_edge=400, sheet=False)
        self.assertEqual(len(result["pages"]), 2)
        self.assertEqual(result["total_pages"], 4)
        self.assertNotIn("sheet", result)


class OfficeRenderTests(unittest.TestCase):
    def test_a_deck_without_libreoffice_is_approximated_and_says_so(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "deck.pptx"
            _deck(deck)
            with mock.patch.object(preview_document, "to_pdf", side_effect=OfficeUnavailableError("missing")):
                result = preview_document.render(deck, Path(tmp) / "previews")
            self.assertTrue(Path(result["sheet"]).is_file())
        self.assertEqual(result["mode"], "approximate")
        self.assertIn("LibreOffice is not installed", result["note"])
        self.assertEqual(len(result["pages"]), 2)

    def test_a_docx_without_libreoffice_is_refused_with_the_install_hint(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            docx = Path(tmp) / "report.docx"
            docx.write_bytes(b"PK")
            with mock.patch.object(preview_document, "to_pdf", side_effect=OfficeUnavailableError("missing")):
                with self.assertRaises(RuntimeError) as caught:
                    preview_document.render(docx, Path(tmp) / "previews")
        message = str(caught.exception)
        self.assertIn("cannot be rendered without LibreOffice", message)
        self.assertIn("doc_check", message)

    @unittest.skipUnless(_pdf.renderer_available(), "no PDF renderer")
    def test_with_libreoffice_the_office_file_goes_through_a_real_pdf(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            docx = Path(tmp) / "report.docx"
            docx.write_bytes(b"PK")

            def fake_to_pdf(source, outdir, soffice=None, timeout=0):
                outdir.mkdir(parents=True, exist_ok=True)
                pdf = outdir / "report.pdf"
                _text_pdf(pdf, pages=2)
                return pdf

            with mock.patch.object(preview_document, "to_pdf", side_effect=fake_to_pdf):
                result = preview_document.render(docx, Path(tmp) / "previews", max_edge=300)
        self.assertEqual(result["mode"], "real")
        self.assertEqual(result["total_pages"], 2)
        self.assertEqual(len(result["pages"]), 2)


class ErrorTests(unittest.TestCase):
    def test_unsupported_and_missing_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            other = Path(tmp) / "notes.txt"
            other.write_text("x", encoding="utf-8")
            with self.assertRaises(ValueError):
                preview_document.render(other, Path(tmp) / "p")
            with self.assertRaises(FileNotFoundError):
                preview_document.render(Path(tmp) / "nope.pdf", Path(tmp) / "p")
            self.assertEqual(preview_document.main([str(other), str(Path(tmp) / "p")]), 1)


if __name__ == "__main__":
    unittest.main()
