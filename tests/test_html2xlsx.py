"""Liveness contract for the HTML to XLSX converter."""

from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from navin.documents.html2xlsx import (
    SheetData,
    _day_first,
    _grid,
    build_workbook,
    export_csv,
    parse_value,
)


def _style(**overrides) -> dict:
    style = {
        "font": "Calibri",
        "families": ["Calibri", "Segoe UI"],
        "size": 16.0,
        "bold": False,
        "italic": False,
        "color": "1F2937",
        "fill": None,
        "align": "left",
        "valign": "middle",
        "wrap": False,
        "borders": None,
    }
    style.update(overrides)
    return style


def _cell(text: str, **overrides) -> dict:
    cell = {
        "text": text,
        "header": False,
        "colspan": 1,
        "rowspan": 1,
        "width": 140.0,
        "style": _style(),
    }
    cell.update(overrides)
    return cell


def _table(rows: list[dict]) -> dict:
    return {"kind": "table", "rows": rows, "top": 100.0}


def _sheet(blocks: list[dict], name: str = "Budget") -> SheetData:
    return SheetData(name=name, blocks=blocks)


def _build(sheets: list[SheetData], **kwargs):
    with tempfile.TemporaryDirectory() as raw:
        output = Path(raw) / "out.xlsx"
        build_workbook(sheets, output, **kwargs)
        return load_workbook(str(output))


class ValueTypingTests(unittest.TestCase):
    def test_currency_becomes_a_number_with_a_money_format(self) -> None:
        value, number_format = parse_value("€1,284,500", None)
        self.assertEqual(value, 1284500)
        self.assertIn("€", number_format)

    def test_french_notation_is_read_correctly(self) -> None:
        self.assertEqual(parse_value("1 284,50", None)[0], 1284.5)

    def test_percentage_is_stored_as_a_ratio(self) -> None:
        value, number_format = parse_value("38.4%", None)
        self.assertAlmostEqual(value, 0.384)
        self.assertIn("%", number_format)

    def test_parentheses_mean_negative(self) -> None:
        self.assertEqual(parse_value("(1,200)", None)[0], -1200)

    def test_iso_date_is_a_real_date(self) -> None:
        self.assertEqual(parse_value("2026-07-25", None)[0], date(2026, 7, 25))

    def test_ambiguous_date_stays_text_without_a_known_order(self) -> None:
        self.assertEqual(parse_value("08/01/2026", None)[0], "08/01/2026")

    def test_ambiguous_date_is_resolved_when_the_order_is_known(self) -> None:
        self.assertEqual(parse_value("08/01/2026", True)[0], date(2026, 1, 8))

    def test_plain_label_stays_text(self) -> None:
        self.assertEqual(parse_value("Consulting", None), ("Consulting", None))

    def test_date_order_inferred_from_an_unambiguous_neighbour(self) -> None:
        rows = [{"cells": [_cell("25/07/2026"), _cell("08/01/2026")]}]
        self.assertIs(_day_first(rows, "und"), True)

    def test_us_dates_are_month_first(self) -> None:
        rows = [{"cells": [_cell("07/25/2026")]}]
        self.assertIs(_day_first(rows, "en-us"), False)


class WorkbookTests(unittest.TestCase):
    def test_cells_hold_typed_values_not_strings(self) -> None:
        table = _table([
            {"cells": [_cell("Ligne", header=True), _cell("CA", header=True)], "head": True},
            {"cells": [_cell("Conseil"), _cell("€1,800")], "head": False},
        ])
        sheet = _build([_sheet([table])]).worksheets[0]
        self.assertEqual(sheet["A2"].value, "Conseil")
        self.assertEqual(sheet["B2"].value, 1800)

    def test_header_row_freezes_and_filters(self) -> None:
        table = _table([
            {"cells": [_cell("Produit", header=True), _cell("Total", header=True)], "head": True},
            {"cells": [_cell("Nova"), _cell("120")], "head": False},
        ])
        sheet = _build([_sheet([table])]).worksheets[0]
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:B2")

    def test_header_fill_and_font_color_are_kept(self) -> None:
        header = _cell("Produit", header=True, style=_style(fill="6D28D9", color="FFFFFF", bold=True))
        table = _table([{"cells": [header], "head": True}])
        cell = _build([_sheet([table])]).worksheets[0]["A1"]
        self.assertEqual(cell.fill.fgColor.rgb, "FF6D28D9")
        self.assertTrue(cell.font.bold)

    def test_total_row_becomes_a_sum_formula(self) -> None:
        rows = [
            {"cells": [_cell("Produit", header=True), _cell("CA", header=True)], "head": True},
            {"cells": [_cell("A"), _cell("100")], "head": False},
            {"cells": [_cell("B"), _cell("200")], "head": False},
            {"cells": [_cell("C"), _cell("300")], "head": False},
            {"cells": [_cell("Total"), _cell("600")], "head": False},
        ]
        sheet = _build([_sheet([_table(rows)])]).worksheets[0]
        self.assertEqual(sheet["B5"].value, "=SUM(B2:B4)")

    def test_wrong_total_stays_a_value(self) -> None:
        rows = [
            {"cells": [_cell("Produit", header=True), _cell("CA", header=True)], "head": True},
            {"cells": [_cell("A"), _cell("100")], "head": False},
            {"cells": [_cell("B"), _cell("200")], "head": False},
            {"cells": [_cell("C"), _cell("300")], "head": False},
            {"cells": [_cell("Total"), _cell("999")], "head": False},
        ]
        sheet = _build([_sheet([_table(rows)])]).worksheets[0]
        self.assertEqual(sheet["B5"].value, 999)

    def test_formulas_can_be_turned_off(self) -> None:
        rows = [
            {"cells": [_cell("Produit", header=True), _cell("CA", header=True)], "head": True},
            {"cells": [_cell("A"), _cell("100")], "head": False},
            {"cells": [_cell("B"), _cell("200")], "head": False},
            {"cells": [_cell("C"), _cell("300")], "head": False},
            {"cells": [_cell("Total"), _cell("600")], "head": False},
        ]
        sheet = _build([_sheet([_table(rows)])], formulas=False).worksheets[0]
        self.assertEqual(sheet["B5"].value, 600)

    def test_kpi_cards_land_side_by_side(self) -> None:
        grid = {
            "kind": "grid",
            "top": 40.0,
            "cells": [
                [{"kind": "text", "text": "CA", "style": _style()},
                 {"kind": "text", "text": "€1,200", "style": _style()}],
                [{"kind": "text", "text": "Commandes", "style": _style()},
                 {"kind": "text", "text": "2 312", "style": _style()}],
            ],
        }
        sheet = _build([_sheet([grid])]).worksheets[0]
        self.assertEqual(sheet["A1"].value, "CA")
        self.assertEqual(sheet["B1"].value, "Commandes")
        self.assertEqual(sheet["A2"].value, 1200)
        self.assertEqual(sheet["B2"].value, 2312)

    def test_sheet_name_is_taken_from_the_title_and_kept_legal(self) -> None:
        workbook = _build([_sheet([_table([{"cells": [_cell("x")], "head": False}])], "Ventes/2026")])
        self.assertEqual(workbook.sheetnames, ["Ventes 2026"])


class GridExpansionTests(unittest.TestCase):
    def test_colspan_reserves_the_cells_it_covers(self) -> None:
        rows = [
            {"cells": [_cell("Titre", colspan=2)]},
            {"cells": [_cell("A"), _cell("B")]},
        ]
        grid, merges = _grid(rows)
        self.assertEqual(len(grid[0]), 2)
        self.assertEqual(merges, [(0, 0, 0, 1)])

    def test_rowspan_pushes_the_next_row_aside(self) -> None:
        rows = [
            {"cells": [_cell("Phase", rowspan=2), _cell("S1")]},
            {"cells": [_cell("S2")]},
        ]
        grid, merges = _grid(rows)
        self.assertEqual(grid[1][0], None)
        self.assertEqual(grid[1][1]["text"], "S2")
        self.assertEqual(merges, [(0, 0, 1, 0)])


class CsvTests(unittest.TestCase):
    def test_csv_export_writes_one_file_per_sheet(self) -> None:
        table = _table([
            {"cells": [_cell("Produit", header=True), _cell("CA", header=True)], "head": True},
            {"cells": [_cell("Nova"), _cell("1 200")], "head": False},
        ])
        with tempfile.TemporaryDirectory() as raw:
            output = Path(raw) / "out.xlsx"
            build_workbook([_sheet([table], "Ventes")], output)
            written = export_csv(output, Path(raw) / "csv")
            self.assertEqual(len(written), 1)
            content = written[0].read_text(encoding="utf-8-sig")
        self.assertIn("Produit,CA", content)
        self.assertIn("Nova,1200", content)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
