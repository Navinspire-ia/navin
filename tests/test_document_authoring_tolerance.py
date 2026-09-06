"""What a model actually writes must not vanish from the document.

Layout names it invents, section keys it prefers, a callout as an object, a
totals row, a nested cover, speaker notes: each used to be dropped silently or
to crash the run. These tests pin the tolerant behaviour.
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from navin.documents import ppt_design, word_render


class LayoutResolutionTests(unittest.TestCase):
    def test_known_ids_pass_and_aliases_map(self) -> None:
        self.assertEqual(ppt_design.resolve_layout({"layout": "chart"}), ("chart", None))
        for alias, expected in (
            ("bullets", "text"), ("title", "cover"), ("thanks", "closing"), ("Title Slide", "cover"),
            ("bullet_points", "text"), ("kpis", "kpi"), ("two columns", "comparison"), ("sommaire", "agenda"),
        ):
            layout, note = ppt_design.resolve_layout({"layout": alias})
            self.assertEqual(layout, expected, alias)
            self.assertIsNone(note, alias)

    def test_unknown_names_fall_back_to_the_content_with_a_note(self) -> None:
        layout, note = ppt_design.resolve_layout({"layout": "zorglub", "chart": {"series": [{"name": "a", "values": [1]}]}})
        self.assertEqual(layout, "chart")
        self.assertIn("unknown layout 'zorglub'", note)
        layout, _ = ppt_design.resolve_layout({"layout": "zorglub", "table": {"headers": ["a"]}})
        self.assertEqual(layout, "table")
        layout, _ = ppt_design.resolve_layout({"layout": "zorglub", "items": ["a", "b", "c"]})
        self.assertEqual(layout, "cards")

    def test_a_deck_is_canonicalised_before_the_passes(self) -> None:
        slides = ppt_design.canonical_layouts(
            [{"layout": "title", "title": "Hi"}, {"layout": "bullets", "items": ["a"]}, {"title": "no layout"}]
        )
        self.assertEqual([s.get("layout") for s in slides], ["cover", "text", None])

    def test_render_slide_no_longer_dies_on_an_alias(self) -> None:
        html = ppt_design.render_slide("aurora_glass", {"layout": "bullets", "title": "Three things", "items": ["a", "b", "c"]})
        self.assertIn("Three things", html)
        self.assertIn('data-layout="text"', html)


class InvertedSlideTests(unittest.TestCase):
    def test_the_theme_carries_an_accent_ink_for_inverted_paper(self) -> None:
        css = ppt_design.theme_css(ppt_design.load_theme("aurora_glass"))
        self.assertIn("--nv-accent-ink-inverted:", css)

    def test_the_engine_rebinds_surface_tokens_on_inverted_slides(self) -> None:
        css = (ppt_design.engine_dir() / "navin-ppt.css").read_text(encoding="utf-8")
        start = css.index(".nv-slide.nv-tone-dark,\n.nv-slide.nv-tone-hero-dark {")
        block = css[start: css.index("}", start)]
        for token in ("--nv-surface:", "--nv-on-surface:", "--nv-border:", "--nv-muted:", "--nv-accent-ink:"):
            self.assertIn(token, block, token)


class SpeakerNotesTests(unittest.TestCase):
    def test_notes_stamped_on_the_html_land_in_the_deck(self) -> None:
        from pptx import Presentation

        from navin.documents import html2pptx

        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            slides = []
            for index, note in enumerate(("Say hello", "", "Wrap up &amp; thank"), start=1):
                page = folder / f"slide_{index:02d}.html"
                attr = f' data-notes="{note}"' if note else ""
                page.write_text(f'<html><body><div class="nv-slide"{attr}>x</div></body></html>', encoding="utf-8")
                slides.append(page)
            deck = folder / "deck.pptx"
            presentation = Presentation()
            for _ in range(3):
                presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(str(deck))
            written = html2pptx._apply_notes(slides, deck)
            notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(deck)).slides]
        self.assertEqual(written, 2)
        self.assertEqual(notes, ["Say hello", "", "Wrap up & thank"])


class ApproximatePreviewChartTests(unittest.TestCase):
    def test_a_native_chart_is_sketched_instead_of_left_blank(self) -> None:
        from PIL import Image
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        from navin.documents.preview_pptx import render_approximate

        with tempfile.TemporaryDirectory() as tmp:
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            data = CategoryChartData()
            data.categories = ["A", "B"]
            data.add_series("S", (1, 3))
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data)
            deck = Path(tmp) / "deck.pptx"
            presentation.save(str(deck))
            pages = render_approximate(deck, Path(tmp) / "out")
            with Image.open(pages[0]) as image:
                colours = image.getcolors(maxcolors=100000)
        self.assertGreater(len(colours), 2, "the chart box must paint bars, not stay white")


class WordRenderToleranceTests(unittest.TestCase):
    def test_alternate_keys_render_the_same_content(self) -> None:
        html = word_render._section_html(
            {
                "title": "Scope",
                "paragraphs": ["First paragraph.", "Second paragraph."],
                "bullets": ["Migration", "Connector"],
                "callout": {"title": "Commitment", "text": "Firm price for twelve months."},
                "table": {"headers": ["Lot", "Amount"], "rows": [["1", "18 000"]], "total": ["Total", "18 000"]},
            }
        )
        for needle in (
            "<h2>Scope</h2>", "First paragraph.", "Second paragraph.", "<li>Migration</li>",
            "<strong>Commitment</strong>", "Firm price", "<tfoot><tr><td>Total</td><td>18 000</td></tr></tfoot>",
        ):
            self.assertIn(needle, html, needle)
        self.assertNotIn("{'title'", html)

    def test_body_text_and_steps_are_understood(self) -> None:
        html = word_render._section_html({"heading": "Plan", "text": "One block.\n\nTwo blocks.", "steps": ["Kick-off", "Build"]})
        self.assertEqual(html.count("<p>"), 2)
        self.assertIn('<ol class="steps">', html)
        html = word_render._section_html({"heading": "Q", "quote": "It works.", "author": "A client"})
        self.assertIn("<blockquote><p>It works.</p><cite>A client</cite></blockquote>", html)

    def test_a_nested_cover_object_fills_the_cover(self) -> None:
        html = word_render._cover(
            {
                "title": "Proposal",
                "cover": {"subtitle": "For Globex", "kicker": "Confidential", "meta": ["2 September 2026", "Reference OF-17"]},
            }
        )
        self.assertIn("<h1>Proposal</h1>", html)
        self.assertIn("For Globex", html)
        self.assertIn('<div class="kicker">Confidential</div>', html)
        self.assertIn("2 September 2026", html)
        self.assertIn("Reference OF-17", html)

    def test_root_level_meta_fields_reach_the_cover(self) -> None:
        html = word_render._cover({"title": "Report", "date": "2026-09-02", "author": "Aymen", "version": "1.2"})
        self.assertIn("2026-09-02", html)
        self.assertIn("Aymen", html)
        self.assertIn("<strong>Version</strong>1.2", html)


if __name__ == "__main__":
    unittest.main()
