"""What the final-file validator blocks, warns about and lets through."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from navin.documents import doc_check


def _codes(report: doc_check.Report, severity: str | None = None) -> set[str]:
    return {f.code for f in report.findings if severity is None or f.severity == severity}


class TextRulesTests(unittest.TestCase):
    def _report(self, text: str) -> doc_check.Report:
        report = doc_check.Report("x", "text")
        doc_check.check_text(report, "here", text, seen=set())
        return report

    def test_a_clean_paragraph_has_no_finding(self) -> None:
        self.assertEqual(self._report("Revenue grew 43% in 2026, driven by cloud.").findings, [])

    def test_lorem_ipsum_and_placeholders_block(self) -> None:
        report = self._report("Lorem ipsum dolor sit amet. Dear {{client}}, see [[ref]].")
        self.assertIn("leftover-sample", _codes(report, "block"))
        self.assertIn("unresolved-placeholder", _codes(report, "block"))

    def test_dashes_and_emoji_are_reported(self) -> None:
        report = self._report("Q1 \u2014 Q2 results \U0001F4CA")
        self.assertIn("dash", _codes(report))
        self.assertIn("emoji", _codes(report))

    def test_one_finding_per_code_and_location(self) -> None:
        report = doc_check.Report("x", "text")
        seen: set[str] = set()
        doc_check.check_text(report, "p1", "{{a}} {{b}}", seen=seen)
        doc_check.check_text(report, "p1", "{{c}}", seen=seen)
        self.assertEqual(len([f for f in report.findings if f.code == "unresolved-placeholder"]), 1)


class DocxTests(unittest.TestCase):
    def _docx(self, folder: Path, paragraphs: list[tuple[str, str | None]], title: str | None = "Report") -> Path:
        from docx import Document

        document = Document()
        for text, style in paragraphs:
            document.add_paragraph(text, style=style) if style else document.add_paragraph(text)
        if title:
            document.core_properties.title = title
        path = folder / "doc.docx"
        document.save(str(path))
        return path

    def test_a_finished_report_passes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = self._docx(
                Path(tmp),
                [
                    ("Quarterly review", "Heading 1"),
                    (
                        "Revenue grew 43% year on year, led by the enterprise segment, "
                        "while the mid-market held steady and self-serve kept its pace.",
                        None,
                    ),
                    ("Margins", "Heading 2"),
                    (
                        "Gross margin reached 71%, up four points, as hosting costs fell "
                        "faster than prices and support automation absorbed the growth.",
                        None,
                    ),
                ],
            )
            report = doc_check.check(path)
        self.assertTrue(report.passed, report.as_dict())
        self.assertEqual(report.kind, "docx")

    def test_leftovers_block_and_are_located(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = self._docx(
                Path(tmp),
                [("Title", "Heading 1"), ("Dear {{client}}, lorem ipsum dolor sit amet.", None)],
            )
            report = doc_check.check(path)
        self.assertFalse(report.passed)
        blocked = [f for f in report.findings if f.severity == "block"]
        self.assertTrue(any("paragraph" in f.where for f in blocked))

    def test_a_missing_title_and_an_empty_body_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = self._docx(Path(tmp), [("Just one line of text.", None)], title=None)
            report = doc_check.check(path)
        self.assertIn("no-title", _codes(report, "warn"))
        self.assertIn("empty-document", _codes(report, "block"))

    def test_a_long_document_without_headings_is_reported(self) -> None:
        body = "The provider performs the services with reasonable care and skill. " * 70
        with tempfile.TemporaryDirectory() as tmp:
            path = self._docx(Path(tmp), [(body, None)])
            report = doc_check.check(path)
        self.assertIn("no-heading", _codes(report, "warn"))


class PptxTests(unittest.TestCase):
    def _deck(self, folder: Path, slides: list[tuple[str, str]]) -> Path:
        from pptx import Presentation

        presentation = Presentation()
        for title, body in slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text_frame.text = body
        presentation.core_properties.title = slides[0][0] if slides else "Deck"
        path = folder / "deck.pptx"
        presentation.save(str(path))
        return path

    def test_a_filled_deck_passes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = self._deck(
                Path(tmp),
                [
                    ("Revenue tripled in three years", "From 1.2M in 2023 to 3.6M in 2026\nEnterprise leads\nChurn down to 2.1%"),
                    ("Three priorities for 2027", "Grow ARR to 20M\nOpen the US office\nShip version 5"),
                ],
            )
            report = doc_check.check(path)
        self.assertTrue(report.passed, report.as_dict())
        self.assertEqual(report.stats.get("slides"), 2)

    def test_template_sample_copy_blocks(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = self._deck(Path(tmp), [("Title here", "Click to add text. Lorem ipsum dolor.")])
            report = doc_check.check(path)
        self.assertFalse(report.passed)
        self.assertIn("leftover-sample", _codes(report, "block"))

    def test_an_empty_slide_blocks(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.core_properties.title = "Deck"
            path = Path(tmp) / "deck.pptx"
            presentation.save(str(path))
            report = doc_check.check(path)
        self.assertIn("empty-slide", _codes(report, "block"))


class XlsxTests(unittest.TestCase):
    def test_placeholders_in_cells_block_and_formulas_count(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmp:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Budget"
            sheet.append(["Item", "Amount"])
            sheet.append(["Cloud", 1200])
            sheet.append(["{{next_item}}", 300])
            sheet.append(["Total", "=SUM(B2:B3)"])
            path = Path(tmp) / "budget.xlsx"
            workbook.save(str(path))
            report = doc_check.check(path)
        self.assertFalse(report.passed)
        self.assertIn("unresolved-placeholder", _codes(report, "block"))


class PdfTests(unittest.TestCase):
    def test_a_text_pdf_passes_and_a_blank_one_is_flagged(self) -> None:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        with tempfile.TemporaryDirectory() as tmp:
            good = Path(tmp) / "good.pdf"
            page = canvas.Canvas(str(good), pagesize=A4)
            page.setTitle("Service agreement")
            for i in range(12):
                page.drawString(72, 800 - 18 * i, f"Article {i + 1}. The provider performs the services with care.")
            page.save()
            report = doc_check.check(good)
            self.assertTrue(report.passed, report.as_dict())
            self.assertEqual(report.stats.get("pages"), 1)

            from pypdf import PdfWriter

            blank = Path(tmp) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with blank.open("wb") as handle:
                writer.write(handle)
            report = doc_check.check(blank)
        self.assertFalse(report.passed)
        self.assertIn("no-text-layer", _codes(report, "block"))
        self.assertIn("no-title", _codes(report, "warn"))


class DispatchTests(unittest.TestCase):
    def test_unknown_suffix_and_missing_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            other = Path(tmp) / "notes.txt"
            other.write_text("x", encoding="utf-8")
            with self.assertRaises(ValueError):
                doc_check.check(other)
            with self.assertRaises(FileNotFoundError):
                doc_check.check(Path(tmp) / "missing.docx")

    def test_cli_exit_code_follows_the_verdict(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            document = Document()
            document.add_paragraph("Heading", style="Heading 1")
            document.add_paragraph("Dear {{name}}")
            document.core_properties.title = "Letter"
            path = Path(tmp) / "letter.docx"
            document.save(str(path))
            self.assertEqual(doc_check.main([str(path)]), 1)
            self.assertEqual(doc_check.main([str(path), "--json"]), 1)


if __name__ == "__main__":
    unittest.main()
