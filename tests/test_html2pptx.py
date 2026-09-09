# Copyright (c) 2026-present Navinspire IA
# SPDX-License-Identifier: AGPL-3.0-only

"""Editability contract for the HTML to PPTX converter."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from navin.documents._chromium import instrument, run_chromium
from navin.documents.html2pptx import (
    EMU_PER_PX,
    SLIDE_HEIGHT_PX,
    SLIDE_WIDTH_PX,
    ConversionError,
    SlideData,
    _paragraph_lines,
    build_deck,
    check_critique,
    collect_slides,
    find_chromium,
    office_font,
    report,
)

ROOT = Path(__file__).resolve().parents[1]
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _run(text: str, **overrides) -> dict:
    run = {
        "text": text,
        "font": "Sora",
        "families": ["Sora", "Segoe UI", "Arial", "sans-serif"],
        "size": 48.0,
        "bold": True,
        "italic": False,
        "underline": False,
        "color": "FFFFFF",
        "spacing": 0,
    }
    run.update(overrides)
    return run


def _text_block(**overrides) -> dict:
    block = {
        "x": 120.0,
        "y": 320.0,
        "w": 1400.0,
        "h": 120.0,
        "align": "left",
        "rtl": False,
        "lineHeight": 56.0,
        "singleLine": False,
        "runs": [_run("Titre de la slide")],
    }
    block.update(overrides)
    return block


class FontMappingTests(unittest.TestCase):
    def test_first_office_font_of_the_stack_wins(self) -> None:
        self.assertEqual(office_font(["Sora", "Segoe UI", "Arial"], "Sora"), "Segoe UI")

    def test_unknown_web_font_falls_back_on_its_substitute(self) -> None:
        self.assertEqual(office_font(["Plus Jakarta Sans"], "Plus Jakarta Sans"), "Segoe UI")

    def test_generic_family_decides_when_nothing_else_matches(self) -> None:
        self.assertEqual(office_font(["Whatever Display", "serif"], "Whatever"), "Georgia")

    def test_keep_preserves_the_css_name(self) -> None:
        self.assertEqual(office_font(["Sora", "Arial"], "Sora", keep=True), "Sora")


class ParagraphSplitTests(unittest.TestCase):
    def test_line_breaks_start_new_paragraphs(self) -> None:
        runs = [_run("Ligne 1"), {"newline": True, "text": "\n"}, _run("Ligne 2")]
        lines = _paragraph_lines(runs)
        self.assertEqual([len(line) for line in lines], [1, 1])
        self.assertEqual(lines[1][0]["text"], "Ligne 2")

    def test_embedded_newline_inside_a_run_also_splits(self) -> None:
        lines = _paragraph_lines([_run("Haut\nBas")])
        self.assertEqual([line[0]["text"] for line in lines], ["Haut", "Bas"])


class DeckBuildTests(unittest.TestCase):
    def _deck(self, slides: list[SlideData]):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "deck.pptx"
            build_deck(slides, output)
            return Presentation(str(output))

    def test_slide_canvas_matches_the_html_master(self) -> None:
        deck = self._deck([SlideData(source=Path("slide_01.html"))])
        self.assertEqual(deck.slide_width, SLIDE_WIDTH_PX * EMU_PER_PX)
        self.assertEqual(deck.slide_height, SLIDE_HEIGHT_PX * EMU_PER_PX)

    def test_text_lands_in_editable_boxes_at_the_measured_position(self) -> None:
        deck = self._deck([SlideData(source=Path("s.html"), texts=[_text_block()])])
        shape = list(list(deck.slides)[0].shapes)[0]
        self.assertTrue(shape.has_text_frame)
        self.assertEqual(shape.text_frame.text, "Titre de la slide")
        self.assertEqual(shape.left, int(120 * EMU_PER_PX))
        self.assertEqual(shape.top, int(320 * EMU_PER_PX))

    def test_css_pixels_become_half_points(self) -> None:
        deck = self._deck([SlideData(source=Path("s.html"), texts=[_text_block()])])
        run = list(list(deck.slides)[0].shapes)[0].text_frame.paragraphs[0].runs[0]
        self.assertEqual(run.font.size.pt, 24.0)
        self.assertEqual(run.font.name, "Segoe UI")

    def test_a_measured_title_lands_in_the_native_title_placeholder(self) -> None:
        """The placeholder is what fills the Outline view and names the slide."""
        deck = self._deck(
            [SlideData(source=Path("s.html"), texts=[_text_block(title=True)])]
        )
        slide = list(deck.slides)[0]
        title = slide.shapes.title
        self.assertIsNotNone(title)
        self.assertEqual(title.text_frame.text, "Titre de la slide")
        # Placed where the page drew it, not where the master would have.
        self.assertEqual(title.left, int(120 * EMU_PER_PX))
        self.assertEqual(title.top, int(320 * EMU_PER_PX))

    def test_the_title_is_not_buried_under_the_decor(self) -> None:
        """It is cloned with the layout, so it starts underneath the backdrop."""
        deck = self._deck(
            [
                SlideData(
                    source=Path("s.html"),
                    texts=[_text_block(title=True)],
                    backdrop_color="102030",
                )
            ]
        )
        slide = list(deck.slides)[0]
        shapes = list(slide.shapes)
        # python-pptx hands out a fresh proxy per access, so the XML is the
        # only thing worth comparing.
        self.assertIs(shapes[-1]._element, slide.shapes.title._element)

    def test_a_slide_without_a_title_gets_no_empty_placeholder(self) -> None:
        """An empty 'Click to add title' box would print on the deck."""
        deck = self._deck([SlideData(source=Path("s.html"), texts=[_text_block()])])
        slide = list(deck.slides)[0]
        self.assertIsNone(slide.shapes.title)
        self.assertEqual(len(list(slide.shapes)), 1)

    def test_the_title_keeps_the_measured_type_rather_than_the_masters(self) -> None:
        deck = self._deck(
            [SlideData(source=Path("s.html"), texts=[_text_block(title=True)])]
        )
        run = list(deck.slides)[0].shapes.title.text_frame.paragraphs[0].runs[0]
        self.assertEqual(run.font.size.pt, 24.0)

    def test_single_line_labels_never_wrap(self) -> None:
        deck = self._deck(
            [
                SlideData(
                    source=Path("s.html"),
                    texts=[_text_block(singleLine=True), _text_block(singleLine=False)],
                )
            ]
        )
        frames = [shape.text_frame for shape in list(deck.slides)[0].shapes]
        self.assertFalse(frames[0].word_wrap)
        self.assertTrue(frames[1].word_wrap)

    def test_gradient_text_keeps_a_native_gradient_fill(self) -> None:
        gradient = {"stops": [{"pos": 0, "hex": "A78BFA"}, {"pos": 100, "hex": "22D3EE"}], "angle": 90}
        deck = self._deck(
            [SlideData(source=Path("s.html"), texts=[_text_block(runs=[_run("better", gradient=gradient)])])]
        )
        run = list(list(deck.slides)[0].shapes)[0].text_frame.paragraphs[0].runs[0]
        fill = run.font._rPr.find(f"{NS}gradFill")
        self.assertIsNotNone(fill)
        stops = fill.findall(f".//{NS}srgbClr")
        self.assertEqual([stop.get("val") for stop in stops], ["A78BFA", "22D3EE"])
        # A gradient run must not keep a competing solid fill.
        self.assertIsNone(run.font._rPr.find(f"{NS}solidFill"))

    def test_flat_decor_becomes_a_background_color_instead_of_a_picture(self) -> None:
        deck = self._deck([SlideData(source=Path("s.html"), backdrop_color="E0DBD1")])
        slide = list(deck.slides)[0]
        self.assertEqual(len(slide.shapes), 0)
        self.assertEqual(str(slide.background.fill.fore_color.rgb), "E0DBD1")

    def test_solid_blocks_become_editable_shapes(self) -> None:
        block = {"x": 100, "y": 100, "w": 400, "h": 200, "fill": "0369FF", "radius": 24, "border": None}
        deck = self._deck([SlideData(source=Path("s.html"), shapes=[block])])
        shape = list(list(deck.slides)[0].shapes)[0]
        self.assertEqual(str(shape.fill.fore_color.rgb), "0369FF")
        self.assertEqual(shape.width, int(400 * EMU_PER_PX))


class SlideCollectionTests(unittest.TestCase):
    def test_a_template_folder_yields_its_slides_in_order(self) -> None:
        folder = ROOT / "templates" / "ppt" / "aurora_glass"
        if not folder.is_dir():
            self.skipTest("built-in templates are not present in this checkout")
        slides = collect_slides([str(folder)])
        self.assertGreater(len(slides), 1)
        self.assertEqual([slide.name for slide in slides], sorted(slide.name for slide in slides))

    def test_a_missing_path_is_reported(self) -> None:
        with self.assertRaises(ConversionError):
            collect_slides(["/nowhere/slide_01.html"])


class ChromiumDiscoveryTests(unittest.TestCase):
    def test_an_explicit_binary_is_honored(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            binary = Path(folder) / "chrome"
            binary.write_text("", encoding="utf-8")
            self.assertEqual(find_chromium(str(binary)), str(binary))


class FrameTests(unittest.TestCase):
    def test_the_page_is_pinned_to_the_slide_size(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            workdir = Path(folder)
            source = workdir / "slide_01.html"
            source.write_text("<html><body>Bonjour</body></html>", encoding="utf-8")
            page = instrument(
                source, "void 0;", workdir, frame=(SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX)
            )
            html = page.read_text(encoding="utf-8")
        self.assertIn(f"height: {SLIDE_HEIGHT_PX}px !important", html)
        self.assertIn("overflow: hidden !important", html)

    def test_animations_are_frozen_at_their_resting_state(self) -> None:
        """A staggered item still at opacity 0 when the page is measured would
        be skipped by the measurement pass but painted by the later screenshot
        pass: its text ends up baked into the backdrop instead of editable."""
        with tempfile.TemporaryDirectory() as folder:
            workdir = Path(folder)
            source = workdir / "slide_01.html"
            source.write_text("<html><body>Bonjour</body></html>", encoding="utf-8")
            page = instrument(source, "void 0;", workdir)
            html = page.read_text(encoding="utf-8")
        self.assertIn("animation: none !important", html)
        self.assertIn("transition: none !important", html)

    def test_window_furniture_is_added_on_top_of_the_wanted_viewport(self) -> None:
        recorded: list[str] = []

        def fake_run(command, *args, **kwargs):
            recorded.extend(arg for arg in command if arg.startswith("--window-size"))
            raise ConversionError("stop")

        with mock.patch("navin.documents._chromium.subprocess.run", side_effect=fake_run):
            with self.assertRaises(ConversionError):
                run_chromium("chrome", ["--dump-dom"], 5, (1920, 1080), (0, 87))
        self.assertEqual(recorded, ["--window-size=1920,1167"])


class CritiqueGateTests(unittest.TestCase):
    """A deck the quality pass rejected must not reach the user by default."""

    def _slides(self, folder: str, critique: dict | None) -> list[Path]:
        slide = Path(folder) / "slide_01.html"
        slide.write_text("<html></html>", encoding="utf-8")
        if critique is not None:
            (Path(folder) / "critique.json").write_text(
                json.dumps(critique), encoding="utf-8"
            )
        return [slide]

    def test_a_failing_critique_blocks_the_conversion(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            slides = self._slides(
                folder,
                {"pass": False, "score": 76, "threshold": 85, "notes": ["No data hero."]},
            )
            with self.assertRaises(ConversionError) as caught:
                check_critique(slides)
        self.assertIn("76", str(caught.exception))
        self.assertIn("No data hero.", str(caught.exception))

    def test_force_downgrades_the_gate_to_a_warning(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            slides = self._slides(folder, {"pass": False, "score": 60, "threshold": 85})
            check_critique(slides, force=True)

    def test_a_passing_critique_converts_normally(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            slides = self._slides(folder, {"pass": True, "score": 92, "threshold": 85})
            check_critique(slides)

    def test_slides_without_a_critique_are_not_gated(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            slides = self._slides(folder, None)
            check_critique(slides)


class ReportTests(unittest.TestCase):
    def test_copy_running_past_the_frame_is_reported(self) -> None:
        data = SlideData(source=Path("slide_04.html"), overflow=81.0)
        warnings = report([data])
        self.assertTrue(any("81px past" in warning for warning in warnings))

    def test_a_visual_reused_across_slides_is_reported(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            picture = Path(folder) / "logo.png"
            picture.write_bytes(b"the same bytes")
            twin = Path(folder) / "logo-copy.png"
            twin.write_bytes(b"the same bytes")
            slides = [
                SlideData(source=Path("slide_02.html"), images=[{"path": str(picture)}]),
                SlideData(source=Path("slide_03.html"), images=[{"path": str(twin)}]),
            ]
            warnings = report(slides)
        self.assertTrue(any("appears 2 times" in warning for warning in warnings))

    def test_light_gray_copy_on_a_light_block_is_reported(self) -> None:
        from PIL import Image

        with tempfile.TemporaryDirectory() as folder:
            backdrop = Path(folder) / "bg.png"
            Image.new("RGB", (200, 80), (245, 245, 245)).save(backdrop)
            block = _text_block(
                x=0, y=0, w=200, h=80, runs=[_run("Presque invisible", color="EDEDED", size=18.0)]
            )
            data = SlideData(
                source=Path("slide_06.html"), texts=[block], backdrop=backdrop
            )
            warnings = report([data])
        self.assertTrue(any("unreadable text" in warning for warning in warnings))

    def test_dark_copy_on_a_light_block_passes(self) -> None:
        from PIL import Image

        with tempfile.TemporaryDirectory() as folder:
            backdrop = Path(folder) / "bg.png"
            Image.new("RGB", (200, 80), (245, 245, 245)).save(backdrop)
            block = _text_block(
                x=0, y=0, w=200, h=80, runs=[_run("Bien lisible", color="1A1A1A", size=18.0)]
            )
            data = SlideData(source=Path("slide_06.html"), texts=[block], backdrop=backdrop)
            self.assertEqual(report([data]), [])


if __name__ == "__main__":
    unittest.main()
