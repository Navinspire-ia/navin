"""Charts travel from the slide JSON to a native, editable PowerPoint chart."""

from __future__ import annotations

import html
import json
import re
import tempfile
import unittest
from pathlib import Path

from navin.documents import html2pptx, ppt_render
from navin.documents.html2pptx import SlideData, build_deck, normalize_chart_spec

_CHART_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _payload(markup: str) -> dict:
    match = re.search(r'data-chart="([^"]+)"', markup)
    assert match, markup
    return json.loads(html.unescape(match.group(1)))


class ChartSpecTests(unittest.TestCase):
    def test_points_become_one_series_with_the_variant_as_type(self) -> None:
        slide = {"layout": "chart", "variant": "column", "series": [{"label": "2024", "value": "1,2 M"}, {"label": "2025", "value": "3,6 M"}]}
        spec = ppt_render.chart_spec(slide)
        self.assertEqual(spec["type"], "column")
        self.assertEqual([p["label"] for p in spec["series"]], ["2024", "2025"])

    def test_a_full_chart_object_keeps_its_series_and_options(self) -> None:
        slide = {
            "layout": "chart",
            "chart": {
                "type": "line",
                "categories": ["Q1", "Q2", "Q3"],
                "series": [{"name": "ARR", "values": [10, "12,5", 15]}, {"name": "Costs", "values": [8, 8, 9]}],
                "number_format": "#,##0",
                "legend": True,
            },
        }
        spec = ppt_render.chart_spec(slide)
        self.assertEqual(spec["type"], "line")
        self.assertEqual(spec["categories"], ["Q1", "Q2", "Q3"])
        self.assertEqual(spec["series"][0]["values"], [10, 12.5, 15])
        self.assertEqual(spec["number_format"], "#,##0")
        self.assertTrue(spec["legend"])

    def test_an_unknown_type_falls_back_to_a_column_chart(self) -> None:
        spec = ppt_render.chart_spec({"chart": {"type": "sunburst", "series": [{"name": "a", "values": [1]}]}})
        self.assertEqual(spec["type"], "column")

    def test_a_slide_without_numbers_has_no_chart(self) -> None:
        self.assertIsNone(ppt_render.chart_spec({"layout": "chart", "body": "no data"}))


class HtmlDrawingTests(unittest.TestCase):
    def test_bars_carry_the_spec_and_scale_to_the_peak(self) -> None:
        markup = ppt_render._chart({"variant": "column", "series": [{"label": "A", "value": "50"}, {"label": "B", "value": "100"}]})
        self.assertIn('class="nv-chart nv-anim-stagger"', markup)
        self.assertEqual(_payload(markup)["type"], "column")
        heights = re.findall(r"height:(\d+)%", markup)
        self.assertEqual(heights, ["50", "100"])

    def test_a_donut_is_drawn_as_a_conic_gradient_with_a_legend(self) -> None:
        markup = ppt_render._chart({"chart": {"type": "doughnut", "series": [{"label": "Cloud", "value": 75}, {"label": "On-prem", "value": 25}]}})
        self.assertIn("nv-chart-round", markup)
        self.assertIn("conic-gradient(", markup)
        self.assertIn("<b>75%</b>", markup)
        self.assertEqual(_payload(markup)["type"], "doughnut")

    def test_lines_are_drawn_as_svg_with_one_polyline_per_series(self) -> None:
        markup = ppt_render._chart(
            {
                "chart": {
                    "type": "line",
                    "categories": ["Q1", "Q2"],
                    "series": [{"name": "ARR", "values": [1, 2]}, {"name": "Costs", "values": [1, 1]}],
                }
            }
        )
        self.assertIn("nv-chart-lines", markup)
        self.assertEqual(markup.count("<polyline"), 2)
        self.assertIn("nv-chart-legend", markup)
        self.assertIn("<span>Q2</span>", markup)

    def test_a_chart_slide_is_not_hollow(self) -> None:
        slide = {"layout": "chart", "title": "ARR triples", "chart": {"type": "column", "series": [{"name": "ARR", "values": [1, 3]}]}}
        self.assertFalse(ppt_render._slide_is_hollow(slide))
        self.assertTrue(ppt_render._slide_is_hollow({"layout": "chart", "title": "Nothing here"}))


class NormalizeTests(unittest.TestCase):
    def test_loose_shapes_become_categories_and_numeric_series(self) -> None:
        spec = normalize_chart_spec({"type": "bar", "series": [{"label": "A", "value": "1 200"}, {"label": "B", "value": "3,4k"}]})
        self.assertEqual(spec["type"], "bar")
        self.assertEqual(spec["categories"], ["A", "B"])
        self.assertEqual(spec["series"][0][1], [1200.0, 3400.0])
        self.assertEqual(spec["number_format"], "#,##0.##")

    def test_percentages_pick_a_percent_format_and_short_series_are_padded(self) -> None:
        spec = normalize_chart_spec({"categories": ["A", "B", "C"], "series": [{"name": "Share", "values": ["40%", "60%"]}]})
        self.assertEqual(spec["number_format"], '0"%"')
        self.assertEqual(spec["series"][0][1], [40.0, 60.0, None])

    def test_no_numbers_means_no_chart(self) -> None:
        self.assertIsNone(normalize_chart_spec({"series": [{"label": "A", "value": "n/a"}]}))
        self.assertIsNone(normalize_chart_spec({}))


def _chart_item(spec: dict, **overrides) -> dict:
    item = {"x": 100, "y": 200, "w": 800, "h": 400, "spec": spec, "ink": "222222", "size": 20, "families": ["Segoe UI"]}
    item.update(overrides)
    return item


class NativeChartTests(unittest.TestCase):
    def _deck(self, charts: list[dict]):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "deck.pptx"
            build_deck([SlideData(source=Path("slide_01.html"), charts=charts)], output)
            return Presentation(str(output))

    def test_a_column_chart_is_a_real_chart_with_the_data(self) -> None:
        deck = self._deck([_chart_item({"type": "column", "categories": ["2024", "2025"], "series": [{"name": "ARR", "values": [1.2, 3.6]}], "colors": ["#0369FF"]})])
        shapes = [s for s in deck.slides[0].shapes if s.has_chart]
        self.assertEqual(len(shapes), 1)
        chart = shapes[0].chart
        self.assertEqual(chart.chart_type.__str__().split(" ")[0], "COLUMN_CLUSTERED")
        self.assertEqual(list(chart.plots[0].categories), ["2024", "2025"])
        self.assertEqual(list(chart.plots[0].series[0].values), [1.2, 3.6])
        self.assertEqual(str(chart.plots[0].series[0].format.fill.fore_color.rgb), "0369FF")
        self.assertTrue(chart.plots[0].has_data_labels)
        self.assertFalse(chart.has_legend)
        self.assertEqual(shapes[0].left, int(100 * html2pptx.EMU_PER_PX))
        self.assertEqual(shapes[0].width, int(800 * html2pptx.EMU_PER_PX))

    def test_two_series_bring_a_legend_and_a_doughnut_colours_its_points(self) -> None:
        deck = self._deck(
            [
                _chart_item({"type": "line", "categories": ["Q1", "Q2"], "series": [{"name": "A", "values": [1, 2]}, {"name": "B", "values": [2, 1]}]}),
                _chart_item({"type": "doughnut", "series": [{"label": "Cloud", "value": "75%"}, {"label": "On-prem", "value": "25%"}]}, y=700, palette=["112233", "445566"]),
            ]
        )
        charts = [s.chart for s in deck.slides[0].shapes if s.has_chart]
        self.assertEqual(len(charts), 2)
        line, donut = charts
        self.assertTrue(line.has_legend)
        self.assertEqual(len(line.plots[0].series), 2)
        self.assertIn("DOUGHNUT", str(donut.chart_type))
        points = donut.plots[0].series[0].points
        self.assertEqual(str(points[0].format.fill.fore_color.rgb), "112233")
        self.assertEqual(str(points[1].format.fill.fore_color.rgb), "445566")
        self.assertTrue(donut.plots[0].data_labels.show_percentage)

    def test_the_chart_sits_on_the_slide_decor_without_its_own_box(self) -> None:
        deck = self._deck([_chart_item({"type": "column", "series": [{"label": "A", "value": 1}]})])
        chart = next(s.chart for s in deck.slides[0].shapes if s.has_chart)
        space = chart._chartSpace
        fill = space.find(f"{_CHART_NS}spPr")
        self.assertIsNotNone(fill)
        self.assertIsNotNone(fill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}noFill"))

    def test_a_spec_without_numbers_adds_no_chart(self) -> None:
        deck = self._deck([_chart_item({"series": [{"label": "A", "value": "tbd"}]})])
        self.assertEqual([s for s in deck.slides[0].shapes if s.has_chart], [])


if __name__ == "__main__":
    unittest.main()
