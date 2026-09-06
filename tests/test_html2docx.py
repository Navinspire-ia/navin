"""Editability contract for the HTML to DOCX converter."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from navin.documents._chromium import ConversionError, find_chromium
from navin.documents.html2docx import (
    PAGE_WIDTH_PX,
    PageData,
    _numbered_runs,
    _signature,
    build_document,
    convert,
    strip_page_furniture,
)

ROOT = Path(__file__).resolve().parents[1]


def _run(text: str, **overrides) -> dict:
    run = {
        "text": text,
        "font": "Inter",
        "families": ["Inter", "Segoe UI", "sans-serif"],
        "size": 16.0,
        "bold": False,
        "italic": False,
        "underline": False,
        "color": "1F2937",
        "spacing": 0,
    }
    run.update(overrides)
    return run


def _style(**overrides) -> dict:
    style = {
        "align": "left",
        "rtl": False,
        "indentLeft": 0,
        "indentRight": 0,
        "lineHeight": 24.0,
        "top": 100.0,
        "bottom": 124.0,
    }
    style.update(overrides)
    return style


def _paragraph(text: str, level: int = 0, top: float = 100.0, **overrides) -> dict:
    block = {
        "kind": "paragraph",
        "level": level,
        "runs": [_run(text)],
        "box": None,
        "style": _style(top=top, bottom=top + 24),
        "top": top,
        "bottom": top + 24,
    }
    block.update(overrides)
    return block


def _page(blocks: list[dict], **overrides) -> PageData:
    page = PageData(
        width=float(PAGE_WIDTH_PX),
        height=1123.0,
        margins={"top": 68.0, "right": 68.0, "bottom": 68.0, "left": 68.0},
        blocks=blocks,
    )
    for key, value in overrides.items():
        setattr(page, key, value)
    return page


def _build(pages: list[PageData]) -> Document:
    with tempfile.TemporaryDirectory() as raw:
        output = Path(raw) / "out.docx"
        build_document(pages, output)
        return Document(str(output))


class PageGeometryTests(unittest.TestCase):
    def test_section_is_a4_with_the_template_margins(self) -> None:
        document = _build([_page([_paragraph("Bonjour")])])
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, delta=0.5)
        self.assertAlmostEqual(section.page_height.mm, 297, delta=0.5)
        # 68px at 96dpi is 18mm, the padding of the master.
        self.assertAlmostEqual(section.left_margin.mm, 18, delta=0.5)


class TextTests(unittest.TestCase):
    def test_headings_use_word_heading_styles(self) -> None:
        document = _build([_page([_paragraph("Titre", level=1)])])
        self.assertEqual(document.paragraphs[0].style.name, "Heading 1")

    def test_typography_survives_the_conversion(self) -> None:
        block = _paragraph("Texte")
        block["runs"] = [_run("Texte", bold=True, size=26.0, color="6D28D9")]
        run = _build([_page([block])]).paragraphs[0].runs[0]
        self.assertTrue(run.bold)
        # 26 CSS pixels are 19.5 points.
        self.assertAlmostEqual(run.font.size.pt, 19.5, places=1)
        self.assertEqual(str(run.font.color.rgb), "6D28D9")

    def test_web_font_is_mapped_onto_an_office_font(self) -> None:
        run = _build([_page([_paragraph("Texte")])]).paragraphs[0].runs[0]
        # Inter is a web font Word does not have; Calibri is its substitute.
        self.assertEqual(run.font.name, "Calibri")
        # Complex scripts fall back to the theme font without these slots.
        fonts = run._element.rPr.find(qn("w:rFonts"))
        self.assertEqual(fonts.get(qn("w:cs")), "Calibri")

    def test_badge_background_becomes_run_shading(self) -> None:
        block = _paragraph("Statut")
        block["runs"] = [_run("En cours", highlight="FEF9C3")]
        run = _build([_page([block])]).paragraphs[0].runs[0]
        shading = run._element.rPr.find(qn("w:shd"))
        self.assertEqual(shading.get(qn("w:fill")), "FEF9C3")


class BoxTests(unittest.TestCase):
    def test_callout_keeps_its_fill_and_borders(self) -> None:
        box = {"id": "b1", "fill": "F5F3FF", "borders": {"left": {"width": 2, "color": "E9D5FF"}}}
        document = _build([_page([_paragraph("Décision", box=box)])])
        properties = document.paragraphs[0]._p.pPr
        self.assertEqual(properties.find(qn("w:shd")).get(qn("w:fill")), "F5F3FF")
        borders = properties.find(qn("w:pBdr"))
        self.assertIsNotNone(borders.find(qn("w:left")))

    def test_shared_box_does_not_draw_lines_between_its_paragraphs(self) -> None:
        box = {"id": "b1", "fill": None, "borders": {
            "top": {"width": 1, "color": "EEEEEE"},
            "bottom": {"width": 1, "color": "EEEEEE"},
        }}
        first = _paragraph("Premier", top=100.0, box=box)
        second = _paragraph("Second", top=140.0, box=box)
        document = _build([_page([first, second])])
        top_borders = document.paragraphs[0]._p.pPr.find(qn("w:pBdr"))
        bottom_borders = document.paragraphs[1]._p.pPr.find(qn("w:pBdr"))
        self.assertIsNotNone(top_borders.find(qn("w:top")))
        self.assertIsNone(top_borders.find(qn("w:bottom")))
        self.assertIsNotNone(bottom_borders.find(qn("w:bottom")))


class TableTests(unittest.TestCase):
    def _table(self) -> dict:
        def cell(text: str, header: bool = False, **extra) -> dict:
            data = {
                "runs": [_run(text)],
                "header": header,
                "colspan": 1,
                "rowspan": 1,
                "width": 200.0,
                "align": "left",
                "valign": "top",
                "fill": "6D28D9" if header else None,
                "borders": None,
                "padding": {"top": 8, "right": 11, "bottom": 8, "left": 11},
            }
            data.update(extra)
            return data

        return {
            "kind": "table",
            "rows": [
                {"cells": [cell("Action", True), cell("Pilote", True)], "head": True, "height": 30},
                {"cells": [cell("Spécifier"), cell("C. Dupont")], "head": False, "height": 30},
            ],
            "width": 658.0,
            "indentLeft": 0,
            "top": 200.0,
            "bottom": 260.0,
        }

    def test_table_cells_hold_real_text(self) -> None:
        document = _build([_page([self._table()])])
        table = document.tables[0]
        self.assertEqual(table.cell(0, 0).text, "Action")
        self.assertEqual(table.cell(1, 1).text, "C. Dupont")

    def test_header_row_keeps_its_fill_and_repeats(self) -> None:
        document = _build([_page([self._table()])])
        table = document.tables[0]
        shading = table.cell(0, 0)._tc.tcPr.find(qn("w:shd"))
        self.assertEqual(shading.get(qn("w:fill")), "6D28D9")
        self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn("w:tblHeader")))

    def test_merged_cells_follow_colspan(self) -> None:
        block = self._table()
        block["rows"][1]["cells"][0]["colspan"] = 2
        block["rows"][1]["cells"] = block["rows"][1]["cells"][:1]
        document = _build([_page([block])])
        self.assertEqual(document.tables[0].cell(1, 0).text, "Spécifier")


class ListTests(unittest.TestCase):
    def test_lists_use_native_word_list_styles(self) -> None:
        block = {
            "kind": "list",
            "ordered": True,
            "items": [
                {"runs": [_run("Premier point")], "level": 0, "style": _style()},
                {"runs": [_run("Sous-point")], "level": 1, "style": _style()},
            ],
            "box": None,
            "top": 100.0,
            "bottom": 160.0,
        }
        document = _build([_page([block])])
        self.assertEqual(document.paragraphs[0].style.name, "List Number")
        self.assertEqual(document.paragraphs[1].style.name, "List Number 2")


class RunningPartsTests(unittest.TestCase):
    def test_repeated_banner_moves_into_the_header(self) -> None:
        pages = [
            _page([_paragraph("MODULE JURIDIQUE", top=60.0), _paragraph("Article 1", top=200.0)]),
            _page([_paragraph("MODULE JURIDIQUE", top=1200.0), _paragraph("Article 2", top=1340.0)]),
        ]
        banner = strip_page_furniture(pages)
        self.assertEqual(len(banner), 1)
        self.assertEqual([b["runs"][0]["text"] for b in pages[0].blocks], ["Article 1"])
        self.assertEqual([b["runs"][0]["text"] for b in pages[1].blocks], ["Article 2"])

    def test_single_page_keeps_everything_in_the_flow(self) -> None:
        pages = [_page([_paragraph("Titre", top=60.0)])]
        self.assertEqual(strip_page_furniture(pages), [])
        self.assertEqual(len(pages[0].blocks), 1)

    def test_page_numbers_become_word_fields(self) -> None:
        runs = _numbered_runs([_run("Page 3 / 8")])
        self.assertEqual([run.get("field") for run in runs if run.get("field")], ["PAGE", "NUMPAGES"])

    def test_plain_footer_text_is_left_alone(self) -> None:
        runs = _numbered_runs([_run("Confidentiel")])
        self.assertEqual(len(runs), 1)
        self.assertNotIn("field", runs[0])

    def test_footer_is_written_even_when_it_starts_on_a_later_page(self) -> None:
        pages = [
            _page([_paragraph("Couverture")]),
            _page([_paragraph("Suite")], footer={"runs": [_run("Confidentiel")], "align": "left"}),
        ]
        document = _build(pages)
        self.assertIn("Confidentiel", document.sections[0].footer.paragraphs[0].text)


class SignatureTests(unittest.TestCase):
    def test_same_text_gives_the_same_signature(self) -> None:
        self.assertEqual(_signature(_paragraph("Titre")), _signature(_paragraph("Titre  ")))

    def test_heading_level_changes_the_signature(self) -> None:
        self.assertNotEqual(_signature(_paragraph("Titre")), _signature(_paragraph("Titre", 1)))


class PaginationTests(unittest.TestCase):
    def test_headings_keep_the_next_paragraph(self) -> None:
        document = _build([_page([_paragraph("Titre", level=1), _paragraph("Corps", top=140)])])
        heading = document.paragraphs[0]._p.pPr
        from docx.oxml.ns import qn

        self.assertIsNotNone(heading.find(qn("w:keepNext")))
        self.assertIsNotNone(heading.find(qn("w:keepLines")))
        self.assertIsNotNone(heading.find(qn("w:widowControl")))

    def test_page_break_before_follows_the_directive(self) -> None:
        block = _paragraph("Annexe", level=1)
        block["doc"] = {"break": "page"}
        document = _build([_page([block])])
        from docx.oxml.ns import qn

        self.assertIsNotNone(document.paragraphs[0]._p.pPr.find(qn("w:pageBreakBefore")))

    def test_table_rows_do_not_split_across_pages(self) -> None:
        def cell(text: str, header: bool = False) -> dict:
            return {
                "runs": [_run(text)],
                "header": header,
                "colspan": 1,
                "rowspan": 1,
                "width": 200.0,
                "align": "left",
                "valign": "top",
                "fill": None,
                "borders": None,
                "padding": {"top": 8, "right": 11, "bottom": 8, "left": 11},
            }

        table_block = {
            "kind": "table",
            "rows": [
                {"cells": [cell("A", True), cell("B", True)], "head": True, "height": 30},
                {"cells": [cell("1"), cell("2")], "head": False, "height": 30},
            ],
            "width": 400.0,
            "indentLeft": 0,
            "top": 200.0,
            "bottom": 260.0,
        }
        document = _build([_page([table_block])])
        from docx.oxml.ns import qn

        for row in document.tables[0].rows:
            self.assertIsNotNone(row._tr.trPr.find(qn("w:cantSplit")))


class TocAndReferenceTests(unittest.TestCase):
    def test_toc_block_inserts_a_refreshable_field(self) -> None:
        pages = [_page([
            {
                "kind": "toc",
                "levels": "1-3",
                "runs": [_run("Contents")],
                "top": 80.0,
                "bottom": 120.0,
            },
            _paragraph("Market", level=1, top=160),
        ])]
        document = _build(pages)
        xml = document.element.xml
        self.assertIn("TOC", xml)
        self.assertIn("updateFields", document.settings.element.xml)

    def test_caption_uses_a_sequence_field_and_a_bookmark(self) -> None:
        caption = _paragraph("Architecture", top=200)
        caption["doc"] = {"caption": "Figure", "bookmark": "fig:arch"}
        document = _build([_page([caption])])
        xml = document.element.xml
        self.assertIn("SEQ Figure", xml)
        self.assertIn("fig_arch", xml)

    def test_cross_reference_points_at_the_bookmark(self) -> None:
        caption = _paragraph("Architecture", top=200)
        caption["doc"] = {"caption": "Figure", "bookmark": "fig:arch"}
        see = _paragraph("See ", top=240)
        see["doc"] = {"ref": "fig:arch"}
        document = _build([_page([caption, see])])
        self.assertIn("REF fig_arch", document.element.xml)


class SectionTests(unittest.TestCase):
    def test_landscape_page_becomes_a_separate_section(self) -> None:
        portrait = _page([_paragraph("Intro")])
        landscape = _page([_paragraph("Wide table")])
        landscape.doc = {"section": "landscape"}
        document = _build([portrait, landscape])
        self.assertEqual(len(document.sections), 2)
        second = document.sections[1]
        self.assertGreater(second.page_width, second.page_height)

    def test_heading_numbering_binds_when_asked(self) -> None:
        heading = _paragraph("Market", level=1)
        heading["doc"] = {"numbering": "1-3"}
        document = _build([_page([heading])])
        from docx.oxml.ns import qn

        self.assertIsNotNone(document.styles["Heading 1"].element.pPr.find(qn("w:numPr")))


def _chromium_available() -> bool:
    try:
        return bool(find_chromium(None))
    except ConversionError:
        return False


@unittest.skipUnless(_chromium_available(), "Chromium is required to measure HTML")
class HtmlDirectiveConversionTests(unittest.TestCase):
    def test_annotated_html_produces_toc_landscape_and_captions(self) -> None:
        html = """<!doctype html><html><head><style>
        .page { width: 794px; min-height: 400px; padding: 48px; box-sizing: border-box; }
        h1 { font-size: 28px; }
        p { font-size: 16px; }
        </style></head><body>
        <section class="page" data-doc-numbering="1-3">
          <nav data-doc-toc="1-3">Contents</nav>
          <h1>Market</h1>
          <p>Growth stayed above plan.</p>
        </section>
        <section class="page" data-doc-section="landscape">
          <h1>Appendix</h1>
          <p data-doc-caption="Figure" data-doc-bookmark="fig:one">Revenue chart</p>
        </section>
        </body></html>"""
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "document.html"
            output = Path(raw) / "out.docx"
            source.write_text(html, encoding="utf-8")
            convert(str(source), output)
            document = Document(str(output))
        xml = document.element.xml
        self.assertIn("TOC", xml)
        self.assertIn("SEQ Figure", xml)
        self.assertIn("fig_one", xml)
        self.assertGreaterEqual(len(document.sections), 2)
        self.assertGreater(document.sections[-1].page_width, document.sections[-1].page_height)

    HEADER_PAGE = """<!doctype html><html><head><style>
    * { margin: 0; padding: 0; box-sizing: border-box; }
    .page { width: 794px; min-height: 400px; padding: 48px; }
    .head { display: flex; justify-content: space-between; }
    h1 { font-size: 28px; }
    p { font-size: 16px; }
    </style></head><body>
    <section class="page">
      <div class="head"%s>
        <div><div class="kicker">Requirements</div><h1>Field service app</h1></div>
        <div class="meta">Version 1.2</div>
      </div>
      <p>The scope covers scheduling and offline capture.</p>
    </section>
    </body></html>"""

    def _headings(self, html: str) -> list[tuple[str, str]]:
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "document.html"
            output = Path(raw) / "out.docx"
            source.write_text(html, encoding="utf-8")
            convert(str(source), output)
            document = Document(str(output))
        return [
            (p.style.name, p.text)
            for p in document.paragraphs
            if p.style.name.startswith("Heading")
        ]

    def test_a_title_nested_in_a_header_band_is_lost_without_the_directive(self) -> None:
        """The row becomes a table, and a heading in a cell has no outline level."""
        self.assertEqual(self._headings(self.HEADER_PAGE % ""), [])

    def test_data_doc_row_stack_brings_the_title_back(self) -> None:
        headings = self._headings(self.HEADER_PAGE % ' data-doc-row="stack"')
        self.assertEqual(headings, [("Heading 1", "Field service app")])

    def test_a_heading_that_is_the_rows_own_child_needs_no_directive(self) -> None:
        html = self.HEADER_PAGE.replace(
            '<div><div class="kicker">Requirements</div><h1>Field service app</h1></div>',
            "<h1>Field service app</h1>",
        )
        self.assertEqual(self._headings(html % ""), [("Heading 1", "Field service app")])


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
