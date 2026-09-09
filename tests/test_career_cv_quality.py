"""Career documents keep candidate facts through AI, persistence and Word export."""

from __future__ import annotations

import base64
import copy
import io
import json
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from navin.career.ai import _facts, apply_revision, apply_valid_revisions, polish_pack
from navin.career.desk import apply_one, download_pack, prepare_application
from navin.career.errors import CareerError
from navin.career.export import attach_exports, render_docx
from navin.career.store import CareerStore
from navin.career.writer import build_pack, cv_to_text, job_keywords


@pytest.fixture
def profile():
    return {
        "display_name": "Camille Morel",
        "headline": "Ingénieure data",
        "email": "camille@example.invalid",
        "city": "Lyon",
        "document_language": "fr",
        "languages": ["fr", "en"],
        "stack": ["SQL", "Python", "Spark", "AWS", "Airflow"],
        "summary": "Ingénieure data spécialisée dans les pipelines Python et Spark sur AWS.",
        "master_cv": "Camille Morel\nIngénieure data\nExpérience professionnelle\n"
        "Senior Data Engineer | Atelier Boréal | 2022 - présent\n"
        "Optimisation des pipelines Spark sur AWS : traitement réduit de 4 à 2 heures.\n"
        "Formation de 3 analystes aux contrôles SQL.\n\n"
        "Data Engineer | Orme Conseil | 2019-2022\n"
        "Automatisation de 8 sources Python avec Airflow.\n\n"
        "Analyste | Rivage Études | 2017-2019\nReporting SQL pour les équipes produit.\n"
        "Formation\nMaster informatique | Université du Rhône | 2017\n"
        "Certifications\nAWS Certified Data Engineer | 2024",
        "experiences": [
            {"title": "Analyste", "company": "Rivage Études", "period": "2017-2019", "facts": "Reporting SQL pour les équipes produit."},
            {"title": "Data Engineer", "company": "Orme Conseil", "period": "2019-2022", "facts": "Automatisation de 8 sources Python avec Airflow."},
            {"title": "Senior Data Engineer", "company": "Atelier Boréal", "period": "2022 - présent", "facts": "Formation de 3 analystes aux contrôles SQL.\nOptimisation des pipelines Spark sur AWS : traitement réduit de 4 à 2 heures."},
        ],
        "education": [{"diploma": "Master informatique", "school": "Université du Rhône", "year": "2017"}],
        "projects": ["Contrôles de qualité SQL pour les données produit."],
        "ai_assist": True,
    }


@pytest.fixture
def job():
    return {
        "id": "job-fictional-quality", "title": "Senior Data Engineer",
        "company": "Horizon Mobilité", "country": "FR", "source": "import",
        "url": "https://employer.example.invalid/jobs/data",
        "document_language": "fr",
        "description": "Pour ce poste chez Horizon Mobilité, Python et Spark sont requis. "
        "AWS est indispensable pour des pipelines fiables. Kafka est souhaitable.",
        "must_haves": ["Python", "Spark", "AWS"],
    }


def _document_text(blob: bytes) -> str:
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(blob)).paragraphs)


def _revision(pack):
    row = pack["cv"]["experiences"][0]
    return {
        "summary": {"text": "Ingénieure data : pipelines Python et Spark sur AWS.", "source_ids": ["profile"]},
        "experiences": [{"id": row["id"], "bullets": [
            {"text": "Pipelines Spark sur AWS optimisés : traitement réduit de 4 à 2 heures.", "source_ids": [f"{row['id']}:bullet-0"]},
            {"text": row["bullets"][1], "source_ids": [f"{row['id']}:bullet-1"]},
        ]}],
    }


def test_ats_tracks_requirements_and_phrases_without_offer_stopwords(profile, job):
    pack = build_pack(job, profile)
    assert set(pack["keywords_matched"]) == {"Python", "Spark", "AWS"}
    assert pack["keywords_missing"] == ["Kafka"]
    assert pack["ats_score"] == 90
    assert "Kafka" not in pack["cv"]["skills"]
    assert not {"pour", "chez", "des", "poste"} & {item.lower() for item in pack["keywords_matched"]}
    assert job_keywords({"description": "Gestion de projets avec Power BI et Google Cloud Platform."}) == ["Gestion de projets", "Power BI", "Google Cloud Platform"]


def test_negative_skill_mentions_do_not_become_candidate_skills(job):
    profile = {"display_name": "Camille", "master_cv": "Python et Spark. Aucune expérience avec AWS. Kafka : à apprendre.", "document_language": "fr"}
    pack = build_pack(job, profile)
    assert pack["keywords_matched"] == ["Python", "Spark"]
    assert set(pack["keywords_missing"]) == {"AWS", "Kafka"}
    assert "AWS" not in pack["cv"]["skills"]
    assert "Kafka" not in pack["cv"]["skills"]


def test_reverse_chronology_relevance_and_natural_factual_letter(profile, job):
    pack = build_pack(job, profile)
    cv = pack["cv"]
    assert [row["company"] for row in cv["experiences"]] == ["Atelier Boréal", "Orme Conseil", "Rivage Études"]
    assert "Spark" in cv["experiences"][0]["bullets"][0]
    assert [row["period"] for row in cv["experiences"]] == ["2022 - présent", "2019-2022", "2017-2019"]
    assert cv["name"] == profile["display_name"]
    assert cv["language"] == "fr"
    assert cv["education"] == profile["education"]
    assert "AWS Certified Data Engineer" in pack["cv_text"]
    for phrase in ("Do not invent", "on-file", "not on file", "faits deja", "ATS"):
        assert phrase not in pack["summary"] + pack["cover"]
    first_proof = next(part for part in pack["cover"].split("\n\n") if "Atelier Boréal" in part)
    assert "8 sources" not in first_proof
    assert "8 sources" in pack["cover"]
    assert pack["cv_text"] == cv_to_text(cv, "fr")


def test_same_year_chronology_uses_months_without_rewriting_dates(profile, job):
    profile["experiences"] = [
        {"title": "Analyste", "company": "Janvier", "period": "01/2024 - 04/2024", "facts": "SQL"},
        {"title": "Analyste", "company": "Décembre", "period": "mai 2024 - décembre 2024", "facts": "SQL"},
        {"title": "Analyste", "company": "Actuel", "period": "2023 - en cours", "facts": "SQL"},
    ]
    assert [row["company"] for row in build_pack(job, profile)["cv"]["experiences"]] == ["Actuel", "Décembre", "Janvier"]


def test_complete_unstructured_master_is_preserved_past_4000_characters(profile, job):
    profile["experiences"] = []
    profile["master_cv"] = "Expérience professionnelle\n" + "\n\n".join(
        f"Mission historique {index} | Entreprise {index} | 2001-2002\nResponsabilité réelle distinctive {index}."
        for index in range(60)
    ) + "\n\nPublications\nPublication finale : Méthodes de traçabilité des données."
    pack = build_pack(job, profile)
    assert len(profile["master_cv"]) > 4000
    assert profile["master_cv"] in json.loads(_facts(job, profile))["candidate"]["master_cv"]
    exported = _document_text(render_docx(job, profile, pack, document="cv"))
    for part in ("Mission historique 0", "Mission historique 59", "Publication finale : Méthodes de traçabilité des données."):
        assert part in pack["cv_text"]
        assert part in exported
    assert pack["cv"]["skills"]
    assert "Master informatique" in exported


def test_partial_structured_profile_does_not_erase_master_education(profile, job):
    profile["education"] = []
    pack = build_pack(job, profile)
    assert "Master informatique" in pack["cv_text"]
    assert "Université du Rhône" in _document_text(render_docx(job, profile, pack, document="cv"))


def test_ai_revisions_reach_canonical_cv_text_and_word_with_skills(profile, job):
    original = build_pack(job, profile)
    revision = _revision(original)
    with patch.dict("os.environ", {"NAVIN_CAREER_AI": "on"}), patch("navin.career.ai.ask", return_value=(json.dumps(revision), "fixture-model")) as ask:
        pack = polish_pack(job, profile, original)
    assert pack["generation"]["mode"] == "model", pack["generation"]
    assert pack["model"] == "fixture-model"
    assert {"cv-tailoring", "cv-builder", "cover-letter-writer", "ats-analyzer", "docx-generator"} <= set(pack["generation"]["skills"]["loaded"])
    assert "source_ids" in ask.call_args.args[1]
    revised_bullet = revision["experiences"][0]["bullets"][0]["text"]
    assert revised_bullet == pack["cv"]["experiences"][0]["bullets"][0]
    assert revised_bullet in pack["cv_text"]
    assert revised_bullet in _document_text(render_docx(job, profile, pack))
    assert pack["summary"] == pack["cv"]["summary"] == revision["summary"]["text"]
    assert original["cv"]["experiences"][0]["bullets"][0] != revised_bullet
    assert pack["cv"]["education"] == original["cv"]["education"]


def _provider_fixture():
    return json.loads((Path(__file__).parent / "fixtures" / "career_provider_revision.json").read_text(encoding="utf-8"))


def test_real_provider_valid_edits_reach_word_despite_unsupported_cover():
    fixture = _provider_fixture()
    profile, job = fixture["input"]["profile"], fixture["input"]["job"]
    baseline, revision = fixture["baseline"], fixture["revision"]
    with patch.dict("os.environ", {"NAVIN_CAREER_AI": "on"}), patch("navin.career.ai.ask", side_effect=[(json.dumps(revision), fixture["captured_provider"]), ("", "")]) as ask:
        pack = polish_pack(job, profile, baseline)
    assert ask.call_count == 2
    assert pack["generation"]["mode"] == "model"
    assert pack["generation"]["status"] == "needs_review"
    assert pack["generation"]["rejected_fields"][0]["field"] == "cover"
    assert "industrialiser" in pack["generation"]["rejected_fields"][0]["reason"]
    assert pack["summary"] == revision["summary"]["text"]
    assert pack["cover"] == baseline["cover"]
    assert pack["summary"] in _document_text(render_docx(job, profile, pack, document="cv"))
    assert pack["cv"]["education"] == baseline["cv"]["education"]
    assert "guaranteed" not in pack["cv_text"].lower()


def test_one_repair_only_changes_rejected_fields_and_preserves_valid_provider_work():
    fixture = _provider_fixture()
    profile, job = fixture["input"]["profile"], fixture["input"]["job"]
    baseline, revision = fixture["baseline"], fixture["revision"]
    repair = {
        "cover": [{"text": paragraph, "source_ids": ["profile"]} for paragraph in baseline["cover"].split("\n\n")],
        "summary": {"text": "Invented job at Google.", "source_ids": ["profile"]},
    }
    with patch.dict("os.environ", {"NAVIN_CAREER_AI": "on"}), patch("navin.career.ai.ask", side_effect=[(json.dumps(revision), "fixture-model"), (json.dumps(repair), "fixture-model")]) as ask:
        pack = polish_pack(job, profile, baseline)
    assert ask.call_count == 2
    repair_input = json.loads(ask.call_args_list[1].args[2])["repair"]
    assert repair_input["fields"] == ["cover"]
    assert pack["generation"]["repair_attempted"]
    assert pack["generation"]["status"] == "complete", pack["generation"]
    assert pack["generation"]["rejected_fields"] == []
    assert pack["summary"] == revision["summary"]["text"]
    assert pack["cover"] == baseline["cover"]
    assert "Google" not in pack["cv_text"]


def test_actual_provider_repair_accepts_courtesy_and_exact_signature_without_sources():
    fixture = _provider_fixture()
    profile, job = fixture["input"]["profile"], fixture["input"]["job"]
    baseline = fixture["baseline"]
    with patch.dict("os.environ", {"NAVIN_CAREER_AI": "on"}), patch("navin.career.ai.ask", side_effect=[
        (json.dumps(fixture["revision"]), fixture["captured_provider"]),
        (json.dumps(fixture["repair_revision"]), fixture["captured_provider"]),
    ]) as ask:
        pack = polish_pack(job, profile, baseline)
    assert ask.call_count == 2
    assert pack["generation"]["status"] == "complete", pack["generation"]
    assert pack["generation"]["rejected_fields"] == []
    assert "cover" in pack["generation"]["accepted_fields"]
    assert pack["cover"] == "\n\n".join(item["text"] for item in fixture["repair_revision"]["cover"])
    assert pack["summary"] == fixture["revision"]["summary"]["text"]
    assert "Madame, Monsieur" in _document_text(render_docx(job, profile, pack, document="cover"))


@pytest.mark.parametrize("unreferenced", [
    "Chez Atelier Boréal, automatisation de 8 sources Python.",
    "Compétences Python, Spark et AWS.",
    "Camille Morel Dupont",
    "Je suis manager.",
])
def test_courtesy_exception_never_allows_uncited_candidate_claims(unreferenced):
    fixture = _provider_fixture()
    revision = {"cover": copy.deepcopy(fixture["repair_revision"]["cover"])}
    revision["cover"][-1] = {"text": unreferenced, "source_ids": []}
    with pytest.raises(ValueError):
        apply_revision(fixture["input"]["job"], fixture["input"]["profile"], fixture["baseline"], revision)


def test_invalid_role_does_not_discard_independent_valid_summary(profile, job):
    baseline = build_pack(job, profile)
    revision = _revision(baseline)
    revision["experiences"][0]["bullets"][0]["text"] = "Pipelines Kubernetes."
    pack, accepted, rejected = apply_valid_revisions(job, profile, baseline, revision)
    assert accepted == ["summary"]
    assert rejected[0]["field"] == "experiences.experience-2"
    assert pack["summary"] == revision["summary"]["text"]
    assert pack["cv"]["experiences"] == baseline["cv"]["experiences"]


@pytest.mark.parametrize("claim", [
    "Chez Atelier Boréal, automatisation de 8 sources Python.",
    "Chez Orme Conseil, pipelines Spark sur AWS.",
    "Traitement réduit de 2 à 4 heures.",
])
def test_summary_cannot_move_metrics_or_tools_to_another_employer(profile, job, claim):
    baseline = build_pack(job, profile)
    with pytest.raises(ValueError):
        apply_revision(job, profile, baseline, {"summary": {"text": claim, "source_ids": ["profile"]}})


@pytest.mark.parametrize("fault", ["employer", "education", "new_tool", "metric", "swapped_metrics", "unknown_source", "other_role", "omitted_bullet", "new_skill", "dropped_skill"])
def test_ai_rejects_invented_or_lost_candidate_facts(profile, job, fault):
    original = build_pack(job, profile)
    revision = _revision(original)
    bullet = revision["experiences"][0]["bullets"][0]
    if fault == "employer":
        revision["experiences"][0]["company"] = "Invented Employer"
    elif fault == "education":
        revision["education"] = [{"diploma": "PhD", "school": "Invented University"}]
    elif fault == "new_tool":
        bullet["text"] += " Kubernetes."
    elif fault == "metric":
        bullet["text"] = bullet["text"].replace("2 heures", "1 heure")
    elif fault == "swapped_metrics":
        bullet["text"] = bullet["text"].replace("de 4 à 2", "de 2 à 4")
    elif fault == "unknown_source":
        bullet["source_ids"] = ["unknown"]
    elif fault == "other_role":
        bullet["source_ids"] = ["experience-1:bullet-0"]
    elif fault == "omitted_bullet":
        revision["experiences"][0]["bullets"].pop()
    elif fault == "new_skill":
        revision["skill_order"] = original["cv"]["skills"].split(" · ") + ["Kubernetes"]
    else:
        revision["skill_order"] = ["Spark"]
    baseline = copy.deepcopy(original)
    with pytest.raises(ValueError):
        apply_revision(job, profile, original, revision)
    assert original == baseline


@pytest.mark.parametrize("output", ["plain CV text", "{}", '{"summary":{},"summary":{}}', '{"cv":{"name":"Invented Candidate"}}'])
def test_malformed_or_invented_model_response_keeps_complete_fallback(profile, job, output):
    original = build_pack(job, profile)
    with patch.dict("os.environ", {"NAVIN_CAREER_AI": "on"}), patch("navin.career.ai.ask", return_value=(output, "fixture-model")):
        pack = polish_pack(job, profile, original)
    assert pack["cv"] == original["cv"]
    assert pack["cv_text"] == original["cv_text"]
    assert pack["generation"]["mode"] == "deterministic"
    assert pack["generation"]["status"] == "needs_review"
    assert pack["generation"]["warnings"][0].startswith("Révision")
    assert "fixture-model" != pack.get("model")


def test_ai_opt_out_and_sparse_profile_are_honest(job):
    profile = {"display_name": "Camille", "headline": "Data Engineer", "stack": ["Spark"], "document_language": "fr", "ai_assist": False}
    with patch("navin.career.ai.ask") as ask:
        pack = polish_pack(job, profile, build_pack(job, profile))
    ask.assert_not_called()
    assert pack["generation"]["mode"] == "deterministic"
    assert pack["generation"]["status"] == "needs_review"
    assert pack["generation"]["warnings"]
    assert not pack["cv"]["experiences"]
    assert "AWS" not in pack["cv_text"]


def test_editable_word_has_real_headings_and_lists_and_separate_cover(profile, job):
    pack = build_pack(job, profile)
    cv_blob = render_docx(job, profile, pack, document="cv")
    doc = Document(io.BytesIO(cv_blob))
    assert not doc.tables
    assert any(para.style.name == "Heading 1" and "EXPÉRIENCE" in para.text for para in doc.paragraphs)
    assert any(para.style.name == "List Bullet" and "Spark" in para.text for para in doc.paragraphs)
    assert doc.styles["Normal"].font.name == "Calibri"
    assert doc.core_properties.language == "fr-FR"
    assert doc.paragraphs[0].text == "Camille Morel"
    assert "Madame, Monsieur" not in _document_text(cv_blob)
    cover_text = _document_text(render_docx(job, profile, pack, document="cover"))
    assert "Madame, Monsieur" in cover_text
    assert "FORMATION" not in cover_text
    assert "Rivage Études" not in cover_text
    with zipfile.ZipFile(io.BytesIO(cv_blob)) as archive:
        document = archive.read("word/document.xml").decode()
        assert "<w:tbl>" not in document and "<w:drawing>" not in document and "<w:txbxContent>" not in document
        assert "\u2013" not in document and "\u2014" not in document


def test_prepared_application_opportunity_and_each_download_share_revision(profile, job, tmp_path):
    store = CareerStore(tmp_path)
    store.save_profile(profile)
    store.upsert_opportunities([job])
    revision = _revision(build_pack(job, profile))
    with patch.dict("os.environ", {"NAVIN_CAREER_AI": "on"}), patch("navin.career.ai.ask", return_value=(json.dumps(revision), "fixture-model")):
        prepared = prepare_application(store, job["id"])["prepared"]
    opportunity = store.get_opportunity(job["id"])
    application = store.load_applications()[0]
    for key in ("cv", "cv_text", "cover", "summary", "model", "generation", "language", "exports", "ats_score", "ats_requirements"):
        assert prepared[key] == opportunity[key] == application[key]
    for kind in ("docx", "cv_docx", "cover_docx"):
        downloaded = download_pack(store, job["id"], kind)
        assert downloaded["kind"] == kind
        text = _document_text(base64.b64decode(downloaded["data"]))
        assert "Camille Morel" in text
        if kind != "cover_docx":
            assert revision["summary"]["text"] in text


def test_changed_candidate_name_cannot_download_stale_export(profile, job, tmp_path):
    store = CareerStore(tmp_path)
    first = attach_exports(store, job, profile, build_pack(job, profile))
    profile["display_name"] = "Camille Morel Dupont"
    second = attach_exports(store, job, profile, build_pack(job, profile))
    assert first["exports"]["cv_docx"]["file_id"] != second["exports"]["cv_docx"]["file_id"]
    data = store.read_bytes(second["exports"]["cv_docx"]["file_id"])
    assert _document_text(base64.b64decode(data["data"])).startswith("Camille Morel Dupont\n")


@pytest.mark.parametrize("source", ["greenhouse", "linkedin"])
def test_autopilot_open_is_never_recorded_as_submitted(profile, job, tmp_path, source):
    store = CareerStore(tmp_path)
    store.save_profile({**profile, "ai_assist": False, "apply_mode": "autopilot"})
    store.upsert_opportunities([{**job, "source": source}])
    apply_one(store, job["id"])
    for row in (store.get_opportunity(job["id"]), store.load_applications()[0]):
        assert row["stage"] == "ready"
        assert not row.get("applied_at")
        assert "manually" in row["next_action"]


def test_empty_profile_cannot_create_a_ready_or_applied_pack(job, tmp_path):
    store = CareerStore(tmp_path)
    store.upsert_opportunities([job])
    prepared = prepare_application(store, job["id"])["prepared"]
    assert not prepared["pack_ready"]
    assert not prepared["exports"]
    with pytest.raises(CareerError):
        apply_one(store, job["id"])
